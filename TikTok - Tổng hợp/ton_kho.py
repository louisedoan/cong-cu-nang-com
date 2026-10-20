"""Cập nhật tồn kho file QL ("Quản lý kênh Online") từ đơn TikTok.

Mỗi tháng 1 sheet "Tháng N". Mỗi mặt hàng có 4 cột: Nhập | Xuất TikTok | Xuất Shopee/khác | Tồn.
Công cụ chỉ điền "Xuất TikTok" (số chai rời kho theo ngày giao cho đơn vị vận chuyển, gồm cả chai tặng)
và "Số lượng đơn – Tiktok". Tồn tự tính = Tồn hôm trước + Nhập − Xuất TikTok − Xuất Shopee/khác.
Tên sản phẩm TikTok -> cột trong QL lấy từ file "Ghép tên sản phẩm.xlsx" (sửa được bằng Excel).
"""
import calendar
import re
import shutil
from collections import defaultdict
from copy import copy
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

# đơn đặt xong 0-2 ngày mới giao ĐVVC: ngày gửi D chỉ đủ số khi file TikTok có đơn đặt từ ngày D-3
LAG_DAYS = 3
SUB = ['Nhập', 'Xuất TikTok', 'Xuất Shopee/khác', 'Tồn']
OPENING = 'Tồn đầu tháng'
MAP_FILE = 'Ghép tên sản phẩm.xlsx'
CONFIG_FILE = 'duong-dan-file-QL.txt'
BACKUP_DIR = 'Sao lưu QL'

AUTO_FILL = PatternFill('solid', fgColor='E4DFEC')   # tím nhạt: công cụ tự điền
INPUT_FILL = PatternFill('solid', fgColor='FFFF99')  # vàng: tồn kiểm kho đầu tháng, gõ tay
LEGEND = ('Ô VÀNG: gõ tồn kiểm kho đầu tháng  •  Cột TÍM (Xuất TikTok, Số lượng đơn Tiktok): công cụ tự điền, không gõ tay  '
          '•  Tồn tự tính = Tồn hôm trước + Nhập − Xuất TikTok − Xuất Shopee/khác  •  Hàng hoàn về kho: ghi vào cột Nhập')


def norm(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip()


def key(s):
    return norm(s).lower()


# ---------- cấu hình ----------

def ql_path(tool_dir):
    """Đường dẫn file QL: lưu trong duong-dan-file-QL.txt, lần đầu (hoặc khi file bị chuyển chỗ) hỏi bằng hộp chọn file."""
    cfg = tool_dir / CONFIG_FILE
    if cfg.exists():
        p = cfg.read_text(encoding='utf-8').strip()
        if p and Path(p).exists():
            return Path(p)
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    p = filedialog.askopenfilename(title='Chọn file QUẢN LÝ TỒN KHO (QL- kênh Online) — bấm Cancel nếu không cập nhật',
                                   filetypes=[('Excel', '*.xlsx')])
    if not p:
        return None
    cfg.write_text(p, encoding='utf-8')
    return Path(p)


def read_mapping(path):
    wb = load_workbook(path)
    ws = wb.active
    mapping = {}
    for r in range(2, ws.max_row + 1):
        name = norm(ws.cell(r, 1).value)
        if name:
            f = ws.cell(r, 3).value
            mapping[key(name)] = (norm(ws.cell(r, 2).value), f if isinstance(f, (int, float)) else 1)
    return wb, ws, mapping


def add_unmapped(wb, ws, names, path):
    for name in names:
        r = ws.max_row + 1
        ws.cell(r, 1, name)
        ws.cell(r, 3, 1)
        ws.cell(r, 4, 'MỚI — chưa ghép: gõ tên cột QL vào ô bên trái')
        for c in range(1, 5):
            ws.cell(r, c).fill = copy(INPUT_FILL)
    try:
        wb.save(path)
        return True
    except PermissionError:
        return False


# ---------- đọc / tạo sheet tháng ----------

def is_new(ws):
    return norm(ws['A6'].value) == OPENING


def has_data(ws):
    return any(isinstance(c.value, (int, float)) for row in ws.iter_rows(min_row=6, min_col=2) for c in row)


def old_labels(ws):
    labels = []
    for c in range(2, ws.max_column + 1):
        v = norm(ws.cell(4, c).value)
        if 'số lượng đơn' in v.lower():
            break
        if v:
            labels.append(v)
    return labels


def layout(ws):
    """Sheet mẫu mới -> ({khóa mặt hàng: (cột đầu khối, tên)}, cột Số đơn Tiktok, {ngày: dòng})."""
    blocks, order_col = {}, None
    for c in range(2, ws.max_column + 1):
        sub = key(ws.cell(5, c).value)
        if sub == key('Xuất TikTok'):
            label = norm(ws.cell(4, c - 1).value)
            blocks[key(label)] = (c - 1, label)
        elif sub == 'tiktok':
            order_col = c
    rows = {ws.cell(r, 1).value.date(): r for r in range(7, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, datetime)}
    return blocks, order_col, rows


def capture_styles(wb):
    """Lấy định dạng từ sheet mẫu cũ (VD "Tháng 9") để sheet mới trông giống hệt."""
    src = next((ws for ws in wb.worksheets if key(ws['A4'].value) == 'ngày' and not is_new(ws)), None)
    if src is None:
        return None
    total = next((src.cell(r, 1) for r in range(6, src.max_row + 1) if key(src.cell(r, 1).value) == 'tổng cộng'), src['A6'])
    return {'title': src['A2']._style, 'head': src['A4']._style, 'in': src['B5']._style, 'out': src['C5']._style,
            'cell_in': src['B6']._style, 'cell_out': src['C6']._style, 'date': src['A6']._style, 'total': total._style}


def style(cell, st, name, fill=None, fmt=None):
    if st:
        cell._style = copy(st[name])
    if fill is not None:
        cell.fill = copy(fill)
    if fmt:
        cell.number_format = fmt


def create_month(wb, year, month, labels, index, prev, st):
    from openpyxl.utils import get_column_letter as L
    ws = wb.create_sheet(f'Tháng {month}', index)
    ndays = calendar.monthrange(year, month)[1]
    oc = 2 + 4 * len(labels)  # cột Số lượng đơn Shopee
    last_col = oc + 1
    first, tr = 7, 7 + ndays

    ws['A2'] = f'Quản lý kênh Online — Tháng {month}/{year}'
    style(ws['A2'], st, 'title')
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(last_col, 20))
    ws['A3'] = LEGEND
    ws['A3'].font = Font(name='Arial', size=9, italic=True, color='595959')
    ws.cell(4, 1, 'Ngày')
    for r in (4, 5):
        style(ws.cell(r, 1), st, 'head')
    ws.merge_cells('A4:A5')
    ws.column_dimensions['A'].width = 14
    ws.row_dimensions[4].height = 30
    ws.row_dimensions[5].height = 42

    prev_blocks, prev_rows = ({}, {}) if prev is None else layout(prev)[0::2]
    prev_last = max(prev_rows.values()) if prev_rows else None
    for i, label in enumerate(labels):
        c = 2 + 4 * i
        N, X, K, T = (L(c + j) for j in range(4))
        ws.cell(4, c, label)
        for j, sub in enumerate(SUB):
            style(ws.cell(4, c + j), st, 'head')
            cell = ws.cell(5, c + j, sub)
            style(cell, st, 'out' if sub == 'Tồn' else 'in', AUTO_FILL if sub == 'Xuất TikTok' else None)
            ws.column_dimensions[L(c + j)].width = 8
        ws.merge_cells(start_row=4, start_column=c, end_row=4, end_column=c + 3)
        # dòng Tồn đầu tháng: nối từ tồn cuối tháng trước nếu tháng trước cũng theo mẫu mới, không thì gõ tay (ô vàng)
        for j in range(3):
            style(ws.cell(6, c + j), st, 'cell_in')
        opening = ws.cell(6, c + 3)
        if key(label) in prev_blocks:
            pc = L(prev_blocks[key(label)][0] + 3)
            opening.value = f"='{prev.title}'!{pc}{prev_last}"
            style(opening, st, 'cell_out', fmt='0;[Red]-0')
        else:
            style(opening, st, 'cell_out', INPUT_FILL, '0;[Red]-0')
        for d in range(1, ndays + 1):
            r = 6 + d
            for j, sub in enumerate(SUB[:3]):
                style(ws.cell(r, c + j), st, 'cell_in', AUTO_FILL if sub == 'Xuất TikTok' else None)
            ws.cell(r, c + 3, f'=IF({T}{r - 1}="","",{T}{r - 1}+{N}{r}-{X}{r}-{K}{r})')
            style(ws.cell(r, c + 3), st, 'cell_out', fmt='0;[Red]-0')
        for j, col in enumerate((N, X, K)):
            ws.cell(tr, c + j, f'=SUM({col}{first}:{col}{tr - 1})')
            style(ws.cell(tr, c + j), st, 'cell_in')
        ws.cell(tr, c + 3, f'={T}{tr - 1}')  # tồn cuối tháng
        style(ws.cell(tr, c + 3), st, 'cell_out', fmt='0;[Red]-0')
        for r in (6, tr):
            for j in range(4):
                ws.cell(r, c + j).font = Font(name='Arial', size=11, bold=True)

    ws.cell(4, oc, 'Số lượng đơn')
    for j, sub in enumerate(['Shopee', 'Tiktok']):
        style(ws.cell(4, oc + j), st, 'head')
        style(ws.cell(5, oc + j), st, 'in', AUTO_FILL if sub == 'Tiktok' else None)
        ws.cell(5, oc + j).value = sub
        ws.column_dimensions[L(oc + j)].width = 9
        style(ws.cell(6, oc + j), st, 'cell_in')
        for r in range(first, tr):
            style(ws.cell(r, oc + j), st, 'cell_in', AUTO_FILL if sub == 'Tiktok' else None)
        ws.cell(tr, oc + j, f'=SUM({L(oc + j)}{first}:{L(oc + j)}{tr - 1})')
        style(ws.cell(tr, oc + j), st, 'cell_in')
    ws.merge_cells(start_row=4, start_column=oc, end_row=4, end_column=oc + 1)

    ws.cell(6, 1, OPENING)
    style(ws.cell(6, 1), st, 'date')
    ws.cell(6, 1).font = Font(name='Arial', size=11, bold=True)
    for d in range(1, ndays + 1):
        style(ws.cell(6 + d, 1, datetime(year, month, d)), st, 'date')
    ws.cell(tr, 1, 'TỔNG CỘNG')
    style(ws.cell(tr, 1), st, 'total')
    ws.freeze_panes = 'B7'
    return ws


def ensure_month(wb, year, month, extra_labels, st):
    """Trả về sheet "Tháng N" theo mẫu mới. Sheet mẫu cũ còn trống thì đổi sang mẫu mới; mẫu cũ đã có số thì giữ nguyên (None)."""
    name = f'Tháng {month}'
    labels = None
    if name in wb.sheetnames:
        ws = wb[name]
        if is_new(ws):
            return ws
        if has_data(ws):
            return None
        labels, index = old_labels(ws), wb.index(ws)
        wb.remove(ws)
    else:
        index = len(wb.sheetnames)
    pm, py = (month - 1, year) if month > 1 else (12, year - 1)
    prev = wb[f'Tháng {pm}'] if f'Tháng {pm}' in wb.sheetnames and is_new(wb[f'Tháng {pm}']) else None
    if labels is None:
        if prev is not None:
            labels = [label for _, label in sorted(layout(prev)[0].values())]
        else:
            old = [ws for ws in wb.worksheets if key(ws['A4'].value) == 'ngày']
            labels = old_labels(old[-1]) if old else []
    have = {key(x) for x in labels}
    labels += [x for x in extra_labels if key(x) not in have]
    return create_month(wb, year, month, labels, index, prev, st)


# ---------- cập nhật ----------

def cap_nhat(rows, exported, tool_dir, recalc):
    """Điền Xuất TikTok vào file QL. Trả về các dòng thông báo để in ra màn hình."""
    ql = ql_path(tool_dir)
    if ql is None:
        return ['Bỏ qua cập nhật tồn kho (chưa chọn file QL).']
    msgs = [f'File QL: {ql}']
    start = min(r['created'] for r in rows).date() + timedelta(LAG_DAYS)
    end = exported.date() - timedelta(1)  # ngày xuất file chưa hết ngày
    if start > end:
        msgs.append('• Chưa có ngày nào đủ dữ liệu để điền — cần file TikTok có đơn đặt dài hơn 3 ngày.')

    map_path = tool_dir / MAP_FILE
    mwb, mws, mapping = read_mapping(map_path)
    qty, orders, unmapped = defaultdict(float), defaultdict(set), defaultdict(float)
    for r in rows:
        if r['shipped'] is None or not start <= r['shipped'].date() <= end:
            continue
        d = r['shipped'].date()
        orders[d].add(r['order'])
        label, factor = mapping.get(key(r['prod']), ('', 1))
        if label:
            qty[(d, key(label))] += r['qty'] * factor
        else:
            unmapped[r['prod']] += r['qty']

    wb = load_workbook(ql)
    st = capture_styles(wb)
    extra = list(dict.fromkeys(label for label, _ in mapping.values() if label))
    shown = {key(label): label for label in extra}
    days = [start + timedelta(i) for i in range((end - start).days + 1)]
    by_month = defaultdict(list)
    for d in days:
        by_month[(d.year, d.month)].append(d)
    changed = False
    # luôn chuẩn bị sẵn sheet của tháng hiện tại để gõ Nhập / tồn đầu tháng
    for y, m in sorted(set(by_month) | {(exported.year, exported.month)}):
        ds = by_month.get((y, m), [])
        name = f'Tháng {m}'
        was_new = name in wb.sheetnames and is_new(wb[name])
        ws = ensure_month(wb, y, m, extra, st)
        if ws is None:
            if ds:
                msgs.append(f'• Bỏ qua {ds[0]:%d/%m}–{ds[-1]:%d/%m}: sheet "{name}" đang theo mẫu cũ, giữ nguyên số bạn đã nhập.')
            continue
        if not was_new:
            changed = True
            msgs.append(f'• Đã tạo sheet "{name}" theo mẫu mới — nhớ gõ tồn kiểm kho đầu tháng vào các ô VÀNG (nếu có).')
        if not ds:
            continue
        blocks, oc, drows = layout(ws)
        for (d, k), v in qty.items():
            if d in ds and k not in blocks:
                msgs.append(f'• (!) Sheet "Tháng {m}" không có cột "{shown[k]}" — {v:,.0f} chai ngày {d:%d/%m} chưa trừ kho.')
        for d in ds:
            r = drows[d]
            for k, (c, _) in blocks.items():
                ws.cell(r, c + 1).value = qty.get((d, k)) or None
            ws.cell(r, oc).value = len(orders[d])
        changed = True
        msgs.append(f'• Tháng {m}: đã điền Xuất TikTok ngày {ds[0]:%d/%m} → {ds[-1]:%d/%m}.')

    if unmapped:
        # chỉ thêm tên mới; tên đã có dòng (đang để trống cột QL) thì không thêm lại
        saved = add_unmapped(mwb, mws, [n for n in unmapped if key(n) not in mapping], map_path)
        msgs.append('• (!) Sản phẩm TikTok chưa ghép với cột QL (chưa trừ kho): '
                    + '; '.join(f'{n} ({q:,.0f} chai)' for n, q in unmapped.items()))
        msgs.append(f'  → Mở "{MAP_FILE}", gõ tên cột QL cho các dòng vàng, rồi chạy lại.' if saved else
                    f'  → Không ghi được "{MAP_FILE}" (đang mở?). Đóng file rồi chạy lại.')

    if changed:
        backup_dir = tool_dir / BACKUP_DIR
        backup_dir.mkdir(exist_ok=True)
        backup = backup_dir / f'{ql.stem} (trước khi cập nhật {datetime.now():%d.%m.%Y %Hh%M}).xlsx'
        shutil.copy2(ql, backup)
        try:
            wb.save(ql)
        except PermissionError:
            return msgs + [f'(!) Không ghi được file QL — đang mở trong Excel. Đóng file rồi chạy lại.']
        if not recalc(ql):
            msgs.append('(!) Không tự tính được bằng Excel — mở file QL bằng Excel, bấm Lưu.')
        msgs.append(f'• Đã sao lưu bản cũ vào thư mục "{BACKUP_DIR}".')
    msgs.append(f'Lần xuất file TikTok sau: chọn đơn đặt từ ngày {end - timedelta(LAG_DAYS - 1):%d/%m/%Y} đến hôm nay.')
    return msgs
