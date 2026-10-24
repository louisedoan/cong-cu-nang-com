"""Tạo báo cáo bán hàng TikTok Shop theo mặt hàng từ file "Tất cả đơn hàng" xuất từ Seller Center.

Cách dùng: kéo thả file .xlsx vào Tao-bao-cao.bat
       hoặc: python bao_cao_tiktok.py "<đường dẫn file>.xlsx"
Báo cáo được lưu vào thư mục "Báo cáo" nằm cạnh công cụ, tên file ghi khoảng ngày đặt đơn và giờ xuất dữ liệu.
"""
import math
import os
import re
import subprocess
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

# openpyxl đi kèm trong thư mục "thu-vien" để chạy được trên máy chưa cài thư viện này
sys.path.insert(0, str(Path(__file__).resolve().parent / 'thu-vien'))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

import licensing  # cấp phép offline 20 ngày (bản sao trong thu-vien)
import ton_kho

CANCELLED = {'Đã hủy', 'Cancelled'}
RETURN_TYPE = 'Return/Refund'
OUT_DIR = Path(__file__).resolve().parent / 'Báo cáo'

# ---- tên sheet ----
OVERVIEW = 'Tổng quan'
PROD = 'Theo mặt hàng'
RET = 'Đơn hoàn & giao thất bại'
DATA = 'Dữ liệu chi tiết'
HOW = 'Cách tính'

# ---- màu & định dạng ----
NAVY, BLUE, LIGHT, BAND = '1F3864', '1F4E78', 'DDEBF7', 'F5F8FC'
TOTAL, GREEN, TILE, RED, GRAY, LINE = 'D9E1F2', 'E2EFDA', 'F2F6FC', 'C00000', '595959', 'BFBFBF'
NUM = '#,##0;-#,##0;"-"'
PCT = '0.0%;-0.0%;"-"'
MONEY = '#,##0" đ";-#,##0" đ";"-"'
THIN = Side(style='thin', color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def font(size=10, bold=False, color='262626', italic=False, underline=None):
    return Font(name='Arial', size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(color):
    return PatternFill('solid', fgColor=color)


CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
RIGHT = Alignment(horizontal='right', vertical='center')


def num(x):
    return float(x) if x not in (None, '') else 0.0


def short_name(full):
    # Tên listing TikTok thường là "Tên sản phẩm – mô tả quảng cáo": chỉ giữ phần trước dấu "–"
    head = full.split(' – ')[0].strip()
    return head if len(head) >= 10 else full


def read_orders(src):
    # không dùng read_only: với file TikTok, chế độ đó chỉ đọc được cột A
    wb = load_workbook(src, data_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    hdr = list(next(it))
    data = []
    for r in it:
        d = dict(zip(hdr, r))
        if str(d.get('Order ID') or '').isdigit():  # bỏ dòng mô tả cột (dòng 2) và dòng trống
            data.append(d)
    wb.close()
    return data


def to_rows(data):
    """Mỗi dòng SKU -> 1 dòng cho sheet dữ liệu, kèm các cột phân loại."""
    rows = []
    seen_po, seen_o = set(), set()
    for r in data:
        oid = r['Order ID']
        canc = r['Order Status'] in CANCELLED
        group = ('Giao thất bại' if r['Shipped Time'] else 'Hủy trước khi gửi') if canc else 'Bán'
        live = group == 'Bán'
        qty = num(r['Quantity'])
        price = num(r['SKU Unit Original Price'])
        is_ret = r['Cancelation/Return Type'] == RETURN_TYPE
        retq = num(r['Sku Quantity of return']) if is_ret else 0
        paid = num(r['SKU Subtotal After Discount'])
        full = re.sub(r'\s+', ' ', r['Product Name']).strip()
        prod = short_name(full)
        cnt_po = 1 if live and (oid, prod) not in seen_po else 0
        if live:
            seen_po.add((oid, prod))
        # đơn bị hủy đếm ở dòng đầu; đơn còn hiệu lực đếm ở dòng hàng bán (giá > 0) đầu tiên,
        # vì trạng thái phụ ghi theo từng dòng và dòng quà tặng có trạng thái riêng
        counts_here = not live or price > 0
        cnt_o = 1 if counts_here and oid not in seen_o else 0
        if counts_here:
            seen_o.add(oid)
        rows.append(dict(
            order=oid, created=datetime.strptime(r['Created Time'], '%d/%m/%Y %H:%M:%S'),
            status=r['Order Status'], substatus=r['Order Substatus'], rtype=r['Cancelation/Return Type'] or '',
            reason=r['Cancel Reason'] or '', prod=prod, full=full, sku=r['SKU ID'], qty=qty, price=price,
            kind='Tặng' if price == 0 else 'Bán', gross=num(r['SKU Subtotal Before Discount']),
            seller_disc=num(r['SKU Seller Discount']), plat_disc=num(r['SKU Platform Discount']), paid=paid,
            retq=retq, refund=paid * retq / qty if is_ret and qty else 0, group=group,
            state=('Đã hoàn tiền' if is_ret else r['Order Substatus']) if live and price > 0 else '',
            cnt_po=cnt_po, cnt_o=cnt_o,
            # ngày giao cho đơn vị vận chuyển = ngày hàng rời kho (dùng cho tồn kho)
            shipped=datetime.strptime(r['Shipped Time'], '%d/%m/%Y %H:%M:%S') if r['Shipped Time'] else None))
    return rows


# ---- sheet "Dữ liệu chi tiết": (khóa, tiêu đề, độ rộng, định dạng) ----
DATA_COLS = [
    ('order', 'Mã đơn', 20, None), ('created', 'Ngày đặt', 16, 'dd/mm/yyyy hh:mm'),
    ('status', 'Trạng thái', 14, None), ('substatus', 'Trạng thái phụ', 16, None),
    ('rtype', 'Hủy/Hoàn', 13, None), ('reason', 'Lý do hủy', 24, None), ('prod', 'Mặt hàng', 40, None),
    ('full', 'Tên đầy đủ trên TikTok', 50, None), ('sku', 'SKU ID', 21, None), ('qty', 'Số lượng', 9, NUM),
    ('price', 'Đơn giá gốc', 12, NUM), ('kind', 'Bán/Tặng', 10, None), ('gross', 'Tiền hàng gốc', 13, NUM),
    ('seller_disc', 'Giảm giá shop', 13, NUM), ('plat_disc', 'Giảm giá sàn', 13, NUM),
    ('paid', 'Khách trả (sau giảm giá)', 15, NUM), ('retq', 'SL khách trả lại', 11, NUM),
    ('refund', 'Tiền hoàn cho khách', 14, NUM), ('group', 'Phân loại', 17, None),
    ('state', 'Tình trạng doanh thu', 19, None),
    ('cnt_po', 'Cột phụ: đếm đơn theo mặt hàng', 12, None), ('cnt_o', 'Cột phụ: đếm đơn', 12, None),
]
HELPERS = {'cnt_po', 'cnt_o'}
DCOL = {k: get_column_letter(i) for i, (k, *_) in enumerate(DATA_COLS, 1)}

# ---- sheet "Theo mặt hàng" ----
PROD_COLS = [
    ('stt', 'STT', 6), ('prod', 'Mặt hàng', 44), ('orders', 'Số đơn', 10),
    ('qty', 'Tổng chai xuất', 11), ('sold', 'Chai bán', 10), ('gift', 'Chai tặng', 10),
    ('ret', 'Khách trả lại', 11), ('fail', 'Giao thất bại', 11),
    ('rev', 'Doanh thu (khách trả)', 15), ('refund', 'Tiền hoàn cho khách', 14), ('net', 'Doanh thu thuần', 15),
    ('share', 'Tỷ trọng', 9), ('plat', 'Sàn TikTok bù giảm giá', 14),
    ('shop', 'Doanh thu shop sau hoàn (gồm sàn bù)', 16),
]
PIDX = {k: i for i, (k, *_) in enumerate(PROD_COLS, 1)}
PCOL = {k: get_column_letter(i) for k, i in PIDX.items()}
PROD_GROUPS = [('SỐ LƯỢNG (CHAI)', 'qty', 'gift'), ('HÀNG HOÀN VỀ (CHAI)', 'ret', 'fail'),
               ('DOANH THU (VNĐ)', 'rev', 'share'), ('TÍNH CẢ PHẦN SÀN BÙ (VNĐ)', 'plat', 'shop')]

MEANING = {
    'Đã hoàn tất': 'Đơn đã hoàn tất, không bị hoàn — doanh thu gần như chắc chắn',
    'Đã giao': 'Khách đã nhận hàng nhưng còn trong thời hạn trả hàng — có thể còn bị hoàn',
    'Đang trung chuyển': 'Hàng đang trên đường giao — có thể giao thất bại hoặc bị hoàn',
    'Đã hoàn tiền': 'Khách đã trả hàng, đã hoàn tiền — bị trừ hết ở "Tiền hoàn cho khách"',
}
STATE_ORDER = ['Đã hoàn tất', 'Đã giao', 'Đang trung chuyển', 'Đã hoàn tiền']


def rng(key, n):
    c = DCOL[key]
    return f"'{DATA}'!${c}$2:${c}${n}"


def row_height(text, width, base=15):
    # ước lượng số dòng khi xuống dòng: ~0.95 ký tự / 1 đơn vị độ rộng cột (Arial 10)
    return base * max(1, math.ceil(len(text) / (width * 0.95))) + 4


def setup_print(ws, title_rows=None, landscape=True):
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    if title_rows:
        ws.print_title_rows = title_rows


def title(ws, text, sub, col='A'):
    ws[f'{col}1'] = text
    ws[f'{col}1'].font = font(14, True, NAVY)
    ws[f'{col}2'] = sub
    ws[f'{col}2'].font = font(9, color=GRAY, italic=True)
    ws.row_dimensions[1].height = 24


def sheet_data(wb, rows):
    ws = wb.create_sheet(DATA)
    ws.sheet_properties.tabColor = '7F7F7F'
    ws.append([h for _, h, *_ in DATA_COLS])
    for r in rows:
        ws.append([r[k] for k, *_ in DATA_COLS])
    n = len(rows) + 1
    for i, (k, _, w, fmt) in enumerate(DATA_COLS, 1):
        c = ws.cell(1, i)
        c.font, c.fill, c.alignment, c.border = font(10, True, 'FFFFFF'), fill('A6A6A6' if k in HELPERS else BLUE), CENTER, BOX
        dim = ws.column_dimensions[get_column_letter(i)]
        dim.width = w
        if k in HELPERS:  # cột phụ để tính công thức: ẩn đi, bấm dấu + phía trên để mở
            dim.outlineLevel, dim.hidden = 1, True
        for (cell,) in ws.iter_rows(min_row=2, max_row=n, min_col=i, max_col=i):
            cell.font = font()
            if fmt:
                cell.number_format = fmt
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(DATA_COLS))}{n}'
    return n


def sheet_products(wb, rows, n, period):
    ws = wb.create_sheet(PROD)
    ws.sheet_properties.tabColor = BLUE
    ws.sheet_view.showGridLines = False
    title(ws, 'DOANH THU THEO MẶT HÀNG', f'{period}  •  Không tính đơn đã hủy  •  Sắp xếp theo doanh thu thuần')

    # header 2 tầng (dòng 4-5)
    grouped = set()
    for label, a, b in PROD_GROUPS:
        ca, cb = PIDX[a], PIDX[b]
        for ci in range(ca, cb + 1):
            grouped.add(ci)
            ws.cell(4, ci).fill, ws.cell(4, ci).border = fill(BLUE), BOX
        ws.cell(4, ca, label).font, ws.cell(4, ca).alignment = font(10, True, 'FFFFFF'), CENTER
        ws.merge_cells(start_row=4, start_column=ca, end_row=4, end_column=cb)
    for i, (k, h, w) in enumerate(PROD_COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        if i in grouped:
            c = ws.cell(5, i, h)
            c.font, c.fill, c.alignment, c.border = font(10, True, NAVY), fill(LIGHT), CENTER, BOX
        else:
            for rr in (4, 5):
                ws.cell(rr, i).fill, ws.cell(rr, i).border = fill(BLUE), BOX
            c = ws.cell(4, i, h)
            c.font, c.alignment = font(10, True, 'FFFFFF'), CENTER
            ws.merge_cells(start_row=4, start_column=i, end_row=5, end_column=i)
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 42

    # chỉ hiện mặt hàng có đơn bán hoặc giao thất bại, sắp theo doanh thu thuần
    net = defaultdict(float)
    for r in rows:
        if r['group'] != 'Hủy trước khi gửi':
            net[r['prod']] += (r['paid'] - r['refund']) if r['group'] == 'Bán' else 0
    prods = sorted(net, key=lambda p: -net[p])

    first = 6
    last = first + len(prods) - 1
    tr = last + 1
    P = PCOL
    for k, p in enumerate(prods):
        r = first + k
        base = f'{rng("prod", n)},$B{r},{rng("group", n)},"Bán"'
        f = {
            'stt': k + 1, 'prod': p,
            'orders': f'=SUMIFS({rng("cnt_po", n)},{base})',
            'qty': f'=SUMIFS({rng("qty", n)},{base})',
            'sold': f'=SUMIFS({rng("qty", n)},{base},{rng("kind", n)},"Bán")',
            'gift': f'=SUMIFS({rng("qty", n)},{base},{rng("kind", n)},"Tặng")',
            'ret': f'=SUMIFS({rng("retq", n)},{base})',
            'fail': f'=SUMIFS({rng("qty", n)},{rng("prod", n)},$B{r},{rng("group", n)},"Giao thất bại")',
            'rev': f'=SUMIFS({rng("paid", n)},{base})',
            'refund': f'=SUMIFS({rng("refund", n)},{base})',
            'net': f'={P["rev"]}{r}-{P["refund"]}{r}',
            'share': f'=IF(${P["net"]}${tr}=0,0,{P["net"]}{r}/${P["net"]}${tr})',
            'plat': f'=SUMIFS({rng("plat_disc", n)},{base})',
            # phần sàn bù của các dòng bị hoàn = giảm giá sàn * SL trả / SL
            'shop': (f'={P["net"]}{r}+{P["plat"]}{r}-SUMPRODUCT(({rng("prod", n)}=$B{r})*({rng("group", n)}="Bán")'
                     f'*({rng("retq", n)}>0)*{rng("plat_disc", n)}*{rng("retq", n)}'
                     f'/({rng("qty", n)}+({rng("qty", n)}=0)))'),
        }
        for i, (key, *_) in enumerate(PROD_COLS, 1):
            ws.cell(r, i, f[key])
        ws.row_dimensions[r].height = row_height(p, PROD_COLS[1][2])

    ws.cell(tr, 2, 'TỔNG CỘNG')
    ws.cell(tr, 3, f'=SUMIFS({rng("cnt_o", n)},{rng("group", n)},"Bán")')  # số đơn không trùng
    for key, *_ in PROD_COLS[3:]:
        c = P[key]
        ws[f'{c}{tr}'] = f'=SUM({c}{first}:{c}{last})'
    ws.row_dimensions[tr].height = 22

    for r in range(first, tr + 1):
        is_total = r == tr
        for i, (key, *_) in enumerate(PROD_COLS, 1):
            c = ws.cell(r, i)
            c.border = BOX
            c.alignment = LEFT if key == 'prod' else (CENTER if key == 'stt' else RIGHT)
            c.number_format = PCT if key == 'share' else NUM
            color = RED if key == 'refund' else '262626'
            c.font = font(10, is_total or key == 'net', color)
            if is_total:
                c.fill = fill(TOTAL)
            elif key == 'net':
                c.fill = fill(GREEN)
            elif (r - first) % 2:
                c.fill = fill(BAND)
        if is_total:
            for i in range(1, len(PROD_COLS) + 1):
                ws.cell(r, i).border = Border(left=THIN, right=THIN, top=Side(style='medium', color=NAVY),
                                              bottom=Side(style='medium', color=NAVY))
    note = tr + 2
    ws.cell(note, 2, 'Số liệu TẠM TÍNH tại thời điểm xuất file; chưa trừ phí sàn, hoa hồng, thuế. '
                     'Giải thích từng cột: xem sheet "Cách tính".').font = font(9, color=GRAY, italic=True)
    ws.freeze_panes = 'C6'
    ws.auto_filter.ref = f'A5:{get_column_letter(len(PROD_COLS))}{last}'
    setup_print(ws, '4:5')
    return tr


def sheet_returns(wb, rows, period):
    ws = wb.create_sheet(RET)
    ws.sheet_properties.tabColor = RED
    ws.sheet_view.showGridLines = False
    title(ws, 'ĐƠN HOÀN & GIAO THẤT BẠI', f'{period}  •  Khách trả hàng (đã hoàn tiền) và đơn giao thất bại (hàng quay về kho)')
    cols = [('STT', 6), ('Loại', 16), ('Mã đơn', 21), ('Ngày đặt', 16), ('Mặt hàng', 44), ('Số chai', 9),
            ('Tiền hoàn cho khách', 15), ('Ghi chú', 46)]
    for i, (h, w) in enumerate(cols, 1):
        c = ws.cell(4, i, h)
        c.font, c.fill, c.alignment, c.border = font(10, True, 'FFFFFF'), fill(BLUE), CENTER, BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[4].height = 30

    items = []
    for r in rows:
        name = r['prod'] + (' (quà tặng)' if r['kind'] == 'Tặng' else '')
        if r['group'] == 'Bán' and r['retq'] > 0:
            items.append((0, r['created'], 'Khách trả hàng', r['order'], name, r['retq'], r['refund'],
                          'Đã hoàn tiền cho khách'))
        elif r['group'] == 'Giao thất bại':
            items.append((1, r['created'], 'Giao thất bại', r['order'], name, r['qty'], 0,
                          f'{r["reason"]} — đơn không tính doanh thu, hàng quay về kho'))
    items.sort(key=lambda x: (x[0], x[1]))

    first = 5
    if not items:
        ws.cell(first, 2, 'Không có đơn hoàn hay giao thất bại trong file này.').font = font(10, italic=True)
        return
    for k, (_, created, kind, oid, name, q, money, note) in enumerate(items):
        r = first + k
        vals = [k + 1, kind, oid, created, name, q, money, note]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.font = font(10, color=RED if i == 7 else '262626')
            c.border = BOX
            c.alignment = CENTER if i in (1, 2, 3, 4) else (RIGHT if i in (6, 7) else LEFT)
            if (r - first) % 2:
                c.fill = fill(BAND)
        ws.cell(r, 4).number_format = 'dd/mm/yyyy hh:mm'
        ws.cell(r, 6).number_format = NUM
        ws.cell(r, 7).number_format = NUM
        ws.row_dimensions[r].height = max(row_height(name, 44), row_height(note, 46))
    tr = first + len(items)
    ws.cell(tr, 5, 'TỔNG CỘNG')
    ws.cell(tr, 6, f'=SUM(F{first}:F{tr - 1})')
    ws.cell(tr, 7, f'=SUM(G{first}:G{tr - 1})')
    for i in range(1, len(cols) + 1):
        c = ws.cell(tr, i)
        c.font, c.fill = font(10, True, RED if i == 7 else '262626'), fill(TOTAL)
        c.border = Border(left=THIN, right=THIN, top=Side(style='medium', color=NAVY), bottom=Side(style='medium', color=NAVY))
        c.number_format = NUM
        c.alignment = RIGHT if i in (6, 7) else LEFT
    ws.cell(tr + 2, 2, 'Quà tặng trong đơn khách trả hàng không được TikTok ghi nhận là trả lại nên không có trong danh sách.'
            ).font = font(9, color=GRAY, italic=True)
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f'A4:H{tr - 1}'
    setup_print(ws, '4:4')


def sheet_overview(ws, rows, n, tr, period, exported):
    ws.title = OVERVIEW
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 2
    for c in 'BCDE':
        ws.column_dimensions[c].width = 28
    ws.column_dimensions['F'].width = 66
    ws['B2'] = 'BÁO CÁO BÁN HÀNG TIKTOK SHOP'
    ws['B2'].font = font(18, True, NAVY)
    ws.row_dimensions[2].height = 30
    ws['B3'] = f'{period}  •  Dữ liệu xuất lúc {exported:%d/%m/%Y %H:%M}  •  Đơn vị tiền: VNĐ'
    ws['B3'].font = font(10, color=GRAY, italic=True)

    P = lambda key: f"='{PROD}'!{PCOL[key]}{tr}"
    cnt = lambda col, val: f'=SUMIFS({rng("cnt_o", n)},{rng(col, n)},"{val}")'
    tiles = [
        [('DOANH THU THUẦN', P('net'), MONEY, GREEN, NAVY),
         ('Doanh thu (khách trả)', P('rev'), MONEY, TILE, NAVY),
         ('Tiền hoàn cho khách', P('refund'), MONEY, TILE, RED),
         ('Doanh thu shop (gồm sàn bù)', P('shop'), MONEY, TILE, NAVY)],
        [('Số đơn (không tính đơn hủy)', P('orders'), NUM, TILE, NAVY),
         ('Tổng chai xuất', P('qty'), NUM, TILE, NAVY),
         ('Chai bán', P('sold'), NUM, TILE, NAVY),
         ('Chai tặng', P('gift'), NUM, TILE, NAVY)],
        [('Đơn khách trả hàng', cnt('state', 'Đã hoàn tiền'), NUM, TILE, RED),
         ('Đơn giao thất bại', cnt('group', 'Giao thất bại'), NUM, TILE, RED),
         ('Đơn hủy trước khi gửi', cnt('group', 'Hủy trước khi gửi'), NUM, TILE, GRAY),
         ('Sàn TikTok bù giảm giá', P('plat'), MONEY, TILE, NAVY)],
    ]
    r = 5
    for line in tiles:
        for j, (label, formula, fmt, bg, color) in enumerate(line):
            col = 2 + j
            top, val = ws.cell(r, col, label), ws.cell(r + 1, col, formula)
            top.font, top.fill, top.alignment = font(9, True, GRAY), fill(bg), Alignment(horizontal='left', vertical='bottom', indent=1)
            val.font, val.fill, val.alignment = font(16, True, color), fill(bg), Alignment(horizontal='left', vertical='center', indent=1)
            val.number_format = fmt
            top.border = Border(left=THIN, right=THIN, top=THIN)
            val.border = Border(left=THIN, right=THIN, bottom=THIN)
        ws.row_dimensions[r].height = 20
        ws.row_dimensions[r + 1].height = 34
        ws.row_dimensions[r + 2].height = 8
        r += 3

    # doanh thu theo tình trạng
    r += 1
    ws.cell(r, 2, 'DOANH THU THEO TÌNH TRẠNG ĐƠN').font = font(12, True, NAVY)
    ws.cell(r + 1, 2, 'Cho biết phần doanh thu nào đã chắc, phần nào còn có thể bị hoàn').font = font(9, color=GRAY, italic=True)
    h = r + 2
    for i, t in enumerate(['Tình trạng', 'Số đơn', 'Doanh thu (khách trả)', 'Tỷ trọng', 'Ý nghĩa'], 2):
        c = ws.cell(h, i, t)
        c.font, c.fill, c.alignment, c.border = font(10, True, 'FFFFFF'), fill(BLUE), CENTER, BOX
    ws.row_dimensions[h].height = 24
    found = {x['state'] for x in rows if x['state']}
    states = [s for s in STATE_ORDER if s in found] + sorted(found - set(STATE_ORDER))
    tot = h + len(states) + 1
    for k, st in enumerate(states):
        rr = h + 1 + k
        vals = [st, f'=SUMIFS({rng("cnt_o", n)},{rng("state", n)},$B{rr})',
                f'=SUMIFS({rng("paid", n)},{rng("state", n)},$B{rr})', f'=IF($D${tot}=0,0,D{rr}/$D${tot})',
                MEANING.get(st, 'Chưa gửi hàng hoặc trạng thái khác — chưa nên tính là doanh thu')]
        for i, v in enumerate(vals, 2):
            c = ws.cell(rr, i, v)
            c.font, c.border = font(10, i == 2), BOX
            c.alignment = LEFT if i in (2, 6) else RIGHT
            c.number_format = PCT if i == 5 else NUM
            if k % 2:
                c.fill = fill(BAND)
        ws.row_dimensions[rr].height = 22
    ws.cell(tot, 2, 'TỔNG')
    for i, col in ((3, 'C'), (4, 'D'), (5, 'E')):
        ws.cell(tot, i, f'=SUM({col}{h + 1}:{col}{tot - 1})')
    for i in range(2, 7):
        c = ws.cell(tot, i)
        c.font, c.fill, c.number_format = font(10, True), fill(TOTAL), PCT if i == 5 else NUM
        c.alignment = LEFT if i == 2 else RIGHT
        c.border = Border(left=THIN, right=THIN, top=Side(style='medium', color=NAVY), bottom=Side(style='medium', color=NAVY))
    ws.row_dimensions[tot].height = 22

    # liên kết tới các sheet khác
    r = tot + 2
    ws.cell(r, 2, 'XEM CHI TIẾT').font = font(12, True, NAVY)
    links = [(PROD, 'Doanh thu, số chai bán / tặng / hoàn của từng mặt hàng'),
             (RET, 'Danh sách đơn khách trả hàng và đơn giao thất bại'),
             (DATA, 'Toàn bộ dòng đơn hàng gốc (có thể lọc, tra cứu theo mã đơn)'),
             (HOW, 'Giải thích cách tính từng con số')]
    for k, (sheet, desc) in enumerate(links, 1):
        c = ws.cell(r + k, 2, f'→ {sheet}')
        # liên kết nội bộ phải dùng "location"; gán chuỗi "#..." sẽ thành link ngoài và Excel báo file hỏng
        c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{sheet}'!A1")
        c.font = font(10, True, '0563C1', underline='single')
        ws.cell(r + k, 3, desc).font = font(10, color=GRAY)
    r += len(links) + 2
    ws.cell(r, 2, 'Lưu ý: Doanh thu là số TẠM TÍNH tại lúc xuất file — đơn đang giao hoặc mới giao vẫn có thể bị hoàn. '
                  'Chưa trừ phí sàn, hoa hồng, thuế.').font = font(9, True, RED, italic=True)
    setup_print(ws)


def sheet_howto(wb, src, exported, period):
    ws = wb.create_sheet(HOW)
    ws.sheet_properties.tabColor = '7F7F7F'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 100
    ws['B1'] = 'CÁCH TÍNH CÁC CON SỐ TRONG BÁO CÁO'
    ws['B1'].font = font(14, True, NAVY)
    sections = [
        ('NGUỒN DỮ LIỆU', [
            ('File gốc', f'{src.name} — xuất từ TikTok Shop Seller Center lúc {exported:%d/%m/%Y %H:%M}'),
            ('Khoảng thời gian', period),
            ('Phạm vi', 'Không tính đơn đã hủy. Mỗi dòng trong file gốc là 1 mặt hàng trong 1 đơn.'),
        ]),
        ('CÁC CỘT TRONG BẢNG "THEO MẶT HÀNG"', [
            ('Số đơn', 'Số đơn có chứa mặt hàng đó (kể cả đơn chỉ được tặng). Một đơn nhiều mặt hàng được đếm ở mỗi '
                       'mặt hàng, nên dòng TỔNG CỘNG là số đơn thực (không trùng).'),
            ('Chai bán / Chai tặng', 'Chai tặng = dòng có đơn giá gốc 0đ (VD: mua Long Ngư 2L được tặng chai 60ml). '
                                     'Chai bán = dòng có giá > 0.'),
            ('Khách trả lại', 'Số chai khách trả hàng sau khi nhận (đơn loại Return/Refund).'),
            ('Giao thất bại', 'Đơn bị hủy SAU khi đã gửi hàng (giao không thành công) — hàng quay về kho, '
                              'không có tiền, không tính doanh thu.'),
            ('Doanh thu (khách trả)', 'Tiền hàng khách trả sau mọi giảm giá. Không gồm phí ship.'),
            ('Tiền hoàn cho khách', 'Tiền hàng đã hoàn lại cho khách trả hàng. Đơn hủy COD trong file TikTok vẫn ghi '
                                    '"tiền hoàn" nhưng khách chưa trả tiền nên KHÔNG tính.'),
            # câu chữ không được bắt đầu bằng "=": openpyxl sẽ coi là công thức và Excel báo file hỏng
            ('Doanh thu thuần', 'Bằng Doanh thu (khách trả) − Tiền hoàn cho khách'),
            ('Tỷ trọng', 'Phần trăm doanh thu thuần của mặt hàng so với tổng.'),
            ('Sàn TikTok bù giảm giá', 'Phần voucher do sàn TikTok chịu — sàn trả lại cho shop khi quyết toán.'),
            ('Doanh thu shop sau hoàn', 'Bằng Doanh thu thuần + Sàn bù − phần sàn bù của các đơn bị hoàn.'),
        ]),
        ('LƯU Ý', [
            ('Số tạm tính', 'Đơn "Đang trung chuyển" và "Đã giao" vẫn có thể bị hoàn sau ngày xuất file. '
                            'Muốn chốt chắc, xuất lại file sau khoảng 15 ngày.'),
            ('Chưa trừ phí', 'Chưa trừ phí sàn, hoa hồng, thuế (file đơn hàng không có). Tiền thực nhận xem ở '
                             'Seller Center > Tài chính > Thu nhập.'),
            ('Tên mặt hàng', 'Rút gọn = phần trước dấu "–" trong tên listing. Tên đầy đủ xem ở sheet "Dữ liệu chi tiết".'),
            ('Mặt hàng không hiện', 'Mặt hàng chỉ có đơn hủy trước khi gửi sẽ không hiện trong bảng.'),
            ('Trạng thái phụ', 'TikTok ghi trạng thái theo từng dòng: dòng quà tặng thường là "Đã hoàn tất" dù hàng bán '
                               'cùng đơn vẫn "Đã giao". Vì vậy bảng tình trạng chỉ dựa vào dòng hàng bán.'),
        ]),
    ]
    r = 3
    for head, items in sections:
        ws.cell(r, 2, head).font = font(11, True, 'FFFFFF')
        for i in (2, 3):
            ws.cell(r, i).fill = fill(BLUE)
        ws.row_dimensions[r].height = 22
        r += 1
        for k, (name, text) in enumerate(items):
            a, b = ws.cell(r, 2, name), ws.cell(r, 3, text)
            a.font, b.font = font(10, True, NAVY), font(10)
            a.alignment, b.alignment = LEFT, LEFT
            a.border = b.border = BOX
            if k % 2:
                a.fill = b.fill = fill(BAND)
            ws.row_dimensions[r].height = row_height(text, 100, 14)
            r += 1
        r += 1
    setup_print(ws)
    ws.page_setup.fitToHeight = 1  # in vừa 1 trang


def export_time(src):
    m = re.search(r'(\d{4})-(\d{2})-(\d{2})-(\d{2})_(\d{2})', src.stem)  # VD: ...-2026-10-01-13_26
    return datetime(*map(int, m.groups())) if m else datetime.fromtimestamp(src.stat().st_mtime)


def build_report(rows, src, exported, out):
    created = [r['created'] for r in rows]
    period = f'Đơn đặt từ {min(created):%d/%m/%Y} đến {max(created):%d/%m/%Y}'
    wb = Workbook()
    overview = wb.active
    n = sheet_data(wb, rows)
    tr = sheet_products(wb, rows, n, period)
    sheet_returns(wb, rows, period)
    sheet_overview(overview, rows, n, tr, period, exported)
    sheet_howto(wb, src, exported, period)
    wb.move_sheet(DATA, offset=len(wb.sheetnames))  # dữ liệu chi tiết nằm cuối
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)


def recalc_with_excel(path):
    """Cho Excel tính sẵn công thức rồi lưu, để file xem được cả trên điện thoại / Zalo."""
    ps = ('$x = New-Object -ComObject Excel.Application; $x.DisplayAlerts = $false; '
          'try { $w = $x.Workbooks.Open($env:BAO_CAO_PATH); $x.CalculateFull(); $w.Save(); $w.Close($false) } '
          'finally { $x.Quit(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($x) }')
    res = subprocess.run(['powershell', '-NoProfile', '-NonInteractive', '-Command', ps],
                         env=dict(os.environ, BAO_CAO_PATH=str(path)), capture_output=True)
    return res.returncode == 0


def pick_file():
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    return filedialog.askopenfilename(title='Chọn file "Tất cả đơn hàng" xuất từ TikTok',
                                      filetypes=[('Excel', '*.xlsx')])


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    if not licensing.ensure_licensed('Báo cáo TikTok'):
        return
    src = sys.argv[1] if len(sys.argv) > 1 else pick_file()
    if not src:
        print('Chưa chọn file.')
        return
    src = Path(src)
    print(f'Đang đọc: {src.name}')
    rows = to_rows(read_orders(src))
    exported = export_time(src)
    created = [r['created'] for r in rows]
    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / (f'Báo cáo TikTok {min(created):%d.%m.%Y} - {max(created):%d.%m.%Y} '
                     f'(xuất {exported:%d.%m.%Y %Hh%M}).xlsx')
    print('Đang tạo báo cáo...')
    try:
        build_report(rows, src, exported, out)
    except PermissionError:
        print(f'Không ghi được — file "{out.name}" đang mở trong Excel. Đóng file rồi chạy lại.')
        return
    print('Đang cho Excel tính sẵn số liệu...')
    if not recalc_with_excel(out):
        print('(!) Không tự tính được bằng Excel. Hãy mở file bằng Excel, bấm Lưu rồi mới gửi đi.')
    print()
    print('XONG! Báo cáo đã lưu tại thư mục:')
    print(f'  {OUT_DIR}')
    print(f'  Tên file: {out.name}')
    print()
    print('Đang cập nhật tồn kho...')
    for line in ton_kho.cap_nhat(rows, exported, OUT_DIR.parent, recalc_with_excel):
        print(f'  {line}')
    subprocess.Popen(f'explorer /select,"{out}"')
    os.startfile(out)


if __name__ == '__main__':
    main()
