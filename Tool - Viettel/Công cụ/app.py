"""Công cụ Quản lý Kênh Viettel (Bưu cục + Thương Mại).

Chạy: nhấp đúp "Mo-cong-cu.bat" -> mở giao diện trên trình duyệt (chỉ máy này truy cập được).
- Lần đầu: tự đọc các file Excel trong thư mục cha (D:\\Công cụ Nàng Cơm\\Tool - Viettel):
  danh mục & giá & hoa hồng & tỉnh (sheet "Danh Mục Thiết Lập"), bảng tải trọng và mẫu đơn.
- Mọi dữ liệu lưu trong 1 file Excel: "Dữ liệu\\Dữ liệu Viettel.xlsx" (mở xem được; trước mỗi lần lưu
  tự sao lưu bản cũ vào "Dữ liệu\\Sao lưu").
- File xuất ra (đơn + phiếu xuất kho, báo cáo) nằm trong "Xuất file".

Mọi phép tính tiền đều do máy chủ (file này) làm; giao diện chỉ hiển thị.
"""
import json
import math
import os
import re
import shutil
import socket
import subprocess
import sys
import threading
import unicodedata
import urllib.request
import webbrowser
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parent
SOURCE_DIR = ROOT.parent  # thư mục chứa các file Excel gốc
sys.path.insert(0, str(ROOT / 'thu-vien'))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import licensing  # noqa: E402  (cấp phép offline 48 giờ — bản sao trong thu-vien, xem README)

DATA_DIR = ROOT / 'Dữ liệu'
DATA_FILE = DATA_DIR / 'Dữ liệu Viettel.xlsx'
BACKUP_DIR = DATA_DIR / 'Sao lưu'
TEMPLATE_DIR = DATA_DIR / 'Mẫu'
OUT_DIR = ROOT / 'Xuất file'
WEB_DIR = ROOT / 'giao-dien'
PORT = 8766

CHANNELS = {'buu_cuc': 'Bưu cục Viettel', 'tm': 'Thương Mại'}
# trạng thái đơn (dùng chung; TM dùng thêm trạng thái thanh toán)
STATUS = {'da_len_don': 'Đã lên đơn', 'giao_mot_phan': 'Giao một phần', 'da_giao_du': 'Đã giao đủ',
          'huy': 'Đã hủy'}
PAY_STATUS = {'chua_tt': 'Chưa thanh toán', 'da_tt': 'Thanh toán'}
DEFAULT_SETTINGS = {  # tỷ lệ thật lấy từ file danh mục lúc nhập lần đầu, sửa được ở Cài đặt
    'cp_pct': 0.0,     # hoa hồng người giới thiệu
    'tinh_pct': 0.0,   # hoa hồng bưu cục / tỉnh
}
# cài đặt dạng chữ: tên người nhận hoa hồng giới thiệu lấy từ file danh mục, chỉ lưu trong file dữ liệu trên máy
TEXT_SETTINGS = {'ten_hh_cp': 'Người giới thiệu'}
LOCK = threading.Lock()
EXCEL_TIMEOUT = 300  # giây: Excel treo thì báo lỗi thay vì khóa công cụ mãi
MAX_BODY = 5 * 1024 * 1024  # giới hạn dữ liệu gửi lên máy chủ
APP_ID = 'viettel'  # nhận ra chính công cụ này khi đang chạy ở cửa sổ khác
LICENSE = licensing.LicenseClient()
DATA_MTIME = [0.0]


# =====================================================================
# tiện ích
# =====================================================================

def norm(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip()


def key(s):
    """Khóa so khớp tên (bỏ dấu, chữ thường, bỏ khoảng trắng thừa, bỏ dấu chấm cuối)."""
    s = norm(s).lower().rstrip(' .')
    s = unicodedata.normalize('NFD', s)
    return ''.join(ch for ch in s if unicodedata.category(ch) != 'Mn').replace('đ', 'd').strip()


def num(v):
    if v in (None, '', '-'):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', '').strip())
    except ValueError:
        return 0.0


def opt_num(v):
    return None if v in (None, '') else num(v)


def money(v):
    return int(round(v + 1e-9))


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and v:
        for fmt in ('%Y-%m-%d', '%d/%m/%Y'):
            try:
                return datetime.strptime(v[:10], fmt).date()
            except ValueError:
                pass
    return None


def iso(d):
    return d.isoformat() if d else None


def today():
    return date.today()


# =====================================================================
# đọc file Excel gốc (lần đầu)
# =====================================================================

VOUCHER_ID = 'SP-PHIEU'
VOUCHER_TEN = 'Phiếu mua hàng đổi quà Nước Chấm Long Ngư 1 Lít'


def _sheet(wb, name):
    for s in wb.sheetnames:
        if norm(s) == norm(name):
            return wb[s]
    raise KeyError(name)


def _has_sheet(names, name):
    return any(norm(s) == norm(name) for s in names)


def find_source_files():
    """(file danh mục có sheet 'Danh Mục Thiết Lập' - ưu tiên bản TM; file mẫu có sheet
    'Đơn Hàng Viettel' + 'Đối chiếu tải trọng')."""
    cat = tpl = None
    for f in sorted(SOURCE_DIR.glob('*.xlsx')):
        if f.name.startswith('~$'):
            continue
        try:
            names = load_workbook(f, read_only=True).sheetnames
        except Exception:
            continue
        if _has_sheet(names, 'Đơn Hàng Viettel') and _has_sheet(names, 'Đối chiếu tải trọng') and tpl is None:
            tpl = f
        if _has_sheet(names, 'Danh Mục Thiết Lập') and (cat is None or 'TM' in f.name.upper()):
            cat = f
    if not cat or not tpl:
        raise RuntimeError('Không tìm thấy file danh mục (sheet "Danh Mục Thiết Lập") hoặc file mẫu '
                           f'(sheet "Đơn Hàng Viettel" + "Đối chiếu tải trọng") trong {SOURCE_DIR}')
    return cat, tpl


def _header_cols(ws, must='Mã sản phẩm', max_row=8):
    for r in range(1, max_row + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if must in vals:
            cols = {}
            for i, v in enumerate(vals, 1):
                if v and v not in cols:
                    cols[v] = i
            return r, cols
    return None, {}


def is_combo_name(name):
    n = norm(name).lower()
    return '+' in n or 'hộp quà' in n or 'combo' in n


def is_combo_dvt(dvt):
    """Combo BOM = ĐVT 'Phần' hoặc 'Hộp' (combo bán theo phần/hộp, cần tách ở phiếu kho).
    'Combo nước tương...' có ĐVT 'LÔ 2' -> là mặt hàng lẻ, KHÔNG phải combo."""
    return key(dvt) in ('phan', 'hop')


def import_catalog(path):
    """Đọc sheet 'Danh Mục Thiết Lập': sản phẩm + giá vốn/bán, hoa hồng, tỉnh, trạng thái."""
    wb = load_workbook(path, data_only=True)
    ws = _sheet(wb, 'Danh Mục Thiết Lập')
    hr, cols = _header_cols(ws)
    if not hr:
        raise RuntimeError('Sheet "Danh Mục Thiết Lập" không có dòng tiêu đề "Mã sản phẩm".')
    c_ma, c_ten, c_qc = cols.get('Mã sản phẩm'), cols.get('Tên sản phẩm'), cols.get('Quy cách')
    c_gv, c_gb = cols.get('Giá vốn tham chiếu'), cols.get('Giá bán tham chiếu')
    c_ctkm = cols.get('Chương trình Khuyến mãi')
    products = []
    for r in range(hr + 1, ws.max_row + 1):
        ten = norm(ws.cell(r, c_ten).value) if c_ten else ''
        if not ten:
            break
        ma = norm(ws.cell(r, c_ma).value) if c_ma else ''
        if not re.fullmatch(r'\d{6,}', ma):  # cột B đôi khi là ghi chú CTKM, không phải mã vạch
            ma = ''
        products.append({'id': f'SP{len(products) + 1:02d}', 'ma_vach': ma, 'ten': ten, 'dvt': '',
                         'quy_cach': int(num(ws.cell(r, c_qc).value)) or 1 if c_qc else 1,
                         'gia_von': num(ws.cell(r, c_gv).value) if c_gv else 0.0,
                         'gia_ban': num(ws.cell(r, c_gb).value) if c_gb else 0.0,
                         'ctkm': norm(ws.cell(r, c_ctkm).value) if c_ctkm else '',
                         'la_combo': is_combo_name(ten), 'dang_ban': True})
    commission = {}
    c_doi, c_ty = cols.get('Đối tượng nhận hoa hồng'), cols.get('Tỷ lệ hoa hồng (%)')
    if c_doi and c_ty:
        for r in range(hr + 1, ws.max_row + 1):
            doi = norm(ws.cell(r, c_doi).value).lower()
            if not doi:
                continue
            rate = num(ws.cell(r, c_ty).value)
            if 'bưu cục' in doi or 'tỉnh' in doi:
                commission['tinh_pct'] = rate
            elif 'cp_pct' not in commission:  # người giới thiệu (tên lấy từ file, không ghi trong mã nguồn)
                commission['cp_pct'] = rate
                commission['ten_hh_cp'] = norm(ws.cell(r, c_doi).value)
    provinces = []
    c_t = cols.get('Tỉnh/Thành phố')
    if c_t:
        for r in range(hr + 1, ws.max_row + 1):
            v = norm(ws.cell(r, c_t).value)
            if v:
                provinces.append(v)
    statuses = []
    c_s = cols.get('Trạng thái')
    if c_s:
        for r in range(hr + 1, ws.max_row + 1):
            v = norm(ws.cell(r, c_s).value)
            if v:
                statuses.append(v)
    wb.close()
    return products, commission, provinces, statuses


def import_weights(path):
    """Sheet 'Đối chiếu tải trọng': tên hàng -> {quy cách (tải trọng), KG/thùng}."""
    wb = load_workbook(path, data_only=True)
    ws = _sheet(wb, 'Đối chiếu tải trọng')
    out = {}
    for r in range(1, ws.max_row + 1):
        ten = norm(ws.cell(r, 1).value)
        if not ten or ten.upper().startswith('TÊN') or ten.upper().startswith('DANH'):
            continue
        qc, kg = ws.cell(r, 3).value, ws.cell(r, 4).value
        if qc or kg:
            out[key(ten)] = {'qc': num(qc), 'kg': num(kg), 'ten': ten}
    wb.close()
    return out


def import_order_rows(path):
    """Sheet 'Đơn Hàng Viettel': khóa tên hàng -> số dòng trên mẫu, và khóa -> ĐVT (cột E)."""
    wb = load_workbook(path, data_only=True)
    ws = _sheet(wb, 'Đơn Hàng Viettel')
    rows, dvt = {}, {}
    for r in range(1, ws.max_row + 1):
        stt = ws.cell(r, 1).value
        ten = norm(ws.cell(r, 3).value)
        if ten and stt not in (None, '') and norm(stt) != 'STT':
            rows.setdefault(key(ten), r)
            dvt.setdefault(key(ten), norm(ws.cell(r, 5).value))
    wb.close()
    return rows, dvt


def match_product(frag, products):
    """So khớp một mẩu tên trong combo với 1 sản phẩm (theo số từ chung nhiều nhất)."""
    ft = set(key(frag).split())
    ft.discard('')
    if not ft:
        return None
    best, best_score = None, 0
    for p in products:
        if p.get('la_combo'):  # thành phần/hàng tặng luôn là SP lẻ, không phải combo
            continue
        score = len(ft & set(key(p['ten']).split()))
        if score > best_score and score >= 2 and score >= 0.5 * len(ft):
            best, best_score = p, score
    return best['id'] if best else None


def parse_combo(name, products):
    """Tách tên combo thành công thức: thành phần, phần chọn 1-trong-N ('hoặc'), hàng tặng ('Tặng')."""
    by_id = {p['id']: p for p in products}
    resolved = lambda frag: by_id[match_product(frag, products)]['ten'] if match_product(frag, products) else frag  # noqa: E731
    t = norm(name)
    head, gift = t, []
    for marker in ('Tặng', 'tặng', 'TẶNG'):
        if marker in t:
            head, _, gtext = t.partition(marker)
            gtext = norm(gtext)
            if gtext:
                sp = match_product(gtext, products)
                gift.append({'sp': sp, 'qty': 1, 'ten': by_id[sp]['ten'] if sp else gtext})
            break
    components, choice = [], None
    for frag in head.split('+'):
        frag = frag.strip(' .')
        if not frag:
            continue
        if 'hoặc' in frag.lower():
            choice = []
            for o in re.split(r'\s*ho[ặă]c\s*', frag, flags=re.I):
                o = o.strip(' .()')
                if o:
                    sp = match_product(o, products)
                    choice.append({'sp': sp, 'ten': by_id[sp]['ten'] if sp else o})
        else:
            qty = 1
            mm = re.match(r'(\d+)\s+', frag)
            if mm:
                qty, frag = int(mm.group(1)), frag[mm.end():].strip()
            sp = match_product(frag, products)
            components.append({'sp': sp, 'qty': qty, 'ten': by_id[sp]['ten'] if sp else frag})
    return {'components': components, 'choice': choice, 'gift': gift}


def _fuzzy_val(ten, mapping):
    """Trả về giá trị của khóa khớp gần đúng nhất với 'ten' (theo số từ chung) trong mapping {key: value}."""
    ft = set(key(ten).split())
    ft.discard('')
    best, best_score = None, 0
    for k, v in mapping.items():
        score = len(ft & set(k.split()))
        if score > best_score and score >= 2 and score >= 0.5 * len(ft):
            best, best_score = v, score
    return best


def _fuzzy_row(ten, order_rows):
    """Khớp tên SP với dòng trên mẫu (dùng cho combo tên lệch)."""
    return _fuzzy_val(ten, order_rows)


def initial_import():
    cat, tpl = find_source_files()
    products, commission, provinces, statuses = import_catalog(cat)
    products.append({'id': VOUCHER_ID, 'ma_vach': '', 'ten': VOUCHER_TEN, 'dvt': 'Phiếu', 'quy_cach': 1,
                     'gia_von': 0.0, 'gia_ban': 0.0, 'ctkm': '', 'la_combo': False, 'dang_ban': False})
    weights_by_name = import_weights(tpl)
    order_rows, dvt_by_key = import_order_rows(tpl)
    for p in products:
        if not p.get('dvt'):
            p['dvt'] = dvt_by_key.get(key(p['ten'])) or _fuzzy_val(p['ten'], dvt_by_key) or ''
    for p in products:  # phân loại combo theo ĐVT (Phần/Hộp) sau khi đã có ĐVT từ mẫu đơn
        if p['id'] != VOUCHER_ID:
            p['la_combo'] = is_combo_dvt(p['dvt'])
    combos = {p['id']: parse_combo(p['ten'], products) for p in products if p['la_combo']}
    weights = {}
    for p in products:
        w = weights_by_name.get(key(p['ten']))
        if w:
            weights[p['id']] = {'quy_cach_tt': w['qc'], 'kg_thung': w['kg']}
    for p in products:
        p['row_mau'] = order_rows.get(key(p['ten'])) or _fuzzy_row(p['ten'], order_rows)
    settings = dict(DEFAULT_SETTINGS, **TEXT_SETTINGS)
    settings.update(commission)
    settings.update(nguon_danh_muc=cat.name, nguon_mau=tpl.name, ngay_import=iso(today()))
    return {'settings': settings, 'products': products, 'combos': combos, 'weights': weights,
            'provinces': provinces, 'statuses': statuses, 'order_rows': order_rows,
            'orders': [], 'deliveries': []}


# =====================================================================
# lưu trữ bằng Excel (1 file, mọi sheet)
# =====================================================================
SETTING_LABELS = {'cp_pct': 'Hoa hồng người giới thiệu (tỷ lệ)', 'ten_hh_cp': 'Tên người nhận hoa hồng giới thiệu',
                  'tinh_pct': 'Hoa hồng Bưu cục/Tỉnh (tỷ lệ)',
                  'nguon_danh_muc': 'File danh mục gốc', 'nguon_mau': 'File mẫu gốc', 'ngay_import': 'Ngày nhập ban đầu'}
SP_COLS = [('id', 'Mã SP'), ('ma_vach', 'Mã vạch'), ('ten', 'Tên hàng'), ('dvt', 'ĐVT'), ('quy_cach', 'Quy cách'),
           ('gia_von', 'Giá vốn'), ('gia_ban', 'Giá bán'), ('ctkm', 'Khuyến mãi'), ('la_combo', 'Là combo'),
           ('dang_ban', 'Đang bán'), ('row_mau', 'Dòng mẫu')]
COMBO_COLS = [('combo_id', 'Mã combo'), ('combo_ten', 'Tên combo'), ('loai', 'Loại thành phần'),
              ('sp', 'Mã SP thành phần'), ('ten', 'Tên thành phần'), ('qty', 'Số lượng/1 combo')]
WEIGHT_COLS = [('sp', 'Mã SP'), ('ten', 'Tên hàng'), ('quy_cach_tt', 'Quy cách (tải trọng)'), ('kg_thung', 'KG/thùng')]
ORDER_COLS = [('id', 'ID'), ('ma', 'Mã đơn'), ('kenh', 'Kênh'), ('ngay_dat', 'Ngày đặt'), ('khach', 'Khách hàng'),
              ('tinh', 'Tỉnh'), ('trang_thai', 'Trạng thái'), ('pay_status', 'Thanh toán'), ('ck_pct', 'Chiết khấu %'),
              ('cong_no', 'Công nợ'), ('phi_vc', 'Phí vận chuyển'), ('chi_phi_khac', 'Chi phí khác'),
              ('ghi_chu', 'Ghi chú'), ('tao_luc', 'Tạo lúc')]
LINE_COLS = [('order_id', 'ID đơn'), ('loai', 'Loại dòng'), ('sp', 'Mã SP'), ('combo_id', 'Mã combo'),
             ('choice_sp', 'SP chọn'), ('thung', 'SL thùng'), ('phan', 'SL phần'), ('gia_chai', 'Giá chai (sửa tay)')]
DELIVERY_COLS = [('order_id', 'ID đơn'), ('dot', 'Đợt'), ('ngay', 'Ngày'), ('sp', 'Mã SP'), ('qty', 'Số lượng')]
MONEY_KEYS = {'gia_von', 'gia_ban', 'cong_no', 'phi_vc', 'chi_phi_khac'}
PCT_KEYS = {'ck_pct'}


def _write_sheet(wb, title, cols, rows):
    ws = wb.create_sheet(title)
    for i, (k, h) in enumerate(cols, 1):
        c = ws.cell(1, i, h)
        c.font, c.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F4E78')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(44, len(h) + 4))
    for r, row in enumerate(rows, 2):
        for i, (k, h) in enumerate(cols, 1):
            v = row.get(k)
            if isinstance(v, bool):
                v = 'Có' if v else 'Không'
            c = ws.cell(r, i, v)
            if k in MONEY_KEYS:
                c.number_format = '#,##0'
            elif k in PCT_KEYS and v not in (None, ''):
                c.number_format = '0.##%'
    ws.freeze_panes = 'A2'


def _read_sheet(wb, title, cols):
    if title not in wb.sheetnames:
        return []
    ws = wb[title]
    head = [norm(c.value) for c in ws[1]]
    idx = {k: head.index(h) for k, h in cols if h in head}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, '') for v in row):
            continue
        out.append({k: (row[i] if i < len(row) else None) for k, i in idx.items()})
    return out


def _combo_rows(db):
    byid = {p['id']: p for p in db['products']}
    rows = []
    for cid, bom in db['combos'].items():
        cten = byid.get(cid, {}).get('ten', '')
        for c in bom.get('components', []):
            rows.append({'combo_id': cid, 'combo_ten': cten, 'loai': 'thanh_phan', 'sp': c.get('sp'),
                         'ten': c.get('ten'), 'qty': c.get('qty', 1)})
        for c in (bom.get('choice') or []):
            rows.append({'combo_id': cid, 'combo_ten': cten, 'loai': 'chon', 'sp': c.get('sp'), 'ten': c.get('ten'), 'qty': 1})
        for g in bom.get('gift', []):
            rows.append({'combo_id': cid, 'combo_ten': cten, 'loai': 'tang', 'sp': g.get('sp'),
                         'ten': g.get('ten'), 'qty': g.get('qty', 1)})
    return rows


def _weight_rows(db):
    byid = {p['id']: p for p in db['products']}
    return [{'sp': sp, 'ten': byid.get(sp, {}).get('ten', ''), 'quy_cach_tt': w['quy_cach_tt'], 'kg_thung': w['kg_thung']}
            for sp, w in db['weights'].items()]


def _line_rows(db):
    rows = []
    for o in db['orders']:
        for ln in o.get('lines', []):
            rows.append({'order_id': o['id'], 'loai': 'ban', 'sp': ln.get('sp'), 'thung': ln.get('thung'),
                         'gia_chai': ln.get('gia_chai')})
        for cb in o.get('combos', []):
            rows.append({'order_id': o['id'], 'loai': 'combo', 'combo_id': cb.get('combo_id'),
                         'choice_sp': cb.get('choice_sp'), 'phan': cb.get('phan')})
        for h in o.get('httm', []):
            rows.append({'order_id': o['id'], 'loai': 'httm', 'sp': h.get('sp'), 'thung': h.get('thung')})
    return rows


def save_db(db):
    """Ghi toàn bộ dữ liệu ra 1 file Excel (sao lưu bản trước). File đang mở trong Excel -> PermissionError."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Cài đặt'
    ws.append(['Khóa', 'Giá trị', 'Ý nghĩa'])
    for k, v in db['settings'].items():
        ws.append([k, v, SETTING_LABELS.get(k, '')])
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width, ws.column_dimensions['C'].width = 16, 34, 36
    _write_sheet(wb, 'Sản phẩm', SP_COLS, db['products'])
    _write_sheet(wb, 'Combo', COMBO_COLS, _combo_rows(db))
    _write_sheet(wb, 'Tải trọng', WEIGHT_COLS, _weight_rows(db))
    _write_sheet(wb, 'Tỉnh', [('ten', 'Tỉnh/Thành phố')], [{'ten': t} for t in db['provinces']])
    _write_sheet(wb, 'Trạng thái', [('ten', 'Trạng thái')], [{'ten': t} for t in db.get('statuses', [])])
    _write_sheet(wb, 'Đơn hàng', ORDER_COLS, db['orders'])
    _write_sheet(wb, 'Chi tiết đơn', LINE_COLS, _line_rows(db))
    _write_sheet(wb, 'Đợt giao', DELIVERY_COLS, db['deliveries'])
    tmp = DATA_DIR / '~dang-luu.xlsx'
    wb.save(tmp)
    backup = None
    if DATA_FILE.exists():
        stamp = f'Dữ liệu {datetime.now():%Y-%m-%d %Hh%Mm%Ss}'
        backup, i = BACKUP_DIR / f'{stamp}.xlsx', 1
        while backup.exists():  # lưu 2 lần trong cùng 1 giây: không ghi đè bản sao lưu trước
            backup, i = BACKUP_DIR / f'{stamp} ({i}).xlsx', i + 1
        shutil.copy2(DATA_FILE, backup)
    try:
        os.replace(tmp, DATA_FILE)
    except PermissionError:  # file đang mở trong Excel: bỏ bản sao lưu vừa tạo, KHÔNG xóa bản sao lưu cũ
        tmp.unlink(missing_ok=True)
        if backup:
            backup.unlink(missing_ok=True)
        raise
    for old in sorted(BACKUP_DIR.glob('Dữ liệu *.xlsx'))[:-60]:  # lưu xong mới dọn, giữ 60 bản gần nhất
        old.unlink()
    DATA_MTIME[0] = DATA_FILE.stat().st_mtime


def load_db():
    if not DATA_FILE.exists():
        db = initial_import()
        save_db(db)
        return db
    wb = load_workbook(DATA_FILE, data_only=True)
    settings = dict(DEFAULT_SETTINGS, **TEXT_SETTINGS)
    for row in wb['Cài đặt'].iter_rows(min_row=2, values_only=True):
        if row[0]:
            settings[row[0]] = row[1]
    for k in ('cp_pct', 'tinh_pct'):
        settings[k] = num(settings.get(k))
    for k, v in TEXT_SETTINGS.items():
        settings[k] = norm(settings.get(k)) or v
    products = []
    for p in _read_sheet(wb, 'Sản phẩm', SP_COLS):
        products.append({'id': p['id'], 'ma_vach': norm(p.get('ma_vach')), 'ten': norm(p.get('ten')),
                         'dvt': norm(p.get('dvt')),
                         'quy_cach': int(num(p.get('quy_cach'))) or 1, 'gia_von': num(p.get('gia_von')),
                         'gia_ban': num(p.get('gia_ban')), 'ctkm': norm(p.get('ctkm')),
                         'la_combo': p.get('la_combo') == 'Có', 'dang_ban': p.get('dang_ban') != 'Không',
                         'row_mau': int(num(p.get('row_mau'))) or None})
    for p in products:  # sửa cả file dữ liệu cũ: chỉ HẠ combo có ĐVT rõ là lẻ (CHAI/HŨ/CAN/LÔ...), giữ nếu ĐVT trống
        if p['id'] != VOUCHER_ID and p['dvt'] and not is_combo_dvt(p['dvt']):
            p['la_combo'] = False
    combo_ids = {p['id'] for p in products if p['la_combo']}
    combos = {}
    for r in _read_sheet(wb, 'Combo', COMBO_COLS):
        cid = r.get('combo_id')
        if not cid or cid not in combo_ids:  # bỏ công thức của SP không còn là combo (VD "Combo nước tương")
            continue
        bom = combos.setdefault(cid, {'components': [], 'choice': None, 'gift': []})
        item = {'sp': r.get('sp') or None, 'ten': norm(r.get('ten')), 'qty': int(num(r.get('qty'))) or 1}
        if r.get('loai') == 'thanh_phan':
            bom['components'].append(item)
        elif r.get('loai') == 'chon':
            if bom['choice'] is None:
                bom['choice'] = []
            bom['choice'].append({'sp': item['sp'], 'ten': item['ten']})
        elif r.get('loai') == 'tang':
            bom['gift'].append(item)
    weights = {}
    for r in _read_sheet(wb, 'Tải trọng', WEIGHT_COLS):
        if r.get('sp'):
            weights[r['sp']] = {'quy_cach_tt': num(r.get('quy_cach_tt')), 'kg_thung': num(r.get('kg_thung'))}
    provinces = [norm(r['ten']) for r in _read_sheet(wb, 'Tỉnh', [('ten', 'Tỉnh/Thành phố')]) if norm(r.get('ten'))]
    statuses = [norm(r['ten']) for r in _read_sheet(wb, 'Trạng thái', [('ten', 'Trạng thái')]) if norm(r.get('ten'))]
    orders, by_id = [], {}
    for o in _read_sheet(wb, 'Đơn hàng', ORDER_COLS):
        oo = {'id': o.get('id'), 'ma': norm(o.get('ma')), 'kenh': o.get('kenh'),
              'ngay_dat': iso(to_date(o.get('ngay_dat'))), 'khach': norm(o.get('khach')), 'tinh': norm(o.get('tinh')),
              'trang_thai': o.get('trang_thai') or 'da_len_don', 'pay_status': o.get('pay_status') or 'chua_tt',
              'ck_pct': num(o.get('ck_pct')), 'cong_no': num(o.get('cong_no')), 'phi_vc': num(o.get('phi_vc')),
              'chi_phi_khac': num(o.get('chi_phi_khac')), 'ghi_chu': norm(o.get('ghi_chu')),
              'tao_luc': o.get('tao_luc'), 'lines': [], 'combos': [], 'httm': []}
        orders.append(oo)
        by_id[oo['id']] = oo
    for ln in _read_sheet(wb, 'Chi tiết đơn', LINE_COLS):
        o = by_id.get(ln.get('order_id'))
        if not o:
            continue
        loai = ln.get('loai')
        if loai == 'ban':
            o['lines'].append({'sp': ln.get('sp'), 'thung': num(ln.get('thung')), 'gia_chai': opt_num(ln.get('gia_chai'))})
        elif loai == 'combo':
            o['combos'].append({'combo_id': ln.get('combo_id'), 'choice_sp': ln.get('choice_sp') or None,
                                'phan': num(ln.get('phan'))})
        elif loai == 'httm':
            o['httm'].append({'sp': ln.get('sp'), 'thung': num(ln.get('thung'))})
    deliveries = []
    for d in _read_sheet(wb, 'Đợt giao', DELIVERY_COLS):
        if d.get('order_id'):
            deliveries.append({'order_id': d['order_id'], 'dot': int(num(d.get('dot'))) or 1,
                               'ngay': iso(to_date(d.get('ngay'))), 'sp': d.get('sp'), 'qty': num(d.get('qty'))})
    wb.close()
    DATA_MTIME[0] = DATA_FILE.stat().st_mtime
    return {'settings': settings, 'products': products, 'combos': combos, 'weights': weights,
            'provinces': provinces, 'statuses': statuses, 'order_rows': {}, 'orders': orders, 'deliveries': deliveries}


# =====================================================================
# tính tiền đơn hàng
# =====================================================================

def _prod_map(db):
    return {p['id']: p for p in db['products']}


def quote(db, body):
    """Tính tiền 1 đơn. Mặt hàng lẻ: tiền = thùng × giá chai × quy cách. Combo: tiền = phần × giá bán combo.
    Chiết khấu % trừ trên tổng tiền hàng; công nợ là số tiền chưa thu; HTTM không tính tiền."""
    prods = _prod_map(db)
    lines, combos, httm = [], [], []
    for ln in body.get('lines', []):
        p, q = prods.get(ln.get('sp')), num(ln.get('thung'))
        if not p or q <= 0:
            continue
        gia_tay = ln.get('gia_chai') not in (None, '') and num(ln['gia_chai']) > 0
        gia_chai = num(ln['gia_chai']) if gia_tay else p['gia_ban']
        gia_thung = gia_chai * p['quy_cach']
        lines.append({'sp': p['id'], 'ten': p['ten'], 'dvt': p.get('dvt', ''), 'quy_cach': p['quy_cach'],
                      'gia_chai': gia_chai, 'gia_ban_niemyet': p['gia_ban'], 'gia_tay': gia_tay,
                      'gia_thung': gia_thung, 'thung': q, 'thanh_tien': q * gia_thung, 'tong_giao': q * p['quy_cach']})
    for cb in body.get('combos', []):
        p, phan = prods.get(cb.get('combo_id')), num(cb.get('phan'))
        if not p or phan <= 0:
            continue
        combos.append({'combo_id': p['id'], 'ten': p['ten'], 'choice_sp': cb.get('choice_sp'), 'phan': phan,
                       'gia_ban': p['gia_ban'], 'thanh_tien': phan * p['gia_ban']})
    for h in body.get('httm', []):
        p, q = prods.get(h.get('sp')), num(h.get('thung'))
        if not p or q <= 0:
            continue
        httm.append({'sp': p['id'], 'ten': p['ten'], 'dvt': p.get('dvt', ''), 'quy_cach': p['quy_cach'],
                     'thung': q, 'tong_giao': q * p['quy_cach']})
    tong = sum(x['thanh_tien'] for x in lines) + sum(x['thanh_tien'] for x in combos)
    ck_pct = num(body.get('ck_pct'))
    chiet_khau = money(tong * ck_pct)
    sau_ck = money(tong) - chiet_khau
    cong_no = money(num(body.get('cong_no')))
    return {'lines': lines, 'combos': combos, 'httm': httm, 'tong_tien_hang': money(tong), 'ck_pct': ck_pct,
            'chiet_khau': chiet_khau, 'sau_ck': sau_ck, 'cong_no': cong_no, 'tong_thanh_toan': sau_ck - cong_no}


def order_profit(db, order):
    """Lợi nhuận 1 đơn. Doanh thu = giá bán niêm yết × SL (combo = 1 dòng theo giá combo).
    Lợi nhuận = Doanh thu − Giá vốn − Chiết khấu(DT×CK%) − HH người giới thiệu − HH tỉnh − Phí VC − Chi phí khác.
    Luôn trừ cả 2 khoản hoa hồng (tỷ lệ ở Cài đặt). Doanh thu KHÔNG đổi theo giá sửa tay hay chiết khấu."""
    prods = _prod_map(db)
    s = db['settings']
    doanh_thu = gia_von = san_luong = 0.0
    for ln in order.get('lines', []):
        p, q = prods.get(ln.get('sp')), num(ln.get('thung'))
        if not p or q <= 0:
            continue
        doanh_thu += q * p['gia_ban'] * p['quy_cach']
        gia_von += q * p['gia_von'] * p['quy_cach']
        san_luong += q * p['quy_cach']
    for cb in order.get('combos', []):
        p, phan = prods.get(cb.get('combo_id')), num(cb.get('phan'))
        if not p or phan <= 0:
            continue
        doanh_thu += phan * p['gia_ban']
        gia_von += phan * p['gia_von']
        san_luong += phan
    chiet_khau = doanh_thu * num(order.get('ck_pct'))
    tien_cp = doanh_thu * num(s.get('cp_pct'))
    tien_tinh = doanh_thu * num(s.get('tinh_pct'))
    phi_vc, chi_phi_khac = num(order.get('phi_vc')), num(order.get('chi_phi_khac'))
    loi_nhuan = doanh_thu - gia_von - chiet_khau - tien_cp - tien_tinh - phi_vc - chi_phi_khac
    return {'doanh_thu': money(doanh_thu), 'gia_von': money(gia_von), 'chiet_khau': money(chiet_khau),
            'tien_cp': money(tien_cp), 'tien_tinh': money(tien_tinh), 'phi_vc': money(phi_vc),
            'chi_phi_khac': money(chi_phi_khac), 'loi_nhuan': money(loi_nhuan), 'san_luong': int(round(san_luong)),
            'bien_ln': (loi_nhuan / doanh_thu) if doanh_thu else 0.0}


def _active_orders(db, kenh=None, month=None):
    out = []
    for o in db['orders']:
        if o.get('trang_thai') == 'huy':
            continue
        if kenh and o.get('kenh') != kenh:
            continue
        if month and (o.get('ngay_dat') or '')[:7] != month:
            continue
        out.append(o)
    return out


_PROFIT_SUM = ('doanh_thu', 'gia_von', 'chiet_khau', 'tien_cp', 'tien_tinh', 'phi_vc', 'chi_phi_khac',
               'loi_nhuan', 'san_luong')


def _agg(orders, db, key_fn):
    groups = {}
    for o in orders:
        g = groups.setdefault(key_fn(o), {'so_don': 0, **{k: 0 for k in _PROFIT_SUM}})
        pr = order_profit(db, o)
        g['so_don'] += 1
        for k in _PROFIT_SUM:
            g[k] += pr[k]
    for g in groups.values():
        g['bien_ln'] = (g['loi_nhuan'] / g['doanh_thu']) if g['doanh_thu'] else 0.0
    return groups


def report_by_province(db, kenh=None, month=None):
    orders = _active_orders(db, kenh, month)
    groups = _agg(orders, db, lambda o: o.get('tinh') or '(chưa có tỉnh)')
    return [{'tinh': t, **g} for t, g in sorted(groups.items())]


def report_by_month(db, kenh=None):
    orders = _active_orders(db, kenh)
    groups = _agg(orders, db, lambda o: (o.get('ngay_dat') or '')[:7])
    return [{'thang': m, **g} for m, g in sorted(groups.items()) if m]


def explode_order(db, order):
    """Tách đơn thành danh sách SP lẻ cho Phiếu Xuất Kho tổng.
    Combo -> cộng dồn thành phần (gồm món chọn 1-trong-N + hàng tặng). Quy thùng/lẻ theo quy cách đơn hàng;
    tải trọng theo quy cách của bảng 'Đối chiếu tải trọng' (KHÁC quy cách đơn hàng)."""
    prods = _prod_map(db)
    combos = db['combos']
    chai = {}      # sp_id -> tổng chai/hũ
    order_sp = []  # giữ thứ tự xuất hiện

    def add(sp, n):
        if not sp or n <= 0:
            return
        if sp not in chai:
            order_sp.append(sp)
        chai[sp] = chai.get(sp, 0) + n

    for ln in order.get('lines', []):
        p = prods.get(ln.get('sp'))
        if p:
            add(p['id'], num(ln.get('thung')) * p['quy_cach'])
    for cb in order.get('combos', []):
        bom = combos.get(cb.get('combo_id'))
        phan = num(cb.get('phan'))
        if not bom or phan <= 0:
            continue
        for c in bom.get('components', []):
            add(c.get('sp'), phan * (c.get('qty') or 1))
        if bom.get('choice') and cb.get('choice_sp'):
            add(cb['choice_sp'], phan)
        for g in bom.get('gift', []):
            add(g.get('sp'), phan * (g.get('qty') or 1))
    for h in order.get('httm', []):
        p = prods.get(h.get('sp'))
        if p:
            add(p['id'], num(h.get('thung')) * p['quy_cach'])

    out = []
    for sp in order_sp:
        p = prods.get(sp, {})
        qc = int(p.get('quy_cach') or 1)
        tong = int(round(chai[sp]))
        thung, le = divmod(tong, qc) if qc else (tong, 0)
        w = db['weights'].get(sp)
        tai_trong = (tong / w['quy_cach_tt'] * w['kg_thung']) if (w and w.get('quy_cach_tt')) else 0
        out.append({'sp': sp, 'ten': p.get('ten', sp), 'dvt': p.get('dvt', ''), 'quy_cach': qc,
                    'tong_chai': tong, 'thung': thung, 'le': le, 'tai_trong': round(tai_trong, 3)})
    return out


# =====================================================================
# chia đợt giao hàng
# =====================================================================

def delivery_state(db, order):
    """Trạng thái giao: đã giao & còn nợ hàng theo từng SP lẻ (tách combo), các đợt, và trạng thái đơn."""
    ordered = {r['sp']: r['tong_chai'] for r in explode_order(db, order)}
    delivered = {}
    dots = {}
    for d in db.get('deliveries', []):
        if d.get('order_id') != order['id']:
            continue
        delivered[d['sp']] = delivered.get(d['sp'], 0) + num(d['qty'])
        dots.setdefault(d['dot'], {'dot': d['dot'], 'ngay': d.get('ngay'), 'items': []})
        dots[d['dot']]['items'].append({'sp': d['sp'], 'qty': num(d['qty'])})
    per_sp = {sp: {'ordered': ordered[sp], 'delivered': int(round(delivered.get(sp, 0))),
                   'remaining': int(round(ordered[sp] - delivered.get(sp, 0)))} for sp in ordered}
    any_delivered = any(v['delivered'] > 0 for v in per_sp.values())
    all_done = all(v['remaining'] <= 0 for v in per_sp.values()) and bool(per_sp)
    trang_thai = 'da_giao_du' if (any_delivered and all_done) else 'giao_mot_phan' if any_delivered else 'da_len_don'
    return {'per_sp': per_sp, 'dot': [dots[k] for k in sorted(dots)], 'trang_thai': trang_thai}


def add_delivery(db, order_id, payload):
    """Thêm 1 đợt giao (tối đa 3). Kẹp số lượng mỗi SP không vượt quá số còn nợ."""
    order = next((o for o in db['orders'] if o['id'] == order_id), None)
    if order is None:
        raise ValueError('Không tìm thấy đơn.')
    existing = sorted({d['dot'] for d in db['deliveries'] if d.get('order_id') == order_id})
    if len(existing) >= 3:
        raise ValueError('Đơn đã giao đủ 3 đợt, không thêm được đợt nữa.')
    dot = (max(existing) + 1) if existing else 1
    st = delivery_state(db, order)
    ngay = iso(to_date(payload.get('ngay')) or today())
    added = []
    for it in payload.get('items', []):
        sp = it.get('sp')
        remaining = st['per_sp'].get(sp, {}).get('remaining', 0)
        qty = min(num(it.get('qty')), remaining)
        if qty > 0:
            rec = {'order_id': order_id, 'dot': dot, 'ngay': ngay, 'sp': sp, 'qty': int(round(qty))}
            db['deliveries'].append(rec)
            added.append(rec)
    if not added:
        raise ValueError('Không có hàng để giao trong đợt này (đã giao đủ hoặc số lượng = 0).')
    order['trang_thai'] = delivery_state(db, order)['trang_thai']
    return {'dot': dot, 'items': added}


# =====================================================================
# xuất Đơn hàng + Phiếu Xuất Kho (điền mẫu Excel giữ logo/định dạng)
# =====================================================================
DDH = 'Đơn Hàng Viettel'
PXK = 'Phiếu Xuất Kho'
PXKT = 'PXK tổng SL đơn hàng'


def run_excel_job(job):
    OUT_DIR.mkdir(exist_ok=True)
    job_file = OUT_DIR / '_job.json'
    job_file.write_text(json.dumps(job, ensure_ascii=False), encoding='utf-8')
    try:
        res = subprocess.run(['powershell', '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass',
                              '-File', str(ROOT / 'dien_mau_excel.ps1'), str(job_file)], capture_output=True,
                             timeout=EXCEL_TIMEOUT)
    except subprocess.TimeoutExpired:
        raise RuntimeError(f'Excel không phản hồi sau {EXCEL_TIMEOUT // 60} phút (có thể Excel đang hiện hộp thoại). '
                           'Mở Excel kiểm tra, đóng hộp thoại rồi làm lại.') from None
    finally:
        job_file.unlink(missing_ok=True)
    if res.returncode != 0:
        err = res.stderr.decode('utf-8', 'replace')[-400:] or res.stdout.decode('utf-8', 'replace')[-400:]
        raise RuntimeError('Excel không điền được mẫu: ' + err)


def _ensure_template():
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    out = TEMPLATE_DIR / 'Mau-don-viettel.xlsx'
    if not out.exists():
        _, tpl = find_source_files()
        shutil.copy2(tpl, out)
    return out


def export_order(db, order):
    """Điền mẫu 'Viettel - Đơn hàng mẫu.xlsx' (3 sheet) bằng Excel, xuất PDF.
    Đơn Hàng: điền SL/giá/KM/tặng từng dòng (Phiếu Xuất Kho tự liên kết theo công thức).
    PXK tổng SL: điền SP lẻ sau khi tách combo + tải trọng + đã giao/còn nợ theo đợt."""
    template = _ensure_template()
    prods = _prod_map(db)
    q = quote(db, order)
    folder = OUT_DIR / 'Đơn hàng'
    folder.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r'[\\/:*?"<>|]', '', order.get('khach') or 'KhachLe')
    out = folder / f'{order["ma"]} - {safe}.xlsx'
    d = to_date(order.get('ngay_dat')) or today()
    khach, tinh = order.get('khach', ''), order.get('tinh', '')

    cells = [[DDH, 'E8', khach], [DDH, 'E9', tinh], [DDH, 'J6', f'Ngày đặt: {d:%d/%m/%Y}'],
             [PXK, 'D8', khach], [PXK, 'D9', tinh],
             [PXKT, 'C5', khach], [PXKT, 'C6', tinh], [PXKT, 'G8', f'Ngày: {d:%d/%m/%Y}']]

    rowagg = {}
    for ln in order.get('lines', []):
        p = prods.get(ln.get('sp'))
        r = p.get('row_mau') if p else None
        if not r:
            continue
        gia = num(ln['gia_chai']) if (ln.get('gia_chai') not in (None, '') and num(ln['gia_chai']) > 0) else p['gia_ban']
        a = rowagg.setdefault(r, {'sl': 0, 'gia': gia, 'qc': p['quy_cach'], 'ctkm': p.get('ctkm', ''), 'tang': 0})
        a['sl'] += num(ln.get('thung'))
    for cb in order.get('combos', []):
        p = prods.get(cb.get('combo_id'))
        r = p.get('row_mau') if p else None
        if not r:
            continue
        a = rowagg.setdefault(r, {'sl': 0, 'gia': p['gia_ban'], 'qc': 1,  # combo: đơn giá thùng = đơn giá phần
                                  'ctkm': p.get('ctkm', ''), 'tang': 0})
        a['sl'] += num(cb.get('phan'))
    for r in range(12, 43):  # xóa số lượng mẫu còn sót để không cộng nhầm
        if r not in rowagg:
            cells += [[DDH, f'I{r}', 0], [DDH, f'K{r}', None], [DDH, f'L{r}', None]]
    for r, a in rowagg.items():
        cells += [[DDH, f'G{r}', a['gia']], [DDH, f'H{r}', a['gia'] * a['qc']], [DDH, f'I{r}', a['sl']],
                  [DDH, f'K{r}', a['ctkm']], [DDH, f'L{r}', a['tang']]]

    inserts = [[DDH, 44, 1]]  # chèn dòng "Chiết khấu" sau 'Tổng tiền hàng' (dòng 43)
    cells += [[DDH, 'A44', 'Chiết khấu'], [DDH, 'I44', q['ck_pct']], [DDH, 'J44', '=J43*I44'],
              [DDH, 'J45', q['cong_no']], [DDH, 'I46', '=J43-J44-J45']]

    skus = explode_order(db, order)
    dv = delivery_state(db, order)
    n = len(skus)
    if n > 10:
        inserts.append([PXKT, 21, n - 10])
    for i, s in enumerate(skus):
        r = 11 + i
        delivered = dv['per_sp'].get(s['sp'], {}).get('delivered', 0)
        remaining = dv['per_sp'].get(s['sp'], {}).get('remaining', s['tong_chai'])
        gt, gl = divmod(int(round(delivered)), s['quy_cach']) if s['quy_cach'] else (int(round(delivered)), 0)
        cells += [[PXKT, f'A{r}', i + 1], [PXKT, f'B{r}', s['ten']], [PXKT, f'C{r}', s['dvt']],
                  [PXKT, f'D{r}', s['quy_cach']], [PXKT, f'E{r}', s['thung']], [PXKT, f'F{r}', s['le']],
                  [PXKT, f'G{r}', gt], [PXKT, f'H{r}', gl], [PXKT, f'I{r}', remaining],
                  [PXKT, f'K{r}', round(s['tai_trong'], 2)]]
    for r in range(11 + n, 21):
        cells += [[PXKT, f'{c}{r}', None] for c in 'ABCDEFGHIJK']
    if n > 10:
        tr = 21 + (n - 10)
        for col in 'EFGHK':
            cells.append([PXKT, f'{col}{tr}', f'=SUM({col}11:{col}{tr - 1})'])

    job = {'template': str(template), 'output': str(out), 'pdf': str(out.with_suffix('.pdf')), 'pdf_sheet': DDH,
           'inserts': inserts, 'cells': cells, 'hide_zero': [[DDH, 12, 42, 'I', 'M']]}
    run_excel_job(job)
    wb = load_workbook(out, data_only=True)
    check = num(wb[DDH]['J43'].value)
    wb.close()
    lech = {} if abs(check - q['tong_tien_hang']) <= 1 else {'tong_tien_hang': check}
    return {'file': str(out), 'pdf': str(out.with_suffix('.pdf')), 'lech': lech}


def _report_head(ws, cols):
    for i, h in enumerate(cols, 1):
        c = ws.cell(1, i, h)
        c.font, c.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F4E78')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(30, len(h) + 3))
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'


REPORT_TEMPLATES = {'buu_cuc': ROOT.parent / 'BÁO CÁO VIETTEL BƯU CỤC.xlsx',
                    'tm': ROOT.parent / 'BÁO CÁO VIETTEL TM.xlsx'}


def _norm_hdr(v):
    return re.sub(r'\s+', ' ', str(v or '')).strip().lower()


def _header_map(ws, needed, max_scan=8):
    """Tìm dòng tiêu đề + vị trí cột theo TÊN (khớp chứa chuỗi). needed: {key: [chuỗi con ưu tiên]}.
    Trả (header_row, {key: col}). Thiếu bất kỳ key -> RuntimeError."""
    for r in range(1, max_scan + 1):
        cells = {c: _norm_hdr(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)}
        pos = {}
        for key, subs in needed.items():
            for sub in subs:
                col = next((c for c, t in cells.items() if t and sub in t and c not in pos.values()), None)
                if col:
                    pos[key] = col
                    break
        if all(k in pos for k in needed):
            return r, pos
    missing = [k for k in needed if True]
    raise RuntimeError(f'File mẫu "{ws.title}" thiếu cột tiêu đề cần thiết ({", ".join(missing)}).')


def _report_detail_rows(db, kenh):
    """Mỗi dòng SP/combo của mọi đơn (khớp order_profit: combo = 1 dòng theo giá combo)."""
    prods = _prod_map(db)
    s = db['settings']
    cp_pct, tinh_pct = num(s.get('cp_pct')), num(s.get('tinh_pct'))
    rows = []
    for o in _active_orders(db, kenh):
        first = True
        items = []
        for ln in o.get('lines', []):
            p = prods.get(ln.get('sp'))
            q = num(ln.get('thung'))
            if p and q > 0:
                sl = q * p['quy_cach']
                items.append((p, sl, p['gia_ban'], p['gia_von']))
        for cb in o.get('combos', []):
            p, phan = prods.get(cb.get('combo_id')), num(cb.get('phan'))
            if p and phan > 0:
                items.append((p, phan, p['gia_ban'], p['gia_von']))
        for p, sl, don_gia, gv in items:
            doanh_thu = sl * don_gia
            rows.append({'ngay': to_date(o.get('ngay_dat')), 'ma': o.get('ma'), 'tinh': o.get('tinh'),
                         'kho': o.get('khach'), 'pay': PAY_STATUS.get(o.get('pay_status'), ''),
                         'ma_sp': p.get('ma_vach') or p['id'], 'ten': p['ten'], 'sl': sl, 'don_gia': don_gia,
                         'gia_von_sp': gv, 'tong_gv': sl * gv, 'doanh_thu': doanh_thu, 'phi_vc': num(o.get('phi_vc')) if first else 0,
                         'pct_cp': cp_pct, 'tien_cp': doanh_thu * cp_pct, 'pct_tinh': tinh_pct, 'tien_tinh': doanh_thu * tinh_pct})
            first = False
    return rows


def _write_rows(ws, header_row, pos, data, numkeys):
    """Ghi các dòng `data` (list dict) bắt đầu ngay dưới header; sao style từ dòng dữ liệu mẫu; xóa công thức/dòng thừa."""
    from copy import copy
    maxcol = ws.max_column
    style_row = header_row + 1
    has_sample = ws.max_row > header_row
    styles = {c: copy(ws.cell(style_row, c)._style) for c in pos.values()} if has_sample else {}
    fmts = {c: ws.cell(style_row, c).number_format for c in pos.values()} if has_sample else {}
    r = header_row + 1
    for item in data:
        for col in range(1, maxcol + 1):  # mọi cột ngoài dữ liệu: xóa công thức mẫu (giữ style)
            if col not in pos.values():
                ws.cell(r, col).value = None
        for key, col in pos.items():
            cell = ws.cell(r, col)
            cell.value = item.get(key)  # gán .value để xóa được công thức cũ khi giá trị là None
            if col in styles:
                cell._style = copy(styles[col])
                cell.number_format = fmts[col]
        r += 1
    if ws.max_row >= r:  # xóa dòng mẫu/công thức cũ còn thừa bên dưới
        ws.delete_rows(r, ws.max_row - r + 1)


def fill_report_template(db, kenh):
    """Điền dữ liệu đã tính vào bản sao file mẫu (giữ bố cục/định dạng, ghi THẲNG giá trị — khớp order_profit)."""
    tpl = REPORT_TEMPLATES.get(kenh)
    if not tpl or not Path(tpl).exists():
        raise RuntimeError(f'Không tìm thấy file mẫu báo cáo: {tpl}')
    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / f'Báo cáo Viettel {CHANNELS.get(kenh, "tất cả kênh")} (cập nhật {datetime.now():%d.%m.%Y %Hh%M}).xlsx'
    shutil.copyfile(tpl, out)
    wb = load_workbook(out)
    s = db['settings']
    hh_name = s.get('ten_hh_cp') or TEXT_SETTINGS['ten_hh_cp']

    # 1) Chi tiết đơn ("Đơn Hàng Từng Tỉnh")
    det = _sheet(wb, 'Đơn Hàng Từng Tỉnh')
    need = {'ngay': ['ngày'], 'ma': ['mã đơn'], 'tinh': ['tỉnh'], 'kho': ['kho'], 'ma_sp': ['mã sp'],
            'ten': ['tên sản phẩm'], 'sl': ['số lượng'], 'don_gia': ['đơn giá'], 'gia_von_sp': ['giá vốn (sp)', 'giá vốn'],
            'tong_gv': ['tổng giá vốn'], 'doanh_thu': ['doanh thu'], 'phi_vc': ['phí vận chuyển'],
            'pct_cp': ['% ' + hh_name.lower(), '% chị', '%'], 'tien_cp': ['tiền ' + hh_name.lower(), 'tiền chị']}
    hr, pos = _header_map(det, need)
    if _norm_hdr_has(det, hr, 'trạng thái'):
        pos['pay'] = _col_with(det, hr, 'trạng thái')
    # cột hoa hồng tỉnh nếu mẫu có
    if _norm_hdr_has(det, hr, 'của tỉnh') or _norm_hdr_has(det, hr, 'của bưu cục'):
        c = _col_with(det, hr, 'tiền của tỉnh')
        if c:
            pos['tien_tinh'] = c
        c2 = _col_with(det, hr, '% của')
        if c2:
            pos['pct_tinh'] = c2
    _write_rows(det, hr, pos, _report_detail_rows(db, kenh), ())

    # 2) Tổng Hợp Doanh Số (theo tỉnh) — sinh lại từ report_by_province
    th = _sheet(wb, 'Tổng Hợp Doanh Số')
    needp = {'tinh': ['tỉnh'], 'doanh_thu': ['doanh số'], 'san_luong': ['sản lượng'], 'gia_von': ['giá vốn'],
             'phi_vc': ['phí vận chuyển'], 'pct_cp': ['% chị', '%'], 'tien_cp': ['tiền chị', 'tiền ' + hh_name.lower()],
             'pct_tinh': ['% của'], 'tien_tinh': ['tiền của tỉnh'], 'tong_cp': ['tổng chi phí'],
             'loi_nhuan': ['lợi nhuận'], 'bien': ['biên']}
    hrp, posp = _header_map(th, needp)
    if _norm_hdr_has(th, hrp, 'stt'):
        posp['stt'] = _col_with(th, hrp, 'stt')
    prov = report_by_province(db, kenh)
    cp_pct, tinh_pct = num(s.get('cp_pct')), num(s.get('tinh_pct'))
    prows, tot = [], {'doanh_thu': 0, 'san_luong': 0, 'gia_von': 0, 'phi_vc': 0, 'tien_cp': 0, 'tien_tinh': 0,
                      'tong_cp': 0, 'loi_nhuan': 0}
    for i, g in enumerate(prov, 1):
        tong_cp = g['gia_von'] + g['chiet_khau'] + g['tien_cp'] + g['tien_tinh'] + g['phi_vc'] + g['chi_phi_khac']
        row = {'stt': i, 'tinh': g['tinh'], 'doanh_thu': g['doanh_thu'], 'san_luong': g['san_luong'],
               'gia_von': g['gia_von'], 'phi_vc': g['phi_vc'], 'pct_cp': cp_pct, 'tien_cp': g['tien_cp'],
               'pct_tinh': tinh_pct, 'tien_tinh': g['tien_tinh'], 'tong_cp': tong_cp, 'loi_nhuan': g['loi_nhuan'],
               'bien': g['bien_ln']}
        prows.append(row)
        for k in tot:
            tot[k] += row.get(k, 0)
    prows.append({'tinh': 'TỔNG CỘNG', 'bien': (tot['loi_nhuan'] / tot['doanh_thu']) if tot['doanh_thu'] else 0, **tot})
    _write_rows(th, hrp, posp, prows, ())

    # 3) Theo Dõi Doanh Số Tháng — sinh lại từ report_by_month
    mo = _sheet(wb, 'Theo Dõi Doanh Số Tháng')
    needm = {'thang': ['tháng'], 'doanh_thu': ['doanh số'], 'loi_nhuan': ['lợi nhuận'], 'bien': ['biên']}
    hrm, posm = _header_map(mo, needm)
    mrows, dt, ln = [], 0, 0
    for g in report_by_month(db, kenh):
        lbl = g['thang']
        mrows.append({'thang': f'Tháng {lbl[5:7]}/{lbl[:4]}' if lbl else '', 'doanh_thu': g['doanh_thu'],
                      'loi_nhuan': g['loi_nhuan'], 'bien': g['bien_ln']})
        dt += g['doanh_thu']
        ln += g['loi_nhuan']
    mrows.append({'thang': 'TỔNG CỘNG', 'doanh_thu': dt, 'loi_nhuan': ln, 'bien': (ln / dt) if dt else 0})
    _write_rows(mo, hrm, posm, mrows, ())

    # 4) Danh Mục Thiết Lập — ghi lại bảng SP (tham khảo)
    _fill_catalog_sheet(wb, db)

    wb.save(out)
    wb.close()
    return str(out)


def _norm_hdr_has(ws, row, sub):
    return any(sub in _norm_hdr(ws.cell(row, c).value) for c in range(1, ws.max_column + 1))


def _col_with(ws, row, sub):
    return next((c for c in range(1, ws.max_column + 1) if sub in _norm_hdr(ws.cell(row, c).value)), None)


def _fill_catalog_sheet(wb, db):
    try:
        ws = _sheet(wb, 'Danh Mục Thiết Lập')
    except RuntimeError:
        return
    need = {'stt': ['stt'], 'ma': ['mã sản phẩm'], 'ten': ['tên sản phẩm'], 'qc': ['quy cách'],
            'gv': ['giá vốn'], 'gb': ['giá bán'], 'ctkm': ['khuyến mã', 'khuyến mại']}
    try:
        hr, pos = _header_map(ws, need)
    except RuntimeError:
        return
    from copy import copy
    style_row = hr + 1
    styles = {c: copy(ws.cell(style_row, c)._style) for c in pos.values()}
    r = hr + 1
    for i, p in enumerate([p for p in db['products'] if p['id'] != VOUCHER_ID], 1):
        for key, val in (('stt', i), ('ma', p.get('ma_vach')), ('ten', p['ten']), ('qc', p.get('quy_cach')),
                         ('gv', p.get('gia_von')), ('gb', p.get('gia_ban')), ('ctkm', p.get('ctkm'))):
            cell = ws.cell(r, pos[key], val)
            cell._style = copy(styles[pos[key]])
        r += 1


# =====================================================================
# API
# =====================================================================

def find_order(db, oid):
    o = next((o for o in db['orders'] if o['id'] == oid), None)
    if not o:
        raise ValueError('Không tìm thấy đơn.')
    return o


def prov_abbrev(tinh):
    return ''.join(w[0] for w in key(tinh).split() if w).upper() or 'X'


def next_code(db, kenh, tinh, d):
    tag = 'TM' if kenh == 'tm' else prov_abbrev(tinh)
    base = f'DHVT-{tag}-{d:%d.%m.%Y}'
    same = [o for o in db['orders'] if str(o.get('ma', '')).startswith(base)]
    return base if not same else f'{base}-{len(same) + 1}'


def order_view(db, o):
    q = quote(db, o)
    pr = order_profit(db, o)
    dv = delivery_state(db, o)
    return {**o, 'tong_tien_hang': q['tong_tien_hang'], 'chiet_khau': q['chiet_khau'], 'sau_ck': q['sau_ck'],
            'cong_no_tien': q['cong_no'], 'tong_thanh_toan': q['tong_thanh_toan'], 'profit': pr, 'delivery': dv}


def public_state(db):
    return {'today': iso(today()), 'channels': CHANNELS, 'status': STATUS, 'pay_status': PAY_STATUS,
            'settings': db['settings'], 'products': db['products'], 'combos': db['combos'], 'weights': db['weights'],
            'provinces': db['provinces'], 'statuses': db.get('statuses', []),
            'orders': [order_view(db, o) for o in db['orders']], 'data_file': str(DATA_FILE)}


def _build_order(db, body, existing=None):
    q = quote(db, body)
    if not q['lines'] and not q['combos']:
        raise ValueError('Đơn chưa có mặt hàng hay combo nào.')
    d = to_date(body.get('ngay_dat')) or today()
    cur = existing if existing is not None else {}
    upd = {'kenh': body.get('kenh') if body.get('kenh') in CHANNELS else 'buu_cuc',
              'ngay_dat': iso(d), 'khach': norm(body.get('khach')), 'tinh': norm(body.get('tinh')),
              'pay_status': body.get('pay_status') if body.get('pay_status') in PAY_STATUS else cur.get('pay_status', 'chua_tt'),
              'ck_pct': num(body.get('ck_pct')), 'cong_no': num(body.get('cong_no')), 'phi_vc': num(body.get('phi_vc')),
              'chi_phi_khac': num(body.get('chi_phi_khac')), 'ghi_chu': norm(body.get('ghi_chu')),
              'lines': [{'sp': ln['sp'], 'thung': num(ln.get('thung')), 'gia_chai': opt_num(ln.get('gia_chai'))}
                        for ln in body.get('lines', []) if ln.get('sp') and num(ln.get('thung')) > 0],
              'combos': [{'combo_id': cb['combo_id'], 'choice_sp': cb.get('choice_sp') or None, 'phan': num(cb.get('phan'))}
                         for cb in body.get('combos', []) if cb.get('combo_id') and num(cb.get('phan')) > 0],
              'httm': [{'sp': h['sp'], 'thung': num(h.get('thung'))}
                       for h in body.get('httm', []) if h.get('sp') and num(h.get('thung')) > 0]}
    prods = _prod_map(db)
    for cb in upd['combos']:  # kiểm tra trên bản mới; lỗi thì đơn đang có không bị đổi
        bom = db['combos'].get(cb['combo_id'])
        if bom and bom.get('choice') and not cb.get('choice_sp'):
            ten = (prods.get(cb['combo_id']) or {}).get('ten', '')[:40]
            raise ValueError(f'Combo "{ten}…" phải chọn 1 món (chọn 1 trong 3) trước khi lưu.')
    cur.update(upd)
    return cur


def valid_channel(kenh):
    """Kênh gửi lên phải là 1 kênh có thật (hoặc trống = mọi kênh)."""
    if kenh in (None, ''):
        return None
    if kenh not in CHANNELS:
        raise ValueError('Kênh không hợp lệ.')
    return kenh


def handle(db, path, body):
    if path == '/api/quote':
        return quote(db, body)
    if path == '/api/order/save':
        if body.get('id'):
            o = find_order(db, body['id'])
            if o.get('trang_thai') not in ('da_len_don', None):
                raise ValueError('Chỉ sửa được đơn chưa giao (Đã lên đơn).')
            _build_order(db, body, o)
        else:
            d = to_date(body.get('ngay_dat')) or today()
            ma = next_code(db, body.get('kenh'), norm(body.get('tinh')), d)
            o = _build_order(db, body)
            o.update(id=f'O{datetime.now():%Y%m%d%H%M%S%f}'[:18], ma=ma, trang_thai='da_len_don',
                     tao_luc=datetime.now().isoformat(timespec='seconds'))
            db['orders'].append(o)
        save_db(db)
        return {'ok': True, 'order': order_view(db, o)}
    if path == '/api/order/deliver':
        o = find_order(db, body['id'])
        res = add_delivery(db, o['id'], body)
        save_db(db)
        return {'ok': True, 'dot': res['dot'], 'order': order_view(db, o)}
    if path == '/api/order/cancel':
        o = find_order(db, body['id'])
        o['trang_thai'] = 'huy'
        save_db(db)
        return {'ok': True}
    if path == '/api/order/set-status':
        o = find_order(db, body['id'])
        if body.get('pay_status') in PAY_STATUS:
            o['pay_status'] = body['pay_status']
        if body.get('trang_thai') in STATUS:
            if (body['trang_thai'] == 'da_len_don' and o.get('trang_thai') != 'da_len_don'
                    and any(d.get('order_id') == o['id'] for d in db.get('deliveries', []))):
                raise ValueError('Đơn đã có đợt giao hàng — không đưa về "Đã lên đơn" được '
                                 '(sửa số lượng lúc này sẽ lệch với số đã giao).')
            o['trang_thai'] = body['trang_thai']
        save_db(db)
        return {'ok': True, 'order': order_view(db, o)}
    if path == '/api/order/export':
        return export_order(db, find_order(db, body['id']))
    if path == '/api/report':
        kenh = valid_channel(body.get('kenh'))
        return {'province': report_by_province(db, kenh, body.get('month')), 'month': report_by_month(db, kenh)}
    if path == '/api/report/export':
        if body.get('tong'):  # báo cáo tổng: xuất cả 2 kênh theo mẫu
            return {'files': [fill_report_template(db, k) for k in CHANNELS]}
        return {'file': fill_report_template(db, valid_channel(body.get('kenh')))}
    if path == '/api/catalog/save':
        _save_catalog(db, body)
        save_db(db)
        return {'ok': True}
    if path == '/api/open':
        target = Path(body['path']).resolve()
        if target.exists() and (OUT_DIR.resolve() in target.parents or target == DATA_FILE.resolve()):
            os.startfile(target)  # noqa: S606
        return {'ok': True}
    raise ValueError('Không rõ yêu cầu: ' + path)


def _new_pid(db):
    """Sinh id sản phẩm dạng 'SP%02d' không trùng id nào hiện có."""
    used = {p['id'] for p in db['products']}
    i = len(used) + 1
    while f'SP{i:02d}' in used:
        i += 1
    return f'SP{i:02d}'


def _catalog_refs(db, pid):
    """Các chỗ đang dùng sản phẩm/combo `pid` (đơn hàng + BOM combo khác). Rỗng = xóa được."""
    refs = []
    for o in db.get('orders', []):
        if (any(l.get('sp') == pid for l in o.get('lines', [])) or any(h.get('sp') == pid for h in o.get('httm', []))
                or any(c.get('combo_id') == pid or c.get('choice_sp') == pid for c in o.get('combos', []))):
            refs.append(f"đơn {o.get('ma') or o.get('id')}")
    for cid, bom in db.get('combos', {}).items():
        if cid == pid:
            continue
        used = ([x.get('sp') for x in bom.get('components', [])] + [x.get('sp') for x in (bom.get('choice') or [])]
                + [x.get('sp') for x in bom.get('gift', [])])
        if pid in used:
            name = next((p['ten'] for p in db['products'] if p['id'] == cid), cid)
            refs.append(f"combo {name}")
    return refs


def _delete_products(db, keep_ids):
    """Xóa product không nằm trong keep_ids; chặn nếu đang được tham chiếu."""
    for p in list(db['products']):
        if p['id'] in keep_ids or p['id'] == VOUCHER_ID:
            continue
        refs = _catalog_refs(db, p['id'])
        if refs:
            raise ValueError(f'Không xóa được "{p["ten"]}" — đang dùng ở: {", ".join(refs)}.')
        db['products'].remove(p)
        db['combos'].pop(p['id'], None)
        db['weights'].pop(p['id'], None)


def _apply_product_fields(p, up):
    if up.get('ten'):
        p['ten'] = norm(up['ten'])
    if 'ma_vach' in up:
        p['ma_vach'] = norm(up['ma_vach'])
    if 'dvt' in up:
        p['dvt'] = norm(up['dvt'])
    if 'quy_cach' in up:
        p['quy_cach'] = int(num(up['quy_cach'])) or p['quy_cach']
    if 'gia_von' in up:
        p['gia_von'] = num(up['gia_von'])
    if 'gia_ban' in up:
        p['gia_ban'] = num(up['gia_ban'])
    if 'ctkm' in up:
        p['ctkm'] = norm(up['ctkm'])
    if 'dang_ban' in up:
        p['dang_ban'] = bool(up['dang_ban'])


def _save_catalog(db, body):
    """Cập nhật danh mục sửa tay từ giao diện. CRUD sản phẩm & combo; chỉ các phần được gửi lên."""
    if 'settings' in body:
        for k in ('cp_pct', 'tinh_pct'):
            if k in body['settings']:
                db['settings'][k] = num(body['settings'][k])
        for k, v in TEXT_SETTINGS.items():
            if k in body['settings']:
                db['settings'][k] = norm(body['settings'][k]) or v
    if 'products' in body:
        cur = {p['id']: p for p in db['products']}
        for up in body['products']:  # thêm id mới, cập nhật id cũ (chỉ các trường gửi lên)
            p = cur.get(up.get('id'))
            if p:
                _apply_product_fields(p, up)
            elif up.get('ten'):  # dòng mới (id rỗng/lạ): tạo sản phẩm thường
                np = {'id': _new_pid(db), 'ma_vach': '', 'ten': '', 'dvt': '', 'quy_cach': 1, 'gia_von': 0.0,
                      'gia_ban': 0.0, 'ctkm': '', 'la_combo': False, 'dang_ban': True, 'row_mau': None}
                _apply_product_fields(np, up)
                db['products'].append(np)
        keep = {up.get('id') for up in body['products'] if cur.get(up.get('id'))}
        keep |= {p['id'] for p in db['products'] if p['id'] not in cur}  # vừa tạo
        keep |= {p['id'] for p in db['products'] if p['la_combo']}  # combo do nhánh combos_meta quản
        _delete_products(db, keep)
    if 'combos_meta' in body:
        cur = {p['id']: p for p in db['products']}
        keep_combo = set()
        for meta in body['combos_meta']:
            p = cur.get(meta.get('id'))
            if p:
                if meta.get('ten'):
                    p['ten'] = norm(meta['ten'])
                if 'dang_ban' in meta:
                    p['dang_ban'] = bool(meta['dang_ban'])
                if 'gia_ban' in meta:
                    p['gia_ban'] = num(meta['gia_ban'])
                if 'gia_von' in meta:
                    p['gia_von'] = num(meta['gia_von'])
                keep_combo.add(p['id'])
            elif meta.get('ten'):  # tạo combo mới = product la_combo, ĐVT combo ('Phần') để load_db giữ là combo
                np = {'id': _new_pid(db), 'ma_vach': norm(meta.get('ma_vach')), 'ten': norm(meta['ten']),
                      'dvt': norm(meta.get('dvt')) or 'Phần', 'quy_cach': 1, 'gia_von': num(meta.get('gia_von')),
                      'gia_ban': num(meta.get('gia_ban')), 'ctkm': '', 'la_combo': True, 'dang_ban': True,
                      'row_mau': None}
                db['products'].append(np)
                db['combos'].setdefault(np['id'], {'components': [], 'choice': None, 'gift': []})
                keep_combo.add(np['id'])
        if body.get('combos_full'):  # xóa combo vắng mặt (có kiểm tham chiếu)
            for p in list(db['products']):
                if p['la_combo'] and p['id'] not in keep_combo:
                    refs = _catalog_refs(db, p['id'])
                    if refs:
                        raise ValueError(f'Không xóa được combo "{p["ten"]}" — đang dùng ở: {", ".join(refs)}.')
                    db['products'].remove(p)
                    db['combos'].pop(p['id'], None)
    if 'weights' in body:
        for sp, w in body['weights'].items():
            db['weights'][sp] = {'quy_cach_tt': num(w.get('quy_cach_tt')), 'kg_thung': num(w.get('kg_thung'))}
    if 'combos' in body:
        db['combos'].update(body['combos'])
    if 'provinces' in body:
        db['provinces'] = [norm(t) for t in body['provinces'] if norm(t)]


# =====================================================================
# máy chủ web
# =====================================================================

def license_info():
    return dict(LICENSE.check_license().to_dict(), policy=licensing.POLICY_TEXT, data_file=str(DATA_FILE),
                app=APP_ID)


class Handler(BaseHTTPRequestHandler):
    db = None

    def log_message(self, *a):
        pass

    def _send(self, code, payload, ctype='application/json; charset=utf-8'):
        data = payload if isinstance(payload, bytes) else json.dumps(payload, ensure_ascii=False, default=str).encode('utf-8')
        self.send_response(code)
        self.send_header('Content-Type', ctype)
        self.send_header('Cache-Control', 'no-store')
        self.end_headers()
        self.wfile.write(data)

    @staticmethod
    def _fresh():
        """Người dùng sửa thẳng file dữ liệu bằng Excel -> đọc lại trước khi làm tiếp."""
        if Handler.db is None:  # chưa có license lúc mở công cụ -> nạp dữ liệu khi license đã hợp lệ
            Handler.db = load_db()
        elif DATA_FILE.exists() and DATA_FILE.stat().st_mtime != DATA_MTIME[0]:
            Handler.db = load_db()

    def _host_ok(self):
        """Chỉ phục vụ yêu cầu gửi tới 127.0.0.1 / localhost: chặn trang web lạ dùng tên miền trỏ về máy này
        (DNS rebinding) để đọc dữ liệu."""
        host = (self.headers.get('Host') or '').strip().lower()
        host = host[1:host.index(']')] if host.startswith('[') and ']' in host else host.rsplit(':', 1)[0]
        if host in ('127.0.0.1', 'localhost', '::1'):
            return True
        self._send(403, {'error': 'Yêu cầu không hợp lệ.'})
        return False

    def _read_json(self, limit=MAX_BODY):
        """Body JSON của POST. Bắt buộc Content-Type application/json: trang web lạ không gửi được yêu cầu loại này
        tới công cụ mà không qua kiểm tra CORS của trình duyệt (chống CSRF). Sai -> ValueError."""
        if not (self.headers.get('Content-Type') or '').startswith('application/json'):
            raise ValueError('content-type')
        n = int(self.headers.get('Content-Length') or 0)
        if n < 0 or n > limit:
            raise ValueError('size')
        body = json.loads(self.rfile.read(n) or b'{}')
        if not isinstance(body, dict):
            raise ValueError('body')
        return body

    def _license_blocked(self):
        """License không hợp lệ -> trả 403 kèm trạng thái; dữ liệu không bị đọc/ghi."""
        st = LICENSE.check_license()
        if st.valid:
            return False
        self._send(403, {'error': st.message, 'license': license_info()})
        return True

    def _activate(self):
        # bắt buộc JSON: trang web lạ không gửi được yêu cầu này mà không qua kiểm tra CORS của trình duyệt
        try:
            body = self._read_json(limit=8192)
        except ValueError as e:
            return self._send(415 if str(e) == 'content-type' else 400, {'error': 'Yêu cầu không hợp lệ.'})
        try:
            st = LICENSE.activate(str(body.get('code') or ''))
        except licensing.ActivationError as e:
            return self._send(400, {'error': str(e), 'reason': e.reason, 'license': license_info()})
        except (ValueError, AttributeError):
            return self._send(400, {'error': 'Yêu cầu không hợp lệ.'})
        print('[License]', st.message)
        self._send(200, {'ok': True, 'license': license_info()})

    def do_GET(self):
        if not self._host_ok():
            return
        u = urlparse(self.path)
        if u.path == '/api/license':
            return self._send(200, license_info())
        if u.path == '/api/state':
            if self._license_blocked():
                return
            with LOCK:
                try:
                    self._fresh()
                except PermissionError:
                    pass
                except Exception as e:  # nạp dữ liệu lần đầu lỗi (thiếu file Excel gốc...)
                    return self._send(500, {'error': f'Không đọc được dữ liệu: {e}'})
                if Handler.db is None:
                    return self._send(500, {'error': 'Chưa đọc được dữ liệu — đóng file dữ liệu trong Excel rồi tải lại trang.'})
                return self._send(200, public_state(Handler.db))
        name = 'index.html' if u.path in ('/', '') else u.path.lstrip('/')
        f = (WEB_DIR / name).resolve()
        if WEB_DIR.resolve() not in f.parents or not f.is_file():
            return self._send(404, {'error': 'Không tìm thấy'})
        ctype = {'.html': 'text/html', '.css': 'text/css', '.js': 'application/javascript'}.get(f.suffix, 'application/octet-stream')
        self._send(200, f.read_bytes(), ctype + '; charset=utf-8')

    @staticmethod
    def _reload():
        """Bỏ mọi thay đổi chưa lưu trong bộ nhớ: đọc lại file dữ liệu (nạp lại khi cần nếu đọc lỗi)."""
        try:
            Handler.db = load_db() if DATA_FILE.exists() else None
        except Exception:
            Handler.db = None

    def do_POST(self):
        if not self._host_ok():
            return
        u = urlparse(self.path)
        if u.path == '/api/license/activate':
            return self._activate()
        if self._license_blocked():
            return
        try:
            body = self._read_json()
        except ValueError as e:
            return self._send(415 if str(e) == 'content-type' else 400, {'error': 'Yêu cầu không hợp lệ.'})
        try:
            with LOCK:
                self._fresh()
                res = handle(Handler.db, u.path, body)
            self._send(200, res)
        except PermissionError:
            with LOCK:
                self._reload()
            self._send(400, {'error': 'File Excel đang mở (file dữ liệu hoặc file xuất) — đóng file trong Excel rồi làm lại.'})
        except (ValueError, RuntimeError, StopIteration, KeyError) as e:
            with LOCK:  # thao tác bị từ chối: không để lại thay đổi dở dang trong bộ nhớ
                self._reload()
            self._send(400, {'error': str(e) or 'Dữ liệu không hợp lệ.'})
        except Exception as e:  # lỗi bất ngờ (file dữ liệu hỏng, đang lưu dở...): vẫn trả lời giao diện
            with LOCK:
                self._reload()
            self._send(500, {'error': f'Lỗi không mong muốn: {e}. Thao tác chưa được lưu; dữ liệu đã đọc lại từ file.'})


class Server(ThreadingHTTPServer):
    """Chiếm cổng độc quyền: trên Windows, SO_REUSEADDR mặc định cho 2 tiến trình nghe cùng 1 cổng (mỗi bên
    giữ dữ liệu riêng trong bộ nhớ -> ghi đè lẫn nhau)."""
    allow_reuse_address = False

    def server_bind(self):
        if hasattr(socket, 'SO_EXCLUSIVEADDRUSE'):
            self.socket.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
        super().server_bind()


def already_running(port):
    """Cổng đang bị chiếm bởi chính công cụ này (mở ở cửa sổ khác)?"""
    try:
        with urllib.request.urlopen(f'http://127.0.0.1:{port}/api/license', timeout=2) as r:
            return json.loads(r.read()).get('app') == APP_ID
    except Exception:
        return False


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    print('Đang mở công cụ Quản lý Kênh Viettel...')
    lic = LICENSE.check_license()
    print('[License]', lic.message)
    if lic.valid:
        try:
            Handler.db = load_db()
        except Exception as e:
            print('(!) Không đọc được dữ liệu:', e)
            input('Bấm Enter để đóng...')
            return
    else:  # chưa/hết license: chỉ mở màn hình nhập mã, không đọc hay ghi dữ liệu
        print('Nhập mã kích hoạt trên giao diện. Mã thiết bị của máy này:', lic.device_id or '-')
        print('Dữ liệu trong thư mục "Dữ liệu" được giữ nguyên.')
    server = None
    for port in range(PORT, PORT + 20):
        try:
            server = Server(('127.0.0.1', port), Handler)
            break
        except OSError:
            if already_running(port):  # đã mở ở cửa sổ khác: dùng lại, không chạy bản thứ hai
                url = f'http://127.0.0.1:{port}/'
                print('Công cụ ĐANG MỞ ở một cửa sổ khác — mở lại giao diện đó:', url)
                if '--khong-mo' not in sys.argv:
                    webbrowser.open(url)
                return
            continue
    if server is None:
        print(f'(!) Không mở được cổng {PORT}–{PORT + 19} (đang bị chương trình khác dùng). Đóng bớt chương trình rồi mở lại.')
        input('Bấm Enter để đóng...')
        return
    url = f'http://127.0.0.1:{port}/'
    print('Giao diện:', url)
    print('ĐỪNG ĐÓNG cửa sổ này trong lúc dùng công cụ. Dùng xong thì đóng cửa sổ này.')
    if '--khong-mo' not in sys.argv:
        threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    server.serve_forever()


if __name__ == '__main__':
    main()
