"""Báo cáo SỐ LƯỢNG hàng TikTok theo từng ngày (không có số tiền).

Ngày D = các đơn ĐẶT từ 14h ngày D-1 đến trước 14h ngày D (giờ chốt đổi ở CUTOFF_HOUR).
Cách dùng:
  - Nhấp đúp Tao-bao-cao.bat  -> đọc tất cả file "Tất cả đơn hàng*.xlsx" nằm trong thư mục công cụ
  - Hoặc kéo 1 hay nhiều file TikTok thả vào Tao-bao-cao.bat
Báo cáo lưu vào thư mục "Báo cáo" cạnh công cụ.
"""
import math
import os
import re
import subprocess
import sys
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent
# openpyxl đi kèm trong thư mục "thu-vien" để chạy được trên máy chưa cài thư viện này
sys.path.insert(0, str(ROOT / 'thu-vien'))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

import licensing  # cấp phép offline 48 giờ (bản sao trong thu-vien)

CUTOFF_HOUR = 14  # giờ chốt đơn mỗi ngày
CANCELLED = {'Đã hủy', 'Cancelled'}
UNPAID = {'Chưa thanh toán', 'Unpaid'}
RETURN_TYPE = 'Return/Refund'
OUT_DIR = ROOT / 'Báo cáo'
# không dùng nhãn ">5": trong COUNTIFS Excel hiểu là phép so sánh "lớn hơn 5"
BUCKETS = ['1', '2', '3', '4', '5', 'Trên 5', 'Chỉ tặng']
REQUIRED = {'Order ID', 'Order Status', 'Product Name', 'Quantity', 'SKU Unit Original Price', 'SKU ID',
            'Sku Quantity of return', 'Cancelation/Return Type', 'Created Time', 'Shipped Time'}

SUMMARY, DATA, HOW = 'Tổng hợp theo ngày', 'Dữ liệu chi tiết', 'Cách tính'

NAVY, BLUE, LIGHT, BAND = '1F3864', '1F4E78', 'DDEBF7', 'F5F8FC'
TOTAL, GREEN, RED, GRAY, LINE, WARN_BG = 'D9E1F2', 'E2EFDA', 'C00000', '595959', 'BFBFBF', 'FCE4D6'
NUM = '#,##0;-#,##0;"-"'
THIN = Side(style='thin', color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
THICK = Border(left=THIN, right=THIN, top=Side(style='medium', color=NAVY), bottom=Side(style='medium', color=NAVY))
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
RIGHT = Alignment(horizontal='right', vertical='center')


def font(size=10, bold=False, color='262626', italic=False, underline=None):
    return Font(name='Arial', size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(color):
    return PatternFill('solid', fgColor=color)


def num(x):
    return float(x) if x not in (None, '') else 0.0


def parse_time(s):
    return datetime.strptime(s, '%d/%m/%Y %H:%M:%S') if s else None


def short_name(full):
    # Tên listing TikTok thường là "Tên sản phẩm – mô tả quảng cáo": chỉ giữ phần trước dấu "–"
    head = re.split(r'\s+[–—]\s+', full, maxsplit=1)[0].strip()
    return head if len(head) >= 10 else full


def report_day(created):
    """Đơn đặt từ 14h ngày D-1 đến trước 14h ngày D thuộc ngày D."""
    return (created + timedelta(hours=24 - CUTOFF_HOUR)).date()


def export_time(src):
    m = re.search(r'(\d{4})-(\d{2})-(\d{2})-(\d{2})_(\d{2})', src.stem)  # VD: ...-2026-10-09-09_19
    return datetime(*map(int, m.groups())) if m else datetime.fromtimestamp(src.stat().st_mtime)


# ---------- đọc & gộp nhiều file ----------

def read_file(src):
    # không dùng read_only: với file TikTok, chế độ đó chỉ đọc được cột A
    wb = load_workbook(src, data_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    hdr = list(next(it))
    missing = REQUIRED - set(hdr)
    if missing:
        wb.close()
        raise ValueError(f'File "{src.name}" không phải file "Tất cả đơn hàng" của TikTok (thiếu cột: {", ".join(sorted(missing))})')
    data = [d for d in (dict(zip(hdr, r)) for r in it) if str(d.get('Order ID') or '').isdigit()]
    wb.close()
    return data


def load_all(files):
    """Gộp các file. Đơn có trong nhiều file thì lấy theo file xuất MỚI NHẤT (trạng thái mới hơn).
    Trả về (các dòng, các khoảng thời gian đặt đơn mà dữ liệu bao phủ, danh sách file)."""
    by_order, covered, used = {}, [], []
    for src in sorted(files, key=export_time):
        data = read_file(src)
        if not data:
            continue
        lines = defaultdict(list)
        for d in data:
            lines[d['Order ID']].append(d)
        for oid, ls in lines.items():
            by_order[oid] = (ls, src.name)
        created = [parse_time(d['Created Time']) for d in data]
        # file TikTok xuất theo NGÀY đặt đơn: coi như bao phủ trọn các ngày từ ngày đặt sớm nhất đến muộn nhất
        start = datetime.combine(min(created).date(), datetime.min.time())
        end = min(datetime.combine(max(created).date() + timedelta(days=1), datetime.min.time()), export_time(src))
        covered.append((start, end))
        used.append((src.name, start, end, len(lines)))
    rows = [(d, name) for ls, name in by_order.values() for d in ls]
    return rows, merge(covered), used


def merge(intervals):
    out = []
    for s, e in sorted(intervals):
        if out and s <= out[-1][1]:
            out[-1] = (out[-1][0], max(out[-1][1], e))
        else:
            out.append((s, e))
    return out


def coverage(day, covered):
    """Khung giờ của ngày + dữ liệu có phủ trọn khung giờ đó không."""
    end = datetime.combine(day, datetime.min.time()) + timedelta(hours=CUTOFF_HOUR)
    start = end - timedelta(days=1)
    full = any(s <= start and end <= e for s, e in covered)
    return start, end, full


def to_rows(raw):
    rows = []
    for d, src_name in raw:
        created, shipped = parse_time(d['Created Time']), parse_time(d['Shipped Time'])
        status = d['Order Status']
        canc = status in CANCELLED
        if status in UNPAID:
            group = 'Chưa thanh toán'
        elif canc and not shipped:
            group = 'Hủy trước khi gửi'
        else:
            group = 'Xuất'  # gồm đơn chờ lấy hàng, đang giao, đã giao, giao thất bại, khách trả hàng
        is_ret = d['Cancelation/Return Type'] == RETURN_TYPE
        full = re.sub(r'\s+', ' ', d['Product Name']).strip()
        rows.append(dict(
            order=d['Order ID'], created=created, day=datetime.combine(report_day(created), datetime.min.time()),
            shipped=shipped, status=status, substatus=d['Order Substatus'], rtype=d['Cancelation/Return Type'] or '',
            reason=d['Cancel Reason'] or '', prod=short_name(full), full=full, sku=d['SKU ID'], qty=num(d['Quantity']),
            kind='Tặng' if num(d['SKU Unit Original Price']) == 0 else 'Bán',
            retq=num(d['Sku Quantity of return']) if is_ret else 0, group=group,
            back='Giao thất bại' if canc and shipped else ('Khách trả hàng' if is_ret else ''),
            bucket='', cnt_po=0, cnt_o=0, src=src_name))
    rows.sort(key=lambda r: (r['created'], r['order']))
    # số chai khách MUA (không tính chai tặng) của từng mặt hàng trong từng đơn
    paid, first_po, seen_o = defaultdict(float), {}, set()
    for i, r in enumerate(rows):
        if r['order'] not in seen_o:
            seen_o.add(r['order'])
            r['cnt_o'] = 1
        if r['group'] == 'Xuất':
            po = (r['order'], r['prod'])
            first_po.setdefault(po, i)
            if r['kind'] == 'Bán':
                paid[po] += r['qty']
    for po, i in first_po.items():
        q = paid[po]
        rows[i]['cnt_po'] = 1
        rows[i]['bucket'] = 'Chỉ tặng' if q == 0 else ('Trên 5' if q > 5 else str(int(q)))
    return rows


# ---------- bố cục ----------

DATA_COLS = [
    ('day', 'Ngày báo cáo', 12, 'dd/mm/yyyy'), ('order', 'Mã đơn', 20, None),
    ('created', 'Giờ đặt đơn', 16, 'dd/mm/yyyy hh:mm'), ('shipped', 'Giờ giao ĐVVC', 16, 'dd/mm/yyyy hh:mm'),
    ('status', 'Trạng thái', 15, None), ('substatus', 'Trạng thái phụ', 18, None), ('rtype', 'Hủy/Hoàn', 13, None),
    ('reason', 'Lý do hủy', 24, None), ('prod', 'Mặt hàng', 40, None), ('full', 'Tên đầy đủ trên TikTok', 50, None),
    ('sku', 'SKU ID', 21, None), ('qty', 'Số lượng', 9, NUM), ('kind', 'Bán/Tặng', 10, None),
    ('retq', 'SL khách trả lại', 11, NUM), ('group', 'Phân loại', 17, None), ('back', 'Hàng hoàn về', 15, None),
    ('bucket', 'Số chai khách mua', 12, None), ('src', 'Lấy từ file', 34, None),
    ('cnt_po', 'Cột phụ: đếm đơn theo mặt hàng', 12, None), ('cnt_o', 'Cột phụ: đếm đơn', 12, None),
]
HELPERS = {'cnt_po', 'cnt_o'}
DCOL = {k: get_column_letter(i) for i, (k, *_) in enumerate(DATA_COLS, 1)}

DAY_COLS = [
    ('stt', 'STT', 6), ('prod', 'Mặt hàng', 42), ('orders', 'Số đơn', 9),
    ('out', 'Chai xuất', 10), ('sold', 'Chai bán', 10), ('gift', 'Chai tặng', 10),
    ('b1', '1 chai', 7), ('b2', '2 chai', 7), ('b3', '3 chai', 7), ('b4', '4 chai', 7), ('b5', '5 chai', 7),
    ('b6', 'Trên 5 chai', 8), ('b7', 'Chỉ nhận tặng', 9),
    ('ret', 'Khách trả lại', 10), ('fail', 'Giao thất bại', 10), ('cancel', 'Chai trong đơn hủy', 10),
]
DIDX = {k: i for i, (k, *_) in enumerate(DAY_COLS, 1)}
DAY_GROUPS = [('SỐ CHAI XUẤT KHO', 'out', 'gift'), ('SỐ ĐƠN THEO SỐ CHAI KHÁCH MUA', 'b1', 'b7'),
              ('HÀNG HOÀN VỀ (CHAI)', 'ret', 'fail')]

SUM_COLS = [
    ('day', 'Ngày', 12), ('window', 'Khung giờ đặt đơn', 30), ('state', 'Dữ liệu', 20), ('orders', 'Số đơn', 12),
    ('out', 'Chai xuất', 12), ('sold', 'Chai bán', 11), ('gift', 'Chai tặng', 11), ('ret', 'Chai khách trả lại', 11),
    ('fail', 'Chai giao thất bại', 11), ('cancel', 'Đơn hủy trước khi gửi', 11), ('unpaid', 'Đơn chưa thanh toán', 11),
]


def rng(key, n):
    c = DCOL[key]
    return f"'{DATA}'!${c}$2:${c}${n}"


def day_sheet_name(day):
    return f'Ngày {day:%d.%m}'


def row_height(text, width, base=15):
    # ước lượng số dòng khi xuống dòng: ~0.95 ký tự / 1 đơn vị độ rộng cột (Arial 10)
    return base * max(1, math.ceil(len(text) / (width * 0.95))) + 4


def setup_print(ws, title_rows=None):
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    if title_rows:
        ws.print_title_rows = title_rows


def header_cell(c, light=False):
    c.font = font(10, True, NAVY if light else 'FFFFFF')
    c.fill = fill(LIGHT if light else BLUE)
    c.alignment, c.border = CENTER, BOX


def window_text(start, end):
    return f'{start:%Hh %d/%m} → trước {end:%Hh %d/%m/%Y}'


def sheet_data(wb, rows):
    ws = wb.create_sheet(DATA)
    ws.sheet_properties.tabColor = '7F7F7F'
    ws.append([h for _, h, *_ in DATA_COLS])
    for r in rows:
        ws.append([r[k] for k, *_ in DATA_COLS])
    n = len(rows) + 1
    for i, (k, _, w, fmt) in enumerate(DATA_COLS, 1):
        header_cell(ws.cell(1, i))
        if k in HELPERS:
            ws.cell(1, i).fill = fill('A6A6A6')
        dim = ws.column_dimensions[get_column_letter(i)]
        dim.width = w
        if k in HELPERS:  # cột phụ để tính công thức: ẩn đi, bấm dấu + phía trên để mở
            dim.outlineLevel, dim.hidden = 1, True
        for (cell,) in ws.iter_rows(min_row=2, max_row=n, min_col=i, max_col=i):
            cell.font = font()
            if fmt:
                cell.number_format = fmt
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(DATA_COLS))}{n}'
    return n


def sheet_day(wb, rows, n, day, cov):
    start, end, full = cov
    ws = wb.create_sheet(day_sheet_name(day))
    ws.sheet_properties.tabColor = BLUE if full else 'ED7D31'
    ws.sheet_view.showGridLines = False
    ws['A1'] = f'SỐ LƯỢNG HÀNG TIKTOK — NGÀY {day:%d/%m/%Y}'
    ws['A1'].font = font(14, True, NAVY)
    ws.row_dimensions[1].height = 24
    ws['A2'] = f'Đơn đặt từ {window_text(start, end)}  •  Không tính đơn chưa thanh toán'
    ws['A2'].font = font(9, color=GRAY, italic=True)
    ws['A3'] = ('Dữ liệu đủ cả khung giờ.' if full else
                'CHƯA ĐỦ DỮ LIỆU: thiếu file TikTok của một phần khung giờ này — số liệu đang thiếu, cần thêm file rồi chạy lại.')
    ws['A3'].font = font(10, True, '2C6A3D' if full else RED)

    grouped = set()
    for label, a, b in DAY_GROUPS:
        for ci in range(DIDX[a], DIDX[b] + 1):
            grouped.add(ci)
            ws.cell(4, ci).fill, ws.cell(4, ci).border = fill(BLUE), BOX
        ws.cell(4, DIDX[a], label).font, ws.cell(4, DIDX[a]).alignment = font(10, True, 'FFFFFF'), CENTER
        ws.merge_cells(start_row=4, start_column=DIDX[a], end_row=4, end_column=DIDX[b])
    for i, (k, h, w) in enumerate(DAY_COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        if i in grouped:
            header_cell(ws.cell(5, i, h), light=True)
        else:
            for rr in (4, 5):
                ws.cell(rr, i).fill, ws.cell(rr, i).border = fill(BLUE), BOX
            header_cell(ws.cell(4, i, h))
            ws.merge_cells(start_row=4, start_column=i, end_row=5, end_column=i)
    ws.row_dimensions[4].height = 30
    ws.row_dimensions[5].height = 34

    out = defaultdict(float)
    for r in rows:
        if r['day'].date() == day and r['group'] != 'Chưa thanh toán':
            out[r['prod']] += r['qty'] if r['group'] == 'Xuất' else 0
    prods = sorted(out, key=lambda p: (-out[p], p))
    first = 6
    if not prods:
        ws.cell(first, 2, 'Không có đơn nào trong khung giờ này.').font = font(10, italic=True)
        return
    last, tr = first + len(prods) - 1, first + len(prods)
    dcell = f'DATE({day.year},{day.month},{day.day})'
    for k, p in enumerate(prods):
        r = first + k
        base = f'{rng("day", n)},{dcell},{rng("prod", n)},$B{r}'
        xuat = f'{base},{rng("group", n)},"Xuất"'
        f = {
            'stt': k + 1, 'prod': p,
            'orders': f'=SUMIFS({rng("cnt_po", n)},{xuat})',
            'out': f'=SUMIFS({rng("qty", n)},{xuat})',
            'sold': f'=SUMIFS({rng("qty", n)},{xuat},{rng("kind", n)},"Bán")',
            'gift': f'=SUMIFS({rng("qty", n)},{xuat},{rng("kind", n)},"Tặng")',
            'ret': f'=SUMIFS({rng("retq", n)},{xuat})',
            'fail': f'=SUMIFS({rng("qty", n)},{xuat},{rng("back", n)},"Giao thất bại")',
            'cancel': f'=SUMIFS({rng("qty", n)},{base},{rng("group", n)},"Hủy trước khi gửi")',
        }
        for j, b in enumerate(BUCKETS, 1):
            f[f'b{j}'] = f'=COUNTIFS({xuat},{rng("bucket", n)},"{b}")'
        for i, (key, *_) in enumerate(DAY_COLS, 1):
            ws.cell(r, i, f[key])
        ws.row_dimensions[r].height = row_height(p, DAY_COLS[1][2])
    ws.cell(tr, 2, 'TỔNG CỘNG')
    ws.cell(tr, 3, f'=SUMIFS({rng("cnt_o", n)},{rng("day", n)},{dcell},{rng("group", n)},"Xuất")')  # số đơn thực
    for key, *_ in DAY_COLS[3:]:
        c = get_column_letter(DIDX[key])
        ws[f'{c}{tr}'] = f'=SUM({c}{first}:{c}{last})'
    ws.row_dimensions[tr].height = 22
    for r in range(first, tr + 1):
        for i, (key, *_) in enumerate(DAY_COLS, 1):
            c = ws.cell(r, i)
            c.border = THICK if r == tr else BOX
            c.alignment = LEFT if key == 'prod' else (CENTER if key == 'stt' else RIGHT)
            c.number_format = NUM
            c.font = font(10, r == tr or key == 'out', RED if key in ('ret', 'fail') else '262626')
            if r == tr:
                c.fill = fill(TOTAL)
            elif key == 'out':
                c.fill = fill(GREEN)
            elif (r - first) % 2:
                c.fill = fill(BAND)
    notes = ['Chai xuất = chai bán + chai tặng của các đơn không bị hủy trước khi gửi (gồm đơn đang chờ lấy hàng).',
             'Số đơn theo số chai khách mua: tính số chai khách TRẢ TIỀN của mặt hàng đó trong đơn, không tính chai tặng. '
             '7 cột cộng lại bằng "Số đơn".',
             'Combo nhiều chai được đếm theo số combo trên TikTok. Giải thích đầy đủ: sheet "Cách tính".']
    for k, t in enumerate(notes):
        ws.cell(tr + 2 + k, 2, t).font = font(9, color=GRAY, italic=True)
    ws.freeze_panes = 'C6'
    setup_print(ws, '4:5')


def sheet_summary(ws, n, days, covs, files, made_at):
    ws.title = SUMMARY
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    ws['A1'] = 'SỐ LƯỢNG HÀNG TIKTOK THEO NGÀY'
    ws['A1'].font = font(16, True, NAVY)
    ws.row_dimensions[1].height = 28
    ws['A2'] = (f'Ngày D = đơn đặt từ {CUTOFF_HOUR}h ngày D-1 đến trước {CUTOFF_HOUR}h ngày D  •  '
                f'Lập lúc {made_at:%d/%m/%Y %H:%M}  •  Đơn vị: chai')
    ws['A2'].font = font(9, color=GRAY, italic=True)
    hr = 4
    for i, (k, h, w) in enumerate(SUM_COLS, 1):
        header_cell(ws.cell(hr, i, h))
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[hr].height = 46
    first = hr + 1
    for k, day in enumerate(days):
        r = first + k
        start, end, full = covs[day]
        dcell = f'$A{r}'
        cnt = lambda val: f'=SUMIFS({rng("cnt_o", n)},{rng("day", n)},{dcell},{rng("group", n)},"{val}")'
        x = f'{rng("day", n)},{dcell},{rng("group", n)},"Xuất"'
        vals = {
            'day': datetime.combine(day, datetime.min.time()), 'window': window_text(start, end),
            'state': 'Đủ dữ liệu' if full else 'Chưa đủ dữ liệu',
            'orders': cnt('Xuất'), 'out': f'=SUMIFS({rng("qty", n)},{x})',
            'sold': f'=SUMIFS({rng("qty", n)},{x},{rng("kind", n)},"Bán")',
            'gift': f'=SUMIFS({rng("qty", n)},{x},{rng("kind", n)},"Tặng")',
            'ret': f'=SUMIFS({rng("retq", n)},{x})',
            'fail': f'=SUMIFS({rng("qty", n)},{x},{rng("back", n)},"Giao thất bại")',
            'cancel': cnt('Hủy trước khi gửi'), 'unpaid': cnt('Chưa thanh toán'),
        }
        for i, (key, *_) in enumerate(SUM_COLS, 1):
            c = ws.cell(r, i, vals[key])
            c.border, c.font = BOX, font(10, key in ('day', 'out'), RED if key in ('ret', 'fail') else '262626')
            c.alignment = CENTER if key in ('day', 'state') else (LEFT if key == 'window' else RIGHT)
            c.number_format = 'dd/mm/yyyy' if key == 'day' else NUM
            if key == 'out':
                c.fill = fill(GREEN)
            elif not full:
                c.fill = fill(WARN_BG)
            elif k % 2:
                c.fill = fill(BAND)
        if not full:
            ws.cell(r, 3).font = font(10, True, RED)
        # bấm vào ngày để mở sheet của ngày đó
        link = ws.cell(r, 1)
        link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{day_sheet_name(day)}'!A1")
        link.font = font(10, True, '0563C1', underline='single')
    last = first + len(days) - 1
    r = last + 2
    ws.cell(r, 1, 'Bấm vào ngày để xem chi tiết từng mặt hàng. Dòng tô cam = chưa đủ dữ liệu (thiếu file của một phần khung giờ), '
                  'cần thêm file TikTok rồi chạy lại.').font = font(9, color=GRAY, italic=True)
    r += 2
    ws.cell(r, 1, 'FILE ĐÃ ĐỌC').font = font(11, True, NAVY)
    for i, h in enumerate(['File', 'Đơn đặt từ', 'Đến trước', 'Số đơn'], 1):
        header_cell(ws.cell(r + 1, i if i == 1 else i + 2, h))
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=3)
    for k, (name, s, e, cnt_orders) in enumerate(files):
        rr = r + 2 + k
        ws.cell(rr, 1, name).font = font()
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3)
        for col, v, fmt in ((4, s, 'dd/mm hh"h"mm'), (5, e, 'dd/mm hh"h"mm'), (6, cnt_orders, NUM)):
            c = ws.cell(rr, col, v)
            c.font, c.number_format, c.alignment = font(), fmt, RIGHT
    ws.freeze_panes = f'A{first}'
    setup_print(ws, f'{hr}:{hr}')


def sheet_howto(wb):
    ws = wb.create_sheet(HOW)
    ws.sheet_properties.tabColor = '7F7F7F'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 100
    ws['B1'] = 'CÁCH TÍNH CÁC CON SỐ'
    ws['B1'].font = font(14, True, NAVY)
    # câu chữ không được bắt đầu bằng "=": openpyxl sẽ coi là công thức và Excel báo file hỏng
    items = [
        ('Ngày', f'Ngày D gồm các đơn khách ĐẶT từ {CUTOFF_HOUR}h00 ngày D-1 đến {CUTOFF_HOUR - 1}h59 ngày D. '
                 f'VD ngày 09/10: đơn đặt từ {CUTOFF_HOUR}h 08/10 đến trước {CUTOFF_HOUR}h 09/10.'),
        ('Dữ liệu đủ / chưa đủ', 'Mỗi file TikTok được coi là phủ trọn các ngày đặt đơn có trong file. Khung giờ của một ngày '
                                 'thường nằm trên 2 file (chiều hôm trước + sáng hôm đó); thiếu 1 file thì ngày đó "chưa đủ dữ liệu".'),
        ('Đơn trùng giữa các file', 'Lấy theo file xuất mới nhất (trạng thái đơn mới hơn), không đếm 2 lần.'),
        ('Chai xuất', 'Chai bán + chai tặng của các đơn KHÔNG bị hủy trước khi gửi: gồm đơn chờ lấy hàng, đang giao, đã giao, '
                      'giao thất bại, khách trả hàng.'),
        ('Chai tặng / Chai bán', 'Chai tặng = dòng có giá gốc 0đ trên TikTok (chỉ dùng để nhận diện, báo cáo không hiện số tiền). Còn lại là chai bán.'),
        ('Số đơn', 'Số đơn có chứa mặt hàng đó. Một đơn nhiều mặt hàng được đếm ở mỗi mặt hàng; dòng TỔNG CỘNG là số đơn thực.'),
        ('1 / 2 / 3 / 4 / 5 / Trên 5 chai', 'Số đơn mà khách MUA đúng chừng ấy chai của mặt hàng đó (cộng các dòng cùng mặt hàng trong đơn, không tính chai tặng).'),
        ('Chỉ nhận tặng', 'Số đơn chỉ được tặng mặt hàng này. 7 cột số đơn cộng lại bằng "Số đơn".'),
        ('Khách trả lại', 'Số chai TikTok ghi nhận khách trả hàng (đơn loại Return/Refund).'),
        ('Giao thất bại', 'Số chai của đơn đã giao cho đơn vị vận chuyển nhưng sau đó bị hủy (giao không thành công, hàng quay về).'),
        ('Chai trong đơn hủy', 'Số chai của đơn bị hủy trước khi gửi hàng — không tính vào chai xuất.'),
        ('Đơn chưa thanh toán', 'Đơn khách chưa trả tiền online — chưa chắc sẽ gửi, không tính vào chai xuất.'),
        ('Lưu ý', 'Hoàn hàng / giao thất bại thường xảy ra vài ngày sau. Muốn số hoàn của một ngày cho chính xác, xuất lại file '
                  'TikTok của ngày đó sau vài ngày rồi chạy lại.'),
        ('Combo', 'Combo nhiều chai được đếm theo số combo trên TikTok (1 combo = 1).'),
    ]
    r = 3
    for k, (name, text) in enumerate(items):
        a, b = ws.cell(r, 2, name), ws.cell(r, 3, text)
        a.font, b.font = font(10, True, NAVY), font(10)
        a.alignment = b.alignment = LEFT
        a.border = b.border = BOX
        if k % 2:
            a.fill = b.fill = fill(BAND)
        ws.row_dimensions[r].height = row_height(text, 100, 14)
        r += 1
    setup_print(ws)
    ws.page_setup.fitToHeight = 1


def build_report(rows, covered, files, out, made_at):
    days = sorted({r['day'].date() for r in rows})
    covs = {d: coverage(d, covered) for d in days}
    wb = Workbook()
    summary = wb.active
    n = sheet_data(wb, rows)
    for d in days:
        sheet_day(wb, rows, n, d, covs[d])
    sheet_summary(summary, n, days, covs, files, made_at)
    sheet_howto(wb)
    wb.move_sheet(DATA, offset=len(wb.sheetnames))  # dữ liệu chi tiết nằm cuối
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    return days, covs


def recalc_with_excel(path):
    """Cho Excel tính sẵn công thức rồi lưu, để file xem được cả trên điện thoại / Zalo."""
    ps = ('$x = New-Object -ComObject Excel.Application; $x.DisplayAlerts = $false; '
          'try { $w = $x.Workbooks.Open($env:BAO_CAO_PATH); $x.CalculateFull(); $w.Save(); $w.Close($false) } '
          'finally { $x.Quit(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($x) }')
    res = subprocess.run(['powershell', '-NoProfile', '-NonInteractive', '-Command', ps],
                         env=dict(os.environ, BAO_CAO_PATH=str(path)), capture_output=True)
    return res.returncode == 0


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    if not licensing.ensure_licensed('Số lượng TikTok theo ngày – V3'):
        return
    files = [Path(a) for a in sys.argv[1:]] or sorted(ROOT.glob('Tất cả đơn hàng*.xlsx'))
    files = [f for f in files if f.is_file() and not f.name.startswith('~$')]
    if not files:
        print(f'Không có file TikTok nào. Bỏ file "Tất cả đơn hàng….xlsx" vào thư mục {ROOT} rồi chạy lại.')
        return
    print('Đang đọc', len(files), 'file:')
    for f in files:
        print('  -', f.name)
    try:
        raw, covered, used = load_all(files)
    except ValueError as e:
        print('(!)', e)
        return
    rows = to_rows(raw)
    made_at = datetime.now()
    days = sorted({r['day'].date() for r in rows})
    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / f'Số lượng TikTok theo ngày {days[0]:%d.%m} - {days[-1]:%d.%m.%Y} (lập {made_at:%d.%m.%Y %Hh%M}).xlsx'
    print('Đang tạo báo cáo...')
    try:
        days, covs = build_report(rows, covered, used, out, made_at)
    except PermissionError:
        print(f'Không ghi được — file "{out.name}" đang mở trong Excel. Đóng file rồi chạy lại.')
        return
    print('Đang cho Excel tính sẵn số liệu...')
    if not recalc_with_excel(out):
        print('(!) Không tự tính được bằng Excel. Hãy mở file bằng Excel, bấm Lưu rồi mới gửi đi.')
    print()
    for d in days:
        s, e, full = covs[d]
        print(f'  Ngày {d:%d/%m}: {"đủ dữ liệu" if full else "CHƯA ĐỦ dữ liệu (thiếu file)"}  ({window_text(s, e)})')
    print()
    print('XONG! Báo cáo đã lưu tại:')
    print(f'  {out}')
    subprocess.Popen(f'explorer /select,"{out}"')
    os.startfile(out)


if __name__ == '__main__':
    main()
