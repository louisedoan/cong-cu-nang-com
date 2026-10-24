"""Công cụ Doanh thu - Đơn hàng - Công nợ (kênh NPP).

Chạy: nhấp đúp "Mo-cong-cu.bat" -> mở giao diện trên trình duyệt (chỉ máy này truy cập được).
- Lần đầu: tự đọc các file Excel trong thư mục cha (D:\\Công cụ Nàng Cơm\\DT - DH - CN): mẫu đơn (sheet "Đơn Đặt Hàng")
  và file NPP/công nợ (sheet "Công nợ chốt").
- Mọi dữ liệu lưu trong 1 file Excel: "Dữ liệu\\Dữ liệu DT - ĐH - CN.xlsx" (mở xem được; trước mỗi lần lưu
  tự sao lưu bản cũ vào "Dữ liệu\\Sao lưu").
- File xuất ra (đơn đặt hàng theo mẫu, sổ doanh thu) nằm trong "Xuất file".
"""
import json
import math
import os
import re
import shutil
import socket
import subprocess
import sys
import threading
import unicodedata
import urllib.request
import webbrowser
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parent
SOURCE_DIR = ROOT.parent  # thư mục chứa các file Excel gốc
sys.path.insert(0, str(ROOT / 'thu-vien'))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import licensing  # noqa: E402  (cấp phép offline 20 ngày — bản sao trong thu-vien, xem README)

DATA_DIR = ROOT / 'Dữ liệu'
DATA_FILE = DATA_DIR / 'Dữ liệu DT - ĐH - CN.xlsx'
BACKUP_DIR = DATA_DIR / 'Sao lưu'
TEMPLATE_DIR = DATA_DIR / 'Mẫu'
OUT_DIR = ROOT / 'Xuất file'
WEB_DIR = ROOT / 'giao-dien'
PORT = 8765

REGIONS = ['Miền Bắc', 'Miền Trung', 'Đông Nam Bộ', 'Tây Nam Bộ']
REGION_MAP = {'miền trung': 'Miền Trung', 'trung': 'Miền Trung', 'miền nam': 'Tây Nam Bộ', 'tây': 'Tây Nam Bộ',
              'tây nam bộ': 'Tây Nam Bộ', 'miền đông': 'Đông Nam Bộ', 'đông': 'Đông Nam Bộ',
              'đông nam bộ': 'Đông Nam Bộ', 'miền bắc': 'Miền Bắc', 'bắc': 'Miền Bắc'}
DEFAULT_SETTINGS = {
    'nhac_tu_ngay': 25,  # nhắc nợ cuối tháng: từ ngày này tới hết tháng
    'ck_km': 0.15,       # chiết khấu KM khác (Sọc Đỏ 520ml, Sọc Phấn 520ml) gợi ý sẵn khi lên đơn
    'cn_pct': 0.5,       # % công nợ gợi ý sẵn trên đơn (mẫu đơn để 50%)
}
# cài đặt dạng chữ: thông tin riêng của công ty, chỉ lưu trong file dữ liệu trên máy (không nằm trong mã nguồn)
TEXT_SETTINGS = {'chuyen_khoan': ''}  # số tài khoản, ngân hàng, tên đơn vị nhận tiền -> in trong tin nhắn nhắc nợ
SETTING_LABELS = {'nhac_tu_ngay': 'Nhắc nợ cuối tháng bắt đầu từ ngày', 'ck_km': 'Chiết khấu KM khác gợi ý (tỷ lệ)',
                  'chuyen_khoan': 'Thông tin chuyển khoản in trong tin nhắn nhắc nợ',
                  'cn_pct': 'Công nợ % gợi ý trên đơn (tỷ lệ)', 'nguon_mau_don': 'File mẫu đơn gốc',
                  'nguon_npp': 'File NPP gốc', 'ngay_cn_chot': 'Ngày lấy công nợ chốt'}
STATUS_VN = {'cho_giao': 'Chờ giao', 'da_giao': 'Đã giao', 'huy': 'Đã hủy', 'lich_su': 'Lịch sử (file cũ)'}
LINE_VN = {'ban': 'Hàng bán', 'tang_km': 'Hàng tặng KM', 'httm': 'Hỗ trợ thị trường'}
KM_VN = {'cung_loai': 'Tặng cùng loại', 'khac': 'Tặng SP khác', 'khong': 'Không', 'hang_tang': 'Dòng hàng tặng'}
LOCK = threading.Lock()
EXCEL_TIMEOUT = 300  # giây: Excel treo thì báo lỗi thay vì khóa công cụ mãi
MAX_BODY = 5 * 1024 * 1024  # giới hạn dữ liệu gửi lên máy chủ
APP_ID = 'dt-dh-cn'  # nhận ra chính công cụ này khi đang chạy ở cửa sổ khác
LICENSE = licensing.LicenseClient()


# =====================================================================
# tiện ích
# =====================================================================

def norm(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip()


def key(s):
    """Khóa so khớp tên NPP: bỏ dấu, chữ thường, bỏ phần '- Trắng/- Xanh' phía sau."""
    s = norm(s).lower()
    s = re.split(r'\s*-\s+|\s*-$', s)[0] if '-' in s else s
    s = unicodedata.normalize('NFD', s)
    return ''.join(ch for ch in s if unicodedata.category(ch) != 'Mn').replace('đ', 'd').strip()


def num(v):
    if v in (None, '', '-'):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', '').strip())
    except ValueError:
        return 0.0


def opt_num(v):
    return None if v in (None, '') else num(v)


def money(v):
    return int(round(v + 1e-9))  # tiền VNĐ làm tròn tới đồng


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and v:
        for fmt in ('%Y-%m-%d', '%d/%m/%Y'):
            try:
                return datetime.strptime(v[:10], fmt).date()
            except ValueError:
                pass
    return None


def iso(d):
    return d.isoformat() if d else None


def today():
    return date.today()


# =====================================================================
# đọc file Excel gốc (lần đầu)
# =====================================================================

def find_source_files():
    order_tpl = npp_file = None
    for f in sorted(SOURCE_DIR.glob('*.xlsx')):
        if f.name.startswith('~$'):
            continue
        try:
            names = load_workbook(f, read_only=True).sheetnames
        except Exception:
            continue
        if 'Đơn Đặt Hàng' in names and order_tpl is None:
            order_tpl = f
        if 'Công nợ chốt' in names and npp_file is None:
            npp_file = f
    return order_tpl, npp_file


def parse_km(text, quy_cach):
    """Chữ khuyến mãi -> quy tắc.
    '6 THÙNG TẶNG 5 CHAI' -> mỗi 6 thùng tặng 5 chai; '1 chai tặng 3 chai Long Ngư 500ml' -> tặng SP khác;
    '6 HỘP TẶNG 1 HỘP' (hộp = chai lẻ, 1 thùng = 6 hộp) -> mỗi 1 thùng tặng 1 hộp (đúng như đơn thực tế)."""
    t = norm(text).lower()
    m = re.match(r'(\d+)\s*chai\s*tặng\s*(\d+)\s*chai\s+(.+)', t)
    if m:
        return {'km_loai': 'khac', 'km_moi_thung_tang': int(m.group(2)) * quy_cach}
    m = re.match(r'(\d+)\s*(thùng|hộp|chai)\s*tặng\s*(\d+)', t)
    if m:
        n, unit, k = int(m.group(1)), m.group(2), int(m.group(3))
        if unit != 'thùng':  # mua tính theo đơn vị lẻ -> đổi ra thùng
            n = max(1, round(n / quy_cach))
        return {'km_loai': 'cung_loai', 'km_n': n, 'km_m': k}
    return {'km_loai': 'khong'}


def import_products(path):
    """Danh mục sản phẩm + bảng giá (giống nhau cho mọi vùng lúc đầu, sửa riêng từng vùng sau)."""
    wb = load_workbook(path, data_only=True)
    ws = wb['Đơn Đặt Hàng']
    products, price, last = [], {}, None
    for r in range(12, ws.max_row + 1):
        if 'tổng tiền hàng' in norm(ws.cell(r, 1).value).lower():
            break
        name = norm(ws.cell(r, 3).value)
        if not name:
            continue
        qc = int(num(ws.cell(r, 5).value)) or 1
        km_text = norm(ws.cell(r, 10).value)
        p = {'id': f'SP{len(products) + 1:02d}', 'nhan': norm(ws.cell(r, 1).value), 'ten': name,
             'dvt': norm(ws.cell(r, 4).value), 'quy_cach': qc, 'dong_mau': r, 'nhom_ck': 'thuong',
             'loai': 'ban', 'tang_cho': None, 'km_sp': None, 'dang_ban': True}
        rule = {'gia_chai': num(ws.cell(r, 6).value), 'km_text': km_text}
        if km_text.lower() == 'hàng tặng':
            p.update(loai='hang_tang', dang_ban=False, tang_cho=last['id'] if last else None, nhom_ck='km')
            rule['km_loai'] = 'hang_tang'
        else:
            rule.update(parse_km(km_text, qc))
        products.append(p)
        price[p['id']] = rule
        last = p
    for p in products:  # nối SP được tặng (VD: Sọc Đỏ 520ml tặng Long Ngư 500ml) + nhóm chiết khấu KM khác
        if price[p['id']].get('km_loai') == 'khac':
            p['nhom_ck'] = 'km'
            gift = next((g for g in products if g.get('tang_cho') == p['id']), None)
            p['km_sp'] = gift['id'] if gift else None
    wb.close()
    return products, {reg: {k: dict(v) for k, v in price.items()} for reg in REGIONS}


def header_map(ws, must='TÊN NPP', max_row=5):
    for r in range(1, max_row + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if must in [v.upper() for v in vals]:
            return r, {v.upper(): i + 1 for i, v in enumerate(vals) if v}
    return None, {}


def import_npps(path):
    wb = load_workbook(path, data_only=True)
    npps, by_key = [], {}
    for sheet in ('Bắc', 'Trung', 'Đông', 'Tây'):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        hr, h = header_map(ws)
        if not hr:
            continue
        for r in range(hr + 1, ws.max_row + 1):
            g = lambda col: ws.cell(r, h[col]).value if col in h else None  # noqa: E731
            name = norm(g('TÊN NPP'))
            if not name:
                continue
            note = ' '.join(norm(ws.cell(r, c).value) for c in range(h.get('CÔNG NỢ', 11) + 1, ws.max_column + 1)
                            if norm(ws.cell(r, c).value))
            cn_m = re.search(r'CN\s*(\d+)\s*%', note, re.I)
            n = {'id': f'NPP{len(npps) + 1:03d}', 'ten': name,
                 'khu_vuc': REGION_MAP.get(norm(g('KHU VỰC')).lower(), REGION_MAP.get(sheet.lower(), sheet)),
                 'gsm': norm(g('GSM')), 'asm': norm(g('ASM QUẢN LÍ')), 'tinh': norm(g('TỈNH')),
                 'dia_chi': norm(g('ĐỊA CHỈ')), 'sdt': norm(g('SĐT')).lstrip(';'),
                 'tinh_trang': norm(g('HÌNH THỨC HOẠT ĐỘNG')) or 'Đang hoạt động',
                 'ck': num(g('CHIẾT KHẤU')), 'cn_pct': int(cn_m.group(1)) / 100 if cn_m else None,
                 'cn_chot': 0.0, 'ghi_chu': note}
            npps.append(n)
            by_key[key(name)] = n
    ws = wb['Công nợ chốt']  # công nợ chốt = số dư công nợ hiện tại của từng NPP
    hr, h = header_map(ws, max_row=3)
    for r in range(hr + 1, ws.max_row + 1):
        name = norm(ws.cell(r, h['TÊN NPP']).value)
        if not name:
            continue
        n = by_key.get(key(name))
        if n is None:  # NPP chỉ có trong sheet công nợ chốt (VD Jerry)
            gsm = norm(ws.cell(r, h['GSM']).value)
            n = {'id': f'NPP{len(npps) + 1:03d}', 'ten': name, 'khu_vuc': 'Tây Nam Bộ' if 'Trà' in gsm else '',
                 'gsm': gsm, 'asm': '', 'tinh': norm(ws.cell(r, h['TỈNH']).value),
                 'dia_chi': norm(ws.cell(r, h['ĐỊA CHỈ']).value), 'sdt': norm(ws.cell(r, h['SĐT']).value),
                 'tinh_trang': 'Đang hoạt động', 'ck': num(ws.cell(r, h['CHIẾT KHẤU']).value), 'cn_pct': None,
                 'cn_chot': 0.0, 'ghi_chu': ''}
            npps.append(n)
            by_key[key(name)] = n
        n['cn_chot'] = num(ws.cell(r, h['CÔNG NỢ']).value)
    history = import_history(wb, by_key)
    for n in npps:  # NPP thiếu chiết khấu: lấy mức chiết khấu gần nhất trong các sheet tháng
        if not n['ck']:
            ck = [x['terms']['ck'] for x in history if x['npp_id'] == n['id'] and x['terms']['ck']]
            n['ck'] = ck[-1] if ck else 0.0
    wb.close()
    return npps, history


def import_history(wb, by_key):
    """Sheet "MM.YYYY": dòng có doanh thu sau chiết khấu -> đơn lịch sử (chỉ để xem doanh thu, không tính công nợ)."""
    history = []
    for sheet in wb.sheetnames:
        m = re.fullmatch(r'(\d{2})\.(\d{4})', sheet)
        if not m:
            continue
        month = date(int(m.group(2)), int(m.group(1)), 1)
        ws = wb[sheet]
        hr, _ = header_map(ws)
        if not hr:
            continue
        # sheet tháng có 2 cột "Chiết khấu" (tỷ lệ %, rồi số tiền) và 2 cột "Công nợ" (công nợ chốt, rồi công nợ của đơn)
        cols = {}
        for c in range(1, ws.max_column + 1):
            cols.setdefault(norm(ws.cell(hr, c).value).upper(), []).append(c)

        def g(r, name, nth=0):
            cs = cols.get(name, [])
            return ws.cell(r, cs[nth]).value if len(cs) > nth else None

        for r in range(hr + 1, ws.max_row + 1):
            sau = num(g(r, 'SAU CHIẾT KHẤU'))
            name = norm(g(r, 'TÊN NPP'))
            if not name or not sau:
                continue
            n = by_key.get(key(name))
            dat, giao = to_date(g(r, 'NGÀY ĐẶT HÀNG')), to_date(g(r, 'NGÀY GIAO HÀNG'))
            history.append(blank_order(
                id=f'LS-{sheet}-{r}', ma=f'LS {sheet} #{r}', loai='lich_su', trang_thai='lich_su',
                npp_id=n['id'] if n else None, npp_ten=name, thang=iso(month), ngay_dat=iso(dat or month),
                ngay_giao=iso(giao or dat or month), terms={'ck': num(g(r, 'CHIẾT KHẤU', 0)), 'ck_km': 0, 'cn_pct': 0},
                truoc_ck=num(g(r, 'TRƯỚC CHIẾT KHẤU')), tong_tien_hang=num(g(r, 'TRƯỚC CHIẾT KHẤU')),
                chiet_khau=num(g(r, 'CHIẾT KHẤU', 1)), sau_ck=sau, cong_no=num(g(r, 'CÔNG NỢ', 1)),
                da_tt=num(g(r, 'ĐÃ THANH TOÁN')), chua_tt=num(g(r, 'CHƯA THANH TOÁN')), ghi_chu=norm(g(r, 'GHI CHÚ'))))
    return history


def blank_order(**kw):
    o = {'id': '', 'ma': '', 'loai': 'don', 'trang_thai': 'cho_giao', 'npp_id': None, 'npp_ten': '', 'thang': None,
         'ngay_dat': None, 'ngay_giao': None, 'terms': {'ck': 0, 'ck_km': 0, 'cn_pct': 0}, 'tong_tien_hang': 0,
         'ck_thuong': 0, 'ck_km': 0, 'truoc_ck': 0, 'chiet_khau': 0, 'sau_ck': 0, 'cong_no': 0, 'thanh_toan': 0,
         'tong_chai_giao': 0, 'da_tt': 0, 'chua_tt': 0, 'ghi_chu': '', 'tao_luc': '', 'lines': [], 'gifts': [], 'httm': []}
    o.update(kw)
    return o


def initial_import():
    order_tpl, npp_file = find_source_files()
    if not order_tpl or not npp_file:
        raise RuntimeError('Không tìm thấy file mẫu đơn (sheet "Đơn Đặt Hàng") hoặc file NPP (sheet "Công nợ chốt") '
                           f'trong {SOURCE_DIR}')
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(order_tpl, TEMPLATE_DIR / 'Mau-don-dat-hang.xlsx')
    products, prices = import_products(order_tpl)
    npps, history = import_npps(npp_file)
    settings = dict(DEFAULT_SETTINGS, **TEXT_SETTINGS, nguon_mau_don=order_tpl.name, nguon_npp=npp_file.name, ngay_cn_chot=iso(today()))
    return {'settings': settings, 'npps': npps, 'products': products, 'prices': prices,
            'orders': history, 'payments': []}


# =====================================================================
# lưu trữ bằng Excel
# =====================================================================
# (khóa, tiêu đề cột, độ rộng) cho từng sheet của file dữ liệu
NPP_COLS = [('id', 'Mã NPP', 9), ('khu_vuc', 'Vùng', 13), ('gsm', 'GSM', 18), ('asm', 'ASM', 12), ('tinh', 'Tỉnh', 14),
            ('ten', 'Tên NPP', 24), ('dia_chi', 'Địa chỉ', 40), ('sdt', 'SĐT', 13), ('tinh_trang', 'Tình trạng', 16),
            ('ck', 'Chiết khấu thường', 10), ('cn_pct', 'Công nợ %', 10), ('cn_chot', 'Công nợ chốt đầu kỳ', 15),
            ('ghi_chu', 'Ghi chú', 40)]
SP_COLS = [('id', 'Mã SP', 8), ('nhan', 'Nhãn', 10), ('ten', 'Tên hàng', 55), ('dvt', 'ĐVT', 10), ('quy_cach', 'Quy cách', 9),
           ('dong_mau', 'Dòng trên mẫu đơn', 9), ('nhom_ck', 'Nhóm chiết khấu', 12), ('loai', 'Loại', 12),
           ('tang_cho', 'Là hàng tặng của', 10), ('km_sp', 'SP tặng kèm', 10), ('dang_ban', 'Đang bán', 9)]
PRICE_COLS = [('vung', 'Vùng', 13), ('sp', 'Mã SP', 8), ('ten', 'Tên hàng', 55), ('gia_chai', 'Giá chai', 11),
              ('km_text', 'Chương trình KM', 32), ('km_loai', 'Loại KM', 14), ('km_n', 'KM: mỗi (thùng)', 10),
              ('km_m', 'KM: tặng (chai/hũ)', 10), ('km_moi_thung_tang', 'KM: mỗi thùng tặng (chai SP khác)', 12)]
ORDER_COLS = [('ma', 'Mã đơn', 14), ('trang_thai', 'Trạng thái', 12), ('ngay_dat', 'Ngày đặt', 11),
              ('ngay_giao', 'Ngày giao', 11), ('npp_id', 'Mã NPP', 9), ('npp_ten', 'Tên NPP', 24), ('khu_vuc', 'Vùng', 13),
              ('gsm', 'GSM', 18), ('ck', 'CK thường', 9), ('ck_km_pct', 'CK KM khác', 9), ('cn_pct', 'Công nợ %', 9),
              ('tong_tien_hang', 'Tổng tiền hàng', 14), ('ck_thuong', 'Tiền CK thường', 13), ('ck_km', 'Tiền CK KM', 12),
              ('chiet_khau', 'Chiết khấu', 13), ('sau_ck', 'Sau chiết khấu', 14), ('cong_no', 'Công nợ', 13),
              ('thanh_toan', 'Tổng thanh toán', 14), ('tong_chai_giao', 'Tổng chai giao', 10),
              ('da_tt', 'Đã TT (file cũ)', 13), ('chua_tt', 'Chưa TT (file cũ)', 13), ('thang', 'Tháng (file cũ)', 11),
              ('ghi_chu', 'Ghi chú', 36), ('tao_luc', 'Tạo lúc', 17), ('id', 'ID', 16)]
LINE_COLS = [('ma', 'Mã đơn', 14), ('loai', 'Loại dòng', 16), ('sp', 'Mã SP', 8), ('ten', 'Tên hàng', 55), ('dvt', 'ĐVT', 10),
             ('quy_cach', 'Quy cách', 8), ('gia_chai', 'Giá chai', 11), ('gia_thung', 'Giá thùng', 12),
             ('thung', 'SL thùng', 9), ('thanh_tien', 'Thành tiền', 13), ('km_text', 'Chương trình KM', 30),
             ('tang', 'Tặng (chai/hũ)', 10), ('tang_tay', 'Tặng nhập tay', 9), ('tong_giao', 'Tổng giao', 10),
             ('gia_tay', 'Giá nhập tay', 9), ('dong_mau', 'Dòng mẫu', 8), ('cho', 'Tặng cho SP', 9), ('order_id', 'ID đơn', 16)]
PAY_COLS = [('id', 'Mã TT', 18), ('ngay', 'Ngày', 11), ('npp_id', 'Mã NPP', 9), ('npp_ten', 'Tên NPP', 24),
            ('so_tien', 'Số tiền', 14), ('tra_cho', 'Trừ vào', 16), ('ghi_chu', 'Ghi chú', 36)]
SHEETS = {'NPP': NPP_COLS, 'Sản phẩm': SP_COLS, 'Bảng giá': PRICE_COLS, 'Đơn hàng': ORDER_COLS,
          'Chi tiết đơn': LINE_COLS, 'Thanh toán': PAY_COLS}
DATE_KEYS = {'ngay_dat', 'ngay_giao', 'ngay', 'thang'}
PCT_KEYS = {'ck', 'cn_pct', 'ck_km_pct'}
MONEY_KEYS = {'cn_chot', 'gia_chai', 'gia_thung', 'thanh_tien', 'tong_tien_hang', 'ck_thuong', 'ck_km', 'chiet_khau',
              'sau_ck', 'cong_no', 'thanh_toan', 'da_tt', 'chua_tt', 'so_tien'}
REV = lambda d: {v: k for k, v in d.items()}  # noqa: E731
DATA_MTIME = [0.0]


def _write_sheet(wb, name, cols, rows):
    ws = wb.create_sheet(name)
    head_font, head_fill = Font(name='Arial', bold=True, color='FFFFFF', size=10), PatternFill('solid', fgColor='1F4E78')
    for i, (k, h, w) in enumerate(cols, 1):
        c = ws.cell(1, i, h)
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(rows, 2):
        for i, (k, _, _) in enumerate(cols, 1):
            v = row.get(k)
            if k in DATE_KEYS and v:
                v = to_date(v)
            c = ws.cell(r, i, v)
            c.font = Font(name='Arial', size=10)
            if k in DATE_KEYS:
                c.number_format = 'dd/mm/yyyy'
            elif k in PCT_KEYS and v is not None:
                c.number_format = '0.##%'
            elif k in MONEY_KEYS:
                c.number_format = '#,##0'
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{max(2, len(rows) + 1)}'


def save_db(db):
    """Ghi toàn bộ dữ liệu ra file Excel (sao lưu bản trước). File đang mở trong Excel -> PermissionError."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    npp = {n['id']: n for n in db['npps']}
    wb = Workbook()
    ws = wb.active
    ws.title = 'Cài đặt'
    ws.append(['Khóa', 'Giá trị', 'Ý nghĩa'])
    for k, v in db['settings'].items():
        ws.append([k, v, SETTING_LABELS.get(k, '')])
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width, ws.column_dimensions['C'].width = 16, 32, 40
    _write_sheet(wb, 'NPP', NPP_COLS, db['npps'])
    _write_sheet(wb, 'Sản phẩm', SP_COLS, [{**p, 'nhom_ck': 'KM khác' if p['nhom_ck'] == 'km' else 'Thường',
                                            'loai': 'Hàng tặng KM' if p['loai'] == 'hang_tang' else 'Hàng bán',
                                            'dang_ban': 'Có' if p['dang_ban'] else 'Không'} for p in db['products']])
    prices = []
    for reg in REGIONS:
        for p in db['products']:
            rule = db['prices'][reg][p['id']]
            prices.append({**rule, 'vung': reg, 'sp': p['id'], 'ten': p['ten'], 'km_loai': KM_VN[rule.get('km_loai', 'khong')]})
    _write_sheet(wb, 'Bảng giá', PRICE_COLS, prices)
    orders, lines = [], []
    for o in db['orders']:
        n = npp.get(o['npp_id'], {})
        orders.append({**o, 'trang_thai': STATUS_VN[o['trang_thai']], 'npp_ten': n.get('ten', o.get('npp_ten')),
                       'khu_vuc': n.get('khu_vuc', ''), 'gsm': n.get('gsm', ''), 'ck': o['terms'].get('ck'),
                       'ck_km_pct': o['terms'].get('ck_km'), 'cn_pct': o['terms'].get('cn_pct')})
        for kind, items in (('ban', o['lines']), ('tang_km', o['gifts']), ('httm', o['httm'])):
            for ln in items:
                lines.append({**ln, 'ma': o['ma'], 'order_id': o['id'], 'loai': LINE_VN[kind],
                              'tang_tay': 'Có' if ln.get('tang_tay') else '', 'gia_tay': 'Có' if ln.get('gia_tay') else ''})
    _write_sheet(wb, 'Đơn hàng', ORDER_COLS, orders)
    _write_sheet(wb, 'Chi tiết đơn', LINE_COLS, lines)
    _write_sheet(wb, 'Thanh toán', PAY_COLS, [{**p, 'npp_ten': npp.get(p['npp_id'], {}).get('ten', ''),
                                              'tra_cho': 'Công nợ chốt' if p['tra_cho'] == 'cn_chot' else 'Nợ theo đơn'}
                                             for p in db['payments']])
    tmp = DATA_DIR / '~dang-luu.xlsx'
    wb.save(tmp)
    backup = None
    if DATA_FILE.exists():
        stamp = f'Dữ liệu {datetime.now():%Y-%m-%d %Hh%Mm%Ss}'
        backup, i = BACKUP_DIR / f'{stamp}.xlsx', 1
        while backup.exists():  # lưu 2 lần trong cùng 1 giây: không ghi đè bản sao lưu trước
            backup, i = BACKUP_DIR / f'{stamp} ({i}).xlsx', i + 1
        shutil.copy2(DATA_FILE, backup)
    try:
        os.replace(tmp, DATA_FILE)
    except PermissionError:  # file đang mở trong Excel: bỏ bản sao lưu vừa tạo, KHÔNG xóa bản sao lưu cũ
        tmp.unlink(missing_ok=True)
        if backup:
            backup.unlink(missing_ok=True)
        raise
    for old in sorted(BACKUP_DIR.glob('Dữ liệu *.xlsx'))[:-60]:  # lưu xong mới dọn, giữ 60 bản gần nhất
        old.unlink()
    DATA_MTIME[0] = DATA_FILE.stat().st_mtime


def _read_sheet(wb, name, cols):
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    head = [norm(c.value) for c in ws[1]]
    idx = {k: head.index(h) for k, h, _ in cols if h in head}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, '') for v in row):
            continue
        d = {}
        for k, i in idx.items():
            v = row[i] if i < len(row) else None
            if k in DATE_KEYS:
                v = iso(to_date(v))
            elif isinstance(v, str):
                v = v.strip()
            d[k] = v
        out.append(d)
    return out


def load_db():
    if not DATA_FILE.exists():
        db = initial_import()
        save_db(db)
        return db
    wb = load_workbook(DATA_FILE, data_only=True)
    settings = dict(DEFAULT_SETTINGS, **TEXT_SETTINGS)
    for row in wb['Cài đặt'].iter_rows(min_row=2, values_only=True):
        if row[0]:
            settings[row[0]] = row[1]
    for k in DEFAULT_SETTINGS:
        settings[k] = num(settings[k])
    for k in TEXT_SETTINGS:
        settings[k] = norm(settings[k])
    npps = []
    for n in _read_sheet(wb, 'NPP', NPP_COLS):
        n.update(ck=num(n.get('ck')), cn_pct=opt_num(n.get('cn_pct')), cn_chot=num(n.get('cn_chot')))
        for k in ('khu_vuc', 'gsm', 'asm', 'tinh', 'ten', 'dia_chi', 'sdt', 'tinh_trang', 'ghi_chu'):
            n[k] = norm(n.get(k))
        npps.append(n)
    products = []
    for p in _read_sheet(wb, 'Sản phẩm', SP_COLS):
        p.update(quy_cach=int(num(p['quy_cach'])) or 1, dong_mau=int(num(p.get('dong_mau'))) or None,
                 nhom_ck='km' if p.get('nhom_ck') == 'KM khác' else 'thuong',
                 loai='hang_tang' if p.get('loai') == 'Hàng tặng KM' else 'ban', dang_ban=p.get('dang_ban') != 'Không',
                 tang_cho=p.get('tang_cho') or None, km_sp=p.get('km_sp') or None, nhan=norm(p.get('nhan')), dvt=norm(p.get('dvt')))
        products.append(p)
    prices = {reg: {} for reg in REGIONS}
    km_rev = REV(KM_VN)
    for r in _read_sheet(wb, 'Bảng giá', PRICE_COLS):
        if r.get('vung') not in prices:
            continue
        prices[r['vung']][r['sp']] = {'gia_chai': num(r.get('gia_chai')), 'km_text': norm(r.get('km_text')),
                                     'km_loai': km_rev.get(r.get('km_loai'), 'khong'), 'km_n': opt_num(r.get('km_n')),
                                     'km_m': opt_num(r.get('km_m')), 'km_moi_thung_tang': opt_num(r.get('km_moi_thung_tang'))}
    st_rev, line_rev = REV(STATUS_VN), REV(LINE_VN)
    orders, by_id = [], {}
    for o in _read_sheet(wb, 'Đơn hàng', ORDER_COLS):
        status = st_rev.get(o.get('trang_thai'), 'cho_giao')
        oo = blank_order(**{k: o.get(k) for k in ('id', 'ma', 'ngay_dat', 'ngay_giao', 'npp_id', 'npp_ten', 'thang', 'tao_luc')},
                         loai='lich_su' if status == 'lich_su' else 'don', trang_thai=status, ghi_chu=norm(o.get('ghi_chu')),
                         terms={'ck': num(o.get('ck')), 'ck_km': num(o.get('ck_km_pct')), 'cn_pct': num(o.get('cn_pct'))},
                         **{k: num(o.get(k)) for k in ('tong_tien_hang', 'ck_thuong', 'ck_km', 'chiet_khau', 'sau_ck',
                                                       'cong_no', 'thanh_toan', 'tong_chai_giao', 'da_tt', 'chua_tt')})
        oo['truoc_ck'] = oo['tong_tien_hang']
        oo['id'] = oo['id'] or oo['ma']
        orders.append(oo)
        by_id[oo['id']] = oo
    for ln in _read_sheet(wb, 'Chi tiết đơn', LINE_COLS):
        o = by_id.get(ln.get('order_id'))
        if not o:
            continue
        kind = line_rev.get(ln.get('loai'), 'ban')
        item = {k: ln.get(k) for k in ('sp', 'ten', 'dvt', 'km_text', 'cho')}
        item.update({k: num(ln.get(k)) for k in ('quy_cach', 'gia_chai', 'gia_thung', 'thung', 'thanh_tien', 'tang', 'tong_giao')},
                    dong_mau=int(num(ln.get('dong_mau'))) or None, tang_tay=ln.get('tang_tay') == 'Có',
                    gia_tay=ln.get('gia_tay') == 'Có')
        o[{'ban': 'lines', 'tang_km': 'gifts', 'httm': 'httm'}[kind]].append(item)
    payments = []
    for p in _read_sheet(wb, 'Thanh toán', PAY_COLS):
        payments.append({'id': p['id'], 'ngay': p['ngay'], 'npp_id': p['npp_id'], 'so_tien': num(p['so_tien']),
                         'tra_cho': 'cn_chot' if p.get('tra_cho') == 'Công nợ chốt' else 'don', 'ghi_chu': norm(p.get('ghi_chu'))})
    wb.close()
    DATA_MTIME[0] = DATA_FILE.stat().st_mtime
    return {'settings': settings, 'npps': npps, 'products': products, 'prices': prices, 'orders': orders, 'payments': payments}


# =====================================================================
# tính đơn hàng
# =====================================================================

def npp_terms(db, npp):
    s = db['settings']
    return {'ck': npp.get('ck') or 0.0, 'ck_km': s['ck_km'],
            'cn_pct': npp['cn_pct'] if npp.get('cn_pct') is not None else s['cn_pct']}


def auto_gift(rule, q):
    """Số tặng tự tính: cùng loại -> mỗi đủ N thùng tặng M; SP khác -> mỗi thùng tặng X chai SP khác."""
    if rule.get('km_loai') == 'cung_loai' and rule.get('km_n'):
        return math.floor(q / rule['km_n'] + 1e-9) * (rule.get('km_m') or 0)
    if rule.get('km_loai') == 'khac':
        return q * (rule.get('km_moi_thung_tang') or 0)
    return 0.0


def quote(db, body):
    """Tính đơn đúng như mẫu "Đơn Đặt Hàng": giá theo vùng của NPP, hàng tặng KM (sửa tay được), chiết khấu nhập tay,
    hàng hỗ trợ thị trường (không tính tiền)."""
    npp = next(n for n in db['npps'] if n['id'] == body['npp_id'])
    region = npp['khu_vuc'] if npp['khu_vuc'] in REGIONS else REGIONS[0]
    price = db['prices'][region]
    terms = npp_terms(db, npp)
    for k in ('ck', 'ck_km', 'cn_pct'):
        if body.get(k) not in (None, ''):
            terms[k] = num(body[k])
    prods = {p['id']: p for p in db['products']}
    out, gifts, httm = [], [], []
    tien_thuong = tien_km = 0.0
    for ln in body.get('lines', []):
        p, q = prods.get(ln.get('sp')), num(ln.get('thung'))
        if not p or q <= 0:
            continue
        rule = price.get(p['id'], {})
        gia_tay = ln.get('gia_chai') not in (None, '') and num(ln['gia_chai']) > 0
        gia_chai = num(ln['gia_chai']) if gia_tay else num(rule.get('gia_chai'))
        gia_thung = gia_chai * p['quy_cach']
        thanh_tien = q * gia_thung
        tang_tay = ln.get('tang') not in (None, '')
        tang = num(ln['tang']) if tang_tay else auto_gift(rule, q)
        if rule.get('km_loai') == 'khac' and p.get('km_sp') in prods:  # tặng SP khác: thành dòng hàng tặng riêng
            g = prods[p['km_sp']]
            gifts.append({'sp': g['id'], 'ten': g['ten'], 'dvt': g['dvt'], 'quy_cach': g['quy_cach'],
                          'gia_chai': num(price.get(g['id'], {}).get('gia_chai')), 'gia_thung': 0, 'thung': tang / g['quy_cach'],
                          'thanh_tien': 0, 'km_text': 'Hàng tặng', 'tang': tang, 'tang_tay': tang_tay, 'tong_giao': tang,
                          'cho': p['id'], 'dong_mau': g.get('dong_mau')})
            own_gift = 0.0
        else:
            own_gift = tang
        if p['nhom_ck'] == 'km':
            tien_km += thanh_tien
        else:
            tien_thuong += thanh_tien
        out.append({'sp': p['id'], 'ten': p['ten'], 'nhan': p['nhan'], 'dvt': p['dvt'], 'quy_cach': p['quy_cach'],
                    'gia_chai': gia_chai, 'gia_tay': gia_tay, 'gia_thung': gia_thung, 'thung': q, 'thanh_tien': thanh_tien,
                    'km_text': rule.get('km_text', ''), 'tang': own_gift, 'tang_tay': tang_tay and own_gift > 0,
                    'tang_tu_dong': auto_gift(rule, q), 'tong_giao': q * p['quy_cach'] + own_gift, 'dong_mau': p.get('dong_mau')})
    for ln in body.get('httm', []):
        p, q = prods.get(ln.get('sp')), num(ln.get('thung'))
        if not p or q <= 0:
            continue
        gia_chai = num(price.get(p['id'], {}).get('gia_chai'))
        httm.append({'sp': p['id'], 'ten': p['ten'], 'dvt': p['dvt'], 'quy_cach': p['quy_cach'], 'gia_chai': gia_chai,
                     'gia_thung': gia_chai * p['quy_cach'], 'thung': q, 'thanh_tien': 0, 'km_text': 'Hỗ trợ thị trường',
                     'tang': 0, 'tong_giao': q * p['quy_cach'], 'dong_mau': None})
    tong = tien_thuong + tien_km
    ck_thuong, ck_khac = tien_thuong * terms['ck'], tien_km * terms['ck_km']
    sau = tong - ck_thuong - ck_khac
    return {'region': region, 'lines': out, 'gifts': gifts, 'httm': httm, 'terms': terms,
            'tong_tien_hang': money(tong), 'ck_thuong': money(ck_thuong), 'ck_km': money(ck_khac),
            'truoc_ck': money(tong), 'chiet_khau': money(ck_thuong) + money(ck_khac),
            'sau_ck': money(tong) - money(ck_thuong) - money(ck_khac), 'cong_no': money(sau * terms['cn_pct']),
            'thanh_toan': money(tong) - money(ck_thuong) - money(ck_khac) - money(sau * terms['cn_pct']),
            'tong_chai_giao': sum(x['tong_giao'] for x in out + gifts + httm),
            'chai_httm': sum(x['tong_giao'] for x in httm)}


def next_code(db, d):
    prefix = f'DH{d:%y%m%d}-'
    used = [int(o['ma'][len(prefix):]) for o in db['orders'] if str(o.get('ma', '')).startswith(prefix)
            and o['ma'][len(prefix):].isdigit()]
    return f'{prefix}{max(used, default=0) + 1:02d}'


# =====================================================================
# công nợ: nhắc mỗi cuối tháng tới khi nợ = 0
# =====================================================================

def debt_state(db, ref=None):
    """Tiền NPP trả (trừ vào nợ theo đơn) trừ dần vào đơn giao sớm nhất; còn dư thì trừ tiếp công nợ chốt."""
    ref = ref or today()
    month_start = ref.replace(day=1)
    per_npp = {n['id']: {'npp_id': n['id'], 'cn_chot_dau': n.get('cn_chot', 0.0), 'cn_chot_da_tra': 0.0,
                         'phai_thu': 0.0, 'da_tra': 0.0, 'no_thang_truoc': 0.0, 'no_trong_thang': 0.0,
                         'no_tu_thang': None} for n in db['npps']}
    pool = {}
    for p in db['payments']:
        st = per_npp.get(p['npp_id'])
        if st is None:
            continue
        if p['tra_cho'] == 'cn_chot':
            st['cn_chot_da_tra'] += p['so_tien']
        else:
            pool[p['npp_id']] = pool.get(p['npp_id'], 0.0) + p['so_tien']
    items = []
    for o in sorted([o for o in db['orders'] if o['trang_thai'] == 'da_giao'], key=lambda o: (o['ngay_giao'], o['ma'])):
        left = pool.get(o['npp_id'], 0.0)
        pay = min(left, o['sau_ck'])
        pool[o['npp_id']] = left - pay
        it = {'order_id': o['id'], 'ma': o['ma'], 'npp_id': o['npp_id'], 'ngay_giao': o['ngay_giao'], 'so_tien': o['sau_ck'],
              'da_tra': pay, 'con_lai': o['sau_ck'] - pay, 'thang': o['ngay_giao'][:7],
              'tu_thang_truoc': to_date(o['ngay_giao']) < month_start}
        items.append(it)
        st = per_npp.get(o['npp_id'])
        if st is None:
            continue
        st['phai_thu'] += it['so_tien']
        st['da_tra'] += pay
        if it['con_lai'] > 0.5:
            st['no_thang_truoc' if it['tu_thang_truoc'] else 'no_trong_thang'] += it['con_lai']
            st['no_tu_thang'] = min(st['no_tu_thang'] or it['thang'], it['thang'])
    for nid, st in per_npp.items():
        left = pool.get(nid, 0.0)
        st['con_no_don'] = st['phai_thu'] - st['da_tra']
        chot = max(0.0, st['cn_chot_dau'] - st['cn_chot_da_tra'])
        st['cn_chot_tu_du'] = min(left, chot)  # tiền trả dư sau khi hết nợ đơn -> trừ tiếp công nợ chốt
        st['cn_chot_con'] = st['cn_chot_dau'] - st['cn_chot_da_tra'] - st['cn_chot_tu_du']
        st['du'] = left - st['cn_chot_tu_du']
        st['tong_no'] = st['cn_chot_con'] + st['con_no_don']
    return items, per_npp


def reminder_period(db, ref=None):
    """Kỳ nhắc nợ: từ ngày X tới hết tháng."""
    ref = ref or today()
    start = int(db['settings']['nhac_tu_ngay'] or 25)
    last = (date(ref.year + ref.month // 12, ref.month % 12 + 1, 1) - date.resolution).day
    start = min(start, last)
    return {'thang': f'{ref:%m/%Y}', 'tu': iso(ref.replace(day=start)), 'den': iso(ref.replace(day=last)),
            'dang_nhac': ref.day >= start}


# =====================================================================
# xuất file
# =====================================================================

def run_excel_job(job):
    OUT_DIR.mkdir(exist_ok=True)
    job_file = OUT_DIR / '_job.json'
    job_file.write_text(json.dumps(job, ensure_ascii=False), encoding='utf-8')
    try:
        res = subprocess.run(['powershell', '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass',
                              '-File', str(ROOT / 'dien_mau_excel.ps1'), str(job_file)], capture_output=True,
                             timeout=EXCEL_TIMEOUT)
    except subprocess.TimeoutExpired:
        raise RuntimeError(f'Excel không phản hồi sau {EXCEL_TIMEOUT // 60} phút (có thể Excel đang hiện hộp thoại). '
                           'Mở Excel kiểm tra, đóng hộp thoại rồi làm lại.') from None
    finally:
        job_file.unlink(missing_ok=True)
    if res.returncode != 0:
        err = res.stderr.decode('utf-8', 'replace')[-400:] or res.stdout.decode('utf-8', 'replace')[-400:]
        raise RuntimeError('Excel không điền được mẫu đơn: ' + err)


def export_order(db, order):
    """Điền mẫu "Đơn Đặt Hàng" gốc bằng Excel (giữ logo, định dạng, công thức), lọc ẩn dòng không đặt, xuất PDF."""
    npp = next(n for n in db['npps'] if n['id'] == order['npp_id'])
    folder = OUT_DIR / 'Đơn đặt hàng'
    folder.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r'[\\/:*?"<>|]', '', npp['ten'])  # bỏ ký tự không được dùng trong tên file
    out = folder / f'{order["ma"]} - {safe_name}.xlsx'
    d = to_date(order['ngay_dat'])
    DDH, PXK = 'Đơn Đặt Hàng', 'Phiếu Xuất Kho'
    cells = []
    for sheet in (DDH, PXK):
        cells += [[sheet, 'D8', npp['ten']], [sheet, 'D9', npp['dia_chi']], [sheet, 'D10', npp['sdt']]]
    cells.append([DDH, 'C5', f'ĐƠN ĐẶT HÀNG          Số: {order["ma"]}          Ngày {d:%d/%m/%Y}'])
    ordered = {ln['dong_mau']: ln for ln in order['lines']}
    gifts = {g['dong_mau']: g for g in order['gifts']}
    for p in db['products']:
        r = p.get('dong_mau')
        if not r:
            continue
        if p['loai'] == 'hang_tang':  # dòng hàng tặng KM (Long Ngư 500ml tặng kèm Sọc Đỏ / Sọc Phấn 520)
            g = gifts.get(r)
            cells += [[DDH, f'H{r}', g['thung'] if g else 0], [DDH, f'K{r}', g['tang'] if g else 0]]
            continue
        ln = ordered.get(r)
        cells.append([DDH, f'H{r}', ln['thung'] if ln else 0])
        if ln:  # giá & KM theo vùng / sửa tay, số tặng do công cụ tính
            cells += [[DDH, f'F{r}', ln['gia_chai']], [DDH, f'J{r}', ln['km_text']], [DDH, f'K{r}', ln['tang']]]
    t = order['terms']
    cells += [[DDH, 'H49', t['ck']], [DDH, 'H50', t['ck_km']], [DDH, 'H52', t['cn_pct']],
              [DDH, 'A55', f'Hỗ Trợ Thương Mại{" " * 52}Phụ trách khu vực: {npp.get("gsm") or "…"}']]
    # hỗ trợ thị trường: mẫu có sẵn 2 dòng (57, 58) -> cần nhiều hơn thì chèn thêm dòng ở cả 2 sheet
    h = order['httm']
    extra = max(0, len(h) - 2)
    inserts = [[DDH, 59, extra], [PXK, 59, extra]] if extra else []
    for i in range(max(2, len(h))):
        r = 57 + i
        ln = h[i] if i < len(h) else None
        if ln:
            cells += [[DDH, f'B{r}', i + 1], [DDH, f'C{r}', ln['ten']], [DDH, f'D{r}', ln['dvt']], [DDH, f'E{r}', ln['quy_cach']],
                      [DDH, f'F{r}', ln['gia_chai']], [DDH, f'G{r}', ln['gia_thung']], [DDH, f'H{r}', ln['thung']],
                      [DDH, f'I{r}', 0], [DDH, f'J{r}', 'Hỗ trợ thị trường'], [DDH, f'L{r}', ln['tong_giao']],
                      [PXK, f'B{r}', i + 1], [PXK, f'C{r}', ln['ten']], [PXK, f'D{r}', ln['dvt']], [PXK, f'E{r}', ln['quy_cach']],
                      [PXK, f'F{r}', ln['thung']], [PXK, f'H{r}', 'Hỗ trợ thị trường'], [PXK, f'J{r}', ln['tong_giao']],
                      [PXK, f'K{r}', int(ln['tong_giao'] // ln['quy_cach'])], [PXK, f'L{r}', ln['tong_giao'] % ln['quy_cach']]]
        else:  # xóa dòng ví dụ có sẵn trong mẫu
            cells += [[DDH, f'{c}{r}', None] for c in 'BCDEFGHIJKL'] + [[PXK, f'{c}{r}', None] for c in 'BCDEFGHIJKL']
    job = {'template': str(TEMPLATE_DIR / 'Mau-don-dat-hang.xlsx'), 'output': str(out), 'pdf': str(out.with_suffix('.pdf')),
           'pdf_sheet': DDH, 'inserts': inserts, 'cells': cells,
           # ẩn các mặt hàng không đặt (cột "Số lượng đặt" = 0) như đơn thực tế, thay bộ lọc cũ còn sót trong mẫu
           'hide_zero': [[DDH, 12, 47, 'H', 'L'], [PXK, 12, 47, 'F', 'P']]}
    run_excel_job(job)
    wb = load_workbook(out, data_only=True)  # đọc lại tổng Excel tính để đối chiếu với công cụ
    ws = wb[DDH]
    check = {'tong_tien_hang': num(ws['I48'].value), 'sau_ck': num(ws['I51'].value), 'cong_no': num(ws['I52'].value)}
    wb.close()
    return {'file': str(out), 'pdf': str(out.with_suffix('.pdf')),
            'lech': {k: v for k, v in check.items() if abs(v - order[k]) > 1}}


THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill('solid', fgColor='1F4E78')
NUMF = '#,##0;-#,##0;"-"'


def _font(bold=False, color='262626', size=10):
    return Font(name='Arial', size=size, bold=bold, color=color)


def _head(ws, cols, row=1):
    for i, (h, w) in enumerate(cols, 1):
        c = ws.cell(row, i, h)
        c.font, c.fill, c.border = _font(True, 'FFFFFF'), HEAD_FILL, BOX
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32


def export_ledger(db):
    """Sổ theo dõi: sheet Nhắc nợ cuối tháng, Công nợ, và mỗi tháng 1 sheet đúng cột như file "Doanh thu, công nợ NPP"."""
    items, per_npp = debt_state(db)
    npps = {n['id']: n for n in db['npps']}
    paid = {}
    for it in items:
        paid[it['order_id']] = it['da_tra']
    by_month = {}
    for o in db['orders']:
        if o['trang_thai'] in ('da_giao', 'lich_su'):
            by_month.setdefault((o['thang'] if o['trang_thai'] == 'lich_su' else o['ngay_giao'])[:7], []).append(o)
    wb = Workbook()
    wb.remove(wb.active)
    reg_order = lambda nid: REGIONS.index(npps[nid]['khu_vuc']) if nid in npps and npps[nid]['khu_vuc'] in REGIONS else 9  # noqa: E731
    # --- nhắc nợ cuối tháng ---
    per = reminder_period(db)
    ws = wb.create_sheet(f'Nhắc nợ {per["thang"].replace("/", ".")}')
    _head(ws, [('GSM', 18), ('Vùng', 13), ('TÊN NPP', 24), ('Tỉnh', 14), ('SĐT', 13), ('Công nợ chốt còn', 15),
               ('Nợ đơn các tháng trước', 15), ('Nợ đơn trong tháng', 15), ('TỔNG CÒN NỢ', 15), ('Nợ đơn từ tháng', 11)])
    r = 2
    for n in sorted(db['npps'], key=lambda n: (n['gsm'], -per_npp[n['id']]['tong_no'])):
        st = per_npp[n['id']]
        if st['tong_no'] <= 0.5:
            continue
        vals = [n['gsm'], n['khu_vuc'], n['ten'], n['tinh'], n['sdt'], st['cn_chot_con'], st['no_thang_truoc'],
                st['no_trong_thang'], f'=F{r}+G{r}+H{r}', st['no_tu_thang'] and f'{st["no_tu_thang"][5:]}/{st["no_tu_thang"][:4]}']
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.font, c.border, c.number_format = _font(i == 9), BOX, NUMF
        r += 1
    for i in range(6, 10):
        L = get_column_letter(i)
        c = ws.cell(r, i, f'=SUM({L}2:{L}{r - 1})')
        c.font, c.number_format = _font(True, 'C00000'), NUMF
    ws.cell(r, 3, 'TỔNG CỘNG').font = _font(True)
    ws.freeze_panes = 'D2'
    # --- sheet tháng ---
    cols = [('STT', 5), ('Khu vực', 13), ('GSM', 18), ('TỈNH', 14), ('TÊN NPP', 24), ('ĐỊA CHỈ', 30), ('SĐT', 13),
            ('Hình thức hoạt động', 15), ('Chiết khấu', 9), ('CÔNG NỢ chốt', 13), ('Ngày đặt hàng', 12),
            ('Ngày giao hàng', 12), ('Trước chiết khấu', 14), ('Chiết khấu', 13), ('Sau chiết khấu', 14),
            ('Công nợ', 13), ('Đã thanh toán', 14), ('Chưa thanh toán', 14), ('Hàng HTTM (chai)', 10), ('Ghi chú', 30)]
    for m in sorted(by_month, reverse=True):
        ws = wb.create_sheet(f'{m[5:7]}.{m[:4]}')
        _head(ws, cols, row=2)
        orders = sorted(by_month[m], key=lambda o: (reg_order(o.get('npp_id')), o['ngay_giao'] or ''))
        for k, o in enumerate(orders, 1):
            r, n = k + 2, npps.get(o.get('npp_id'), {})
            if o['trang_thai'] == 'lich_su':
                da, chua, note = o['da_tt'], o['chua_tt'], ('[Nhập từ file cũ] ' + o['ghi_chu']).strip()
            else:
                da = paid.get(o['id'], 0.0)
                chua, note = o['sau_ck'] - da, f'{o["ma"]} {o["ghi_chu"]}'.strip()
            vals = [k, n.get('khu_vuc', ''), n.get('gsm', ''), n.get('tinh', ''), n.get('ten', o.get('npp_ten')),
                    n.get('dia_chi', ''), n.get('sdt', ''), n.get('tinh_trang', ''), o['terms'].get('ck'), n.get('cn_chot', 0),
                    to_date(o['ngay_dat']), to_date(o['ngay_giao']), o['truoc_ck'], o['chiet_khau'], o['sau_ck'],
                    o['cong_no'], da, chua, sum(x['tong_giao'] for x in o['httm']) or None, note]
            for i, v in enumerate(vals, 1):
                c = ws.cell(r, i, v)
                c.font, c.border = _font(), BOX
                c.number_format = '0.##%' if i == 9 else 'dd/mm/yyyy' if i in (11, 12) else NUMF
        last = max(len(orders) + 2, 3)
        for i in range(13, 20):  # dòng 1 = tổng, giống file gốc
            L = get_column_letter(i)
            c = ws.cell(1, i, f'=SUM({L}3:{L}{last})')
            c.font, c.number_format = _font(True, 'C00000'), NUMF
        ws.freeze_panes = 'F3'
        ws.auto_filter.ref = f'A2:T{last}'
    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / f'Sổ theo dõi doanh thu - công nợ NPP (cập nhật {datetime.now():%d.%m.%Y %Hh%M}).xlsx'
    wb.save(out)
    return str(out)


# =====================================================================
# API
# =====================================================================

def public_state(db):
    items, per_npp = debt_state(db)
    return {'today': iso(today()), 'regions': REGIONS, 'settings': db['settings'], 'npps': db['npps'],
            'products': db['products'], 'prices': db['prices'], 'orders': db['orders'], 'payments': db['payments'],
            'debt_items': items, 'debt_npp': per_npp, 'nhac': reminder_period(db), 'data_file': str(DATA_FILE)}


def license_info():
    return dict(LICENSE.check_license().to_dict(), policy=licensing.POLICY_TEXT, data_file=str(DATA_FILE),
                app=APP_ID)


def find_order(db, oid):
    o = next((o for o in db['orders'] if o['id'] == oid), None)
    if not o:
        raise ValueError('Không tìm thấy đơn.')
    return o


def handle(db, path, body):
    s = db['settings']
    if path == '/api/quote':
        return quote(db, body)
    if path == '/api/order/save':
        q = quote(db, body)
        if not q['lines'] and not q['httm']:
            raise ValueError('Đơn chưa có mặt hàng nào.')
        d = to_date(body.get('ngay_dat')) or today()
        if body.get('id'):
            o = find_order(db, body['id'])
            if o['trang_thai'] != 'cho_giao':
                raise ValueError('Chỉ sửa được đơn đang chờ giao.')
        else:
            ma = next_code(db, d)
            o = blank_order(id=ma, ma=ma, tao_luc=datetime.now().isoformat(timespec='seconds'))
            db['orders'].append(o)
        o.update({'npp_id': body['npp_id'], 'ngay_dat': iso(d), 'ngay_giao': None, 'ghi_chu': norm(body.get('ghi_chu')),
                  **{k: q[k] for k in ('lines', 'gifts', 'httm', 'terms', 'tong_tien_hang', 'ck_thuong', 'ck_km', 'truoc_ck',
                                       'chiet_khau', 'sau_ck', 'cong_no', 'thanh_toan', 'tong_chai_giao')}})
        save_db(db)
        return {'ok': True, 'order': o}
    if path == '/api/order/deliver':
        o = find_order(db, body['id'])
        if o['trang_thai'] != 'cho_giao':
            raise ValueError('Đơn này không ở trạng thái chờ giao.')
        o['trang_thai'], o['ngay_giao'] = 'da_giao', iso(to_date(body.get('ngay_giao')) or today())
        save_db(db)
        return {'ok': True}
    if path == '/api/order/undeliver':
        o = find_order(db, body['id'])
        o['trang_thai'], o['ngay_giao'] = 'cho_giao', None
        save_db(db)
        return {'ok': True}
    if path == '/api/order/cancel':
        o = find_order(db, body['id'])
        if o['trang_thai'] == 'da_giao':
            raise ValueError('Đơn đã giao: bấm "Hoàn tác giao" trước khi hủy.')
        o['trang_thai'] = 'huy'
        save_db(db)
        return {'ok': True}
    if path == '/api/order/export':
        return export_order(db, find_order(db, body['id']))
    if path == '/api/payment/add':
        amount = num(body.get('so_tien'))
        if amount <= 0:
            raise ValueError('Số tiền phải lớn hơn 0.')
        if not any(n['id'] == body.get('npp_id') for n in db['npps']):
            raise ValueError('Chọn NPP.')
        db['payments'].append({'id': f'TT{datetime.now():%Y%m%d%H%M%S%f}'[:20], 'npp_id': body['npp_id'],
                               'ngay': iso(to_date(body.get('ngay')) or today()), 'so_tien': amount,
                               'tra_cho': 'cn_chot' if body.get('tra_cho') == 'cn_chot' else 'don',
                               'ghi_chu': norm(body.get('ghi_chu'))})
        save_db(db)
        return {'ok': True}
    if path == '/api/payment/delete':
        db['payments'] = [p for p in db['payments'] if p['id'] != body['id']]
        save_db(db)
        return {'ok': True}
    if path == '/api/npp/save':
        data = body['npp']
        data.update(ck=num(data.get('ck')), cn_chot=num(data.get('cn_chot')), cn_pct=opt_num(data.get('cn_pct')))
        for k in ('khu_vuc', 'gsm', 'asm', 'tinh', 'ten', 'dia_chi', 'sdt', 'tinh_trang', 'ghi_chu'):
            data[k] = norm(data.get(k))
        if data.get('id'):
            next(n for n in db['npps'] if n['id'] == data['id']).update(data)
        else:
            nums = [int(n['id'][3:]) for n in db['npps'] if n['id'][3:].isdigit()]
            data['id'] = f'NPP{max(nums, default=0) + 1:03d}'
            db['npps'].append(data)
        save_db(db)
        return {'ok': True}
    if path == '/api/product/save':
        data = body['product']
        p = next(p for p in db['products'] if p['id'] == data['id'])
        p.update(ten=norm(data.get('ten')) or p['ten'], dvt=norm(data.get('dvt')) or p['dvt'],
                 quy_cach=int(num(data.get('quy_cach'))) or p['quy_cach'], dang_ban=bool(data.get('dang_ban')))
        save_db(db)
        return {'ok': True}
    if path == '/api/prices/save':
        reg = body['region']
        if reg not in REGIONS:
            raise ValueError('Vùng không hợp lệ.')
        for r in body['rows']:
            rule = db['prices'][reg].setdefault(r['sp'], {})
            rule.update(gia_chai=num(r.get('gia_chai')), km_text=norm(r.get('km_text')))
            if rule.get('km_loai') == 'hang_tang':
                continue
            if rule.get('km_loai') == 'khac':
                rule['km_moi_thung_tang'] = num(r.get('km_moi_thung_tang'))
            else:
                n_, m_ = num(r.get('km_n')), num(r.get('km_m'))
                rule.update(km_loai='cung_loai', km_n=n_, km_m=m_) if n_ > 0 and m_ > 0 else \
                    rule.update(km_loai='khong', km_n=None, km_m=None)
        save_db(db)
        return {'ok': True}
    if path == '/api/prices/copy':
        if body['from'] not in REGIONS or body['to'] not in REGIONS:
            raise ValueError('Vùng không hợp lệ.')
        db['prices'][body['to']] = json.loads(json.dumps(db['prices'][body['from']]))
        save_db(db)
        return {'ok': True}
    if path == '/api/settings/save':
        for k in DEFAULT_SETTINGS:
            if k in body:
                s[k] = num(body[k])
        for k in TEXT_SETTINGS:
            if k in body:
                s[k] = norm(body[k])
        save_db(db)
        return {'ok': True}
    if path == '/api/export/ledger':
        return {'file': export_ledger(db)}
    if path == '/api/open':
        target = Path(body['path']).resolve()
        if target.exists() and (OUT_DIR.resolve() in target.parents or target == DATA_FILE.resolve()):
            os.startfile(target)
        return {'ok': True}
    raise ValueError('Không rõ yêu cầu: ' + path)


class Handler(BaseHTTPRequestHandler):
    db = None

    def log_message(self, *a):
        pass

    def _send(self, code, payload, ctype='application/json; charset=utf-8'):
        data = payload if isinstance(payload, bytes) else json.dumps(payload, ensure_ascii=False, default=str).encode('utf-8')
        self.send_response(code)
        self.send_header('Content-Type', ctype)
        self.send_header('Cache-Control', 'no-store')
        self.end_headers()
        self.wfile.write(data)

    @staticmethod
    def _fresh():
        """Người dùng sửa thẳng file dữ liệu bằng Excel -> đọc lại trước khi làm tiếp."""
        if Handler.db is None:  # chưa có license lúc mở công cụ -> nạp dữ liệu khi license đã hợp lệ
            Handler.db = load_db()
        elif DATA_FILE.exists() and DATA_FILE.stat().st_mtime != DATA_MTIME[0]:
            Handler.db = load_db()

    def _host_ok(self):
        """Chỉ phục vụ yêu cầu gửi tới 127.0.0.1 / localhost: chặn trang web lạ dùng tên miền trỏ về máy này
        (DNS rebinding) để đọc dữ liệu."""
        host = (self.headers.get('Host') or '').strip().lower()
        host = host[1:host.index(']')] if host.startswith('[') and ']' in host else host.rsplit(':', 1)[0]
        if host in ('127.0.0.1', 'localhost', '::1'):
            return True
        self._send(403, {'error': 'Yêu cầu không hợp lệ.'})
        return False

    def _read_json(self, limit=MAX_BODY):
        """Body JSON của POST. Bắt buộc Content-Type application/json: trang web lạ không gửi được yêu cầu loại này
        tới công cụ mà không qua kiểm tra CORS của trình duyệt (chống CSRF). Sai -> ValueError."""
        if not (self.headers.get('Content-Type') or '').startswith('application/json'):
            raise ValueError('content-type')
        n = int(self.headers.get('Content-Length') or 0)
        if n < 0 or n > limit:
            raise ValueError('size')
        body = json.loads(self.rfile.read(n) or b'{}')
        if not isinstance(body, dict):
            raise ValueError('body')
        return body

    def _license_blocked(self):
        """License không hợp lệ -> trả 403 kèm trạng thái; dữ liệu không bị đọc/ghi."""
        st = LICENSE.check_license()
        if st.valid:
            return False
        self._send(403, {'error': st.message, 'license': license_info()})
        return True

    def _activate(self):
        # bắt buộc JSON: trang web lạ không gửi được yêu cầu này mà không qua kiểm tra CORS của trình duyệt
        try:
            body = self._read_json(limit=8192)
        except ValueError as e:
            return self._send(415 if str(e) == 'content-type' else 400, {'error': 'Yêu cầu không hợp lệ.'})
        try:
            st = LICENSE.activate(str(body.get('code') or ''))
        except licensing.ActivationError as e:
            return self._send(400, {'error': str(e), 'reason': e.reason, 'license': license_info()})
        except (ValueError, AttributeError):
            return self._send(400, {'error': 'Yêu cầu không hợp lệ.'})
        print('[License]', st.message)
        self._send(200, {'ok': True, 'license': license_info()})

    def do_GET(self):
        if not self._host_ok():
            return
        u = urlparse(self.path)
        if u.path == '/api/license':
            return self._send(200, license_info())
        if u.path == '/api/state':
            if self._license_blocked():
                return
            with LOCK:
                try:
                    self._fresh()
                except PermissionError:
                    pass
                except Exception as e:  # nạp dữ liệu lần đầu lỗi (thiếu file Excel gốc...)
                    return self._send(500, {'error': f'Không đọc được dữ liệu: {e}'})
                if Handler.db is None:
                    return self._send(500, {'error': 'Chưa đọc được dữ liệu — đóng file dữ liệu trong Excel rồi tải lại trang.'})
                return self._send(200, public_state(Handler.db))
        name = 'index.html' if u.path in ('/', '') else u.path.lstrip('/')
        f = (WEB_DIR / name).resolve()
        if WEB_DIR.resolve() not in f.parents or not f.is_file():
            return self._send(404, {'error': 'Không tìm thấy'})
        ctype = {'.html': 'text/html', '.css': 'text/css', '.js': 'application/javascript'}.get(f.suffix, 'application/octet-stream')
        self._send(200, f.read_bytes(), ctype + '; charset=utf-8')

    @staticmethod
    def _reload():
        """Bỏ mọi thay đổi chưa lưu trong bộ nhớ: đọc lại file dữ liệu (nạp lại khi cần nếu đọc lỗi)."""
        try:
            Handler.db = load_db() if DATA_FILE.exists() else None
        except Exception:
            Handler.db = None

    def do_POST(self):
        if not self._host_ok():
            return
        u = urlparse(self.path)
        if u.path == '/api/license/activate':
            return self._activate()
        if self._license_blocked():
            return
        try:
            body = self._read_json()
        except ValueError as e:
            return self._send(415 if str(e) == 'content-type' else 400, {'error': 'Yêu cầu không hợp lệ.'})
        try:
            with LOCK:
                self._fresh()
                res = handle(Handler.db, u.path, body)
            self._send(200, res)
        except PermissionError:
            with LOCK:
                self._reload()
            self._send(400, {'error': 'File Excel đang mở (file dữ liệu hoặc file xuất) — đóng file trong Excel rồi làm lại.'})
        except (ValueError, RuntimeError, StopIteration, KeyError) as e:
            with LOCK:  # thao tác bị từ chối: không để lại thay đổi dở dang trong bộ nhớ
                self._reload()
            self._send(400, {'error': str(e) or 'Dữ liệu không hợp lệ.'})
        except Exception as e:  # lỗi bất ngờ (file dữ liệu hỏng, đang lưu dở...): vẫn trả lời giao diện
            with LOCK:
                self._reload()
            self._send(500, {'error': f'Lỗi không mong muốn: {e}. Thao tác chưa được lưu; dữ liệu đã đọc lại từ file.'})


class Server(ThreadingHTTPServer):
    """Chiếm cổng độc quyền: trên Windows, SO_REUSEADDR mặc định cho 2 tiến trình nghe cùng 1 cổng (mỗi bên
    giữ dữ liệu riêng trong bộ nhớ -> ghi đè lẫn nhau)."""
    allow_reuse_address = False

    def server_bind(self):
        if hasattr(socket, 'SO_EXCLUSIVEADDRUSE'):
            self.socket.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
        super().server_bind()


def already_running(port):
    """Cổng đang bị chiếm bởi chính công cụ này (mở ở cửa sổ khác)?"""
    try:
        with urllib.request.urlopen(f'http://127.0.0.1:{port}/api/license', timeout=2) as r:
            return json.loads(r.read()).get('app') == APP_ID
    except Exception:
        return False


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    print('Đang mở công cụ Doanh thu - Đơn hàng - Công nợ...')
    lic = LICENSE.check_license()
    print('[License]', lic.message)
    if lic.valid:
        try:
            Handler.db = load_db()
        except Exception as e:
            print('(!) Không đọc được dữ liệu:', e)
            input('Bấm Enter để đóng...')
            return
    else:  # chưa/hết license: chỉ mở màn hình nhập mã, không đọc hay ghi dữ liệu
        print('Nhập mã kích hoạt trên giao diện. Mã thiết bị của máy này:', lic.device_id or '-')
        print('Dữ liệu trong thư mục "Dữ liệu" được giữ nguyên.')
    server = None
    for port in range(PORT, PORT + 20):
        try:
            server = Server(('127.0.0.1', port), Handler)
            break
        except OSError:
            if already_running(port):  # đã mở ở cửa sổ khác: dùng lại, không chạy bản thứ hai
                url = f'http://127.0.0.1:{port}/'
                print('Công cụ ĐANG MỞ ở một cửa sổ khác — mở lại giao diện đó:', url)
                if '--khong-mo' not in sys.argv:
                    webbrowser.open(url)
                return
            continue
    if server is None:
        print(f'(!) Không mở được cổng {PORT}–{PORT + 19} (đang bị chương trình khác dùng). Đóng bớt chương trình rồi mở lại.')
        input('Bấm Enter để đóng...')
        return
    url = f'http://127.0.0.1:{port}/'
    print('Giao diện:', url)
    print('File dữ liệu:', DATA_FILE)
    print('ĐỪNG ĐÓNG cửa sổ này trong lúc dùng công cụ. Dùng xong thì đóng cửa sổ này.')
    if '--khong-mo' not in sys.argv:  # chạy thử: không tự mở trình duyệt
        threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    server.serve_forever()


if __name__ == '__main__':
    main()
