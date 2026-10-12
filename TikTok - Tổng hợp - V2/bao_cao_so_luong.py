"""Báo cáo số lượng TikTok Shop và cập nhật tồn kho cho thư mục V2.

Chạy: python bao_cao_so_luong.py <file đơn hàng.xlsx> [--report-only]
Không đọc hay xuất số tiền; chỉ dùng đơn giá gốc bằng 0 để nhận diện quà tặng.
"""

import os
import re
import subprocess
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / 'thu-vien'))

from openpyxl import Workbook, load_workbook

import licensing  # cấp phép offline 48 giờ (bản sao trong thu-vien)
import ton_kho
from dinh_dang_bao_cao import apply_reference_styles, detail_sheet, grouped_sheet, overview_sheet

CANCELLED = {'Đã hủy', 'Cancelled'}
RETURN_TYPE = 'Return/Refund'
REPORT_DIR = ROOT / 'Báo cáo'

REQUIRED = {
    'Order ID', 'Order Status', 'Product Name', 'Quantity',
    'SKU Unit Original Price', 'Sku Quantity of return',
    'Cancelation/Return Type', 'Created Time', 'Shipped Time',
}


def number(value):
    if value in (None, ''):
        return 0
    return float(value)


def date_time(value):
    if isinstance(value, datetime):
        return value
    if not value:
        return None
    return datetime.strptime(str(value), '%d/%m/%Y %H:%M:%S')


def short_name(value):
    full = re.sub(r'\s+', ' ', str(value or '')).strip()
    # TikTok thêm mô tả quảng cáo sau dấu gạch dài; chấp nhận cả khoảng trắng không đều.
    head = re.split(r'\s+[–—]\s+', full, maxsplit=1)[0].strip()
    return head if len(head) >= 10 else full


def read_orders(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows)
    positions = {name: i for i, name in enumerate(headers)}
    missing = REQUIRED - positions.keys()
    if missing:
        wb.close()
        raise ValueError('File không đúng mẫu TikTok, thiếu cột: ' + ', '.join(sorted(missing)))
    result = []
    for values in rows:
        order = str(values[positions['Order ID']] or '')
        if not order.isdigit():
            continue

        def get(field):
            return values[positions[field]]

        qty = number(get('Quantity'))
        status = str(get('Order Status') or '')
        returned = number(get('Sku Quantity of return')) if get('Cancelation/Return Type') == RETURN_TYPE else 0
        result.append({
            'order': order,
            'created': date_time(get('Created Time')),
            'shipped': date_time(get('Shipped Time')),
            'status': status,
            'cancelled': status in CANCELLED,
            'prod': short_name(get('Product Name')),
            'qty': qty,
            'gift': number(get('SKU Unit Original Price')) == 0,
            'returned': returned,
        })
    wb.close()
    if not result:
        raise ValueError('File không có dòng đơn hàng hợp lệ.')
    if any(r['created'] is None for r in result):
        raise ValueError('Có đơn hàng thiếu ngày tạo đơn.')
    return result


def exported_at(path):
    match = re.search(r'(\d{4})-(\d{2})-(\d{2})-(\d{2})_(\d{2})', path.stem)
    return datetime(*map(int, match.groups())) if match else datetime.fromtimestamp(path.stat().st_mtime)


def product_mapping():
    wb, _, mapping = ton_kho.read_mapping(ROOT / ton_kho.MAP_FILE)
    wb.close()
    return mapping


def summarize(rows, mapping, group_by):
    """Gộp SKU trong cùng đơn trước khi xếp đơn vào nhóm SL 2/3/4/5."""
    groups = {}
    for row in rows:
        label, factor = mapping.get(ton_kho.key(row['prod']), ('', 1))
        factor = factor if factor > 0 else 1
        mapped = label or 'CHƯA GHÉP: ' + row['prod']
        group = mapped if group_by == 'ql' else row['prod']
        entry = groups.setdefault(group, {
            'name': group, 'ql': mapped, 'factor': factor,
            'orders': set(), 'shipped': 0, 'sold': 0, 'gift': 0,
            'returned': 0, 'failed': 0, 'by_order': defaultdict(float),
        })
        bottles = row['qty'] * factor
        if row['shipped']:
            entry['shipped'] += bottles
        if row['cancelled']:
            if row['shipped']:
                entry['failed'] += bottles
        else:
            entry['orders'].add(row['order'])
            if row['gift']:
                entry['gift'] += bottles
            else:
                entry['sold'] += bottles
                entry['by_order'][row['order']] += row['qty']
            entry['returned'] += row['returned'] * factor
    return groups



def create_report(rows, src, exported, out):
    mapping = product_mapping()
    dates = [r['created'] for r in rows]
    period = f'Đơn đặt từ {min(dates):%d/%m/%Y} đến {max(dates):%d/%m/%Y}'
    by_ql = summarize(rows, mapping, 'ql')
    by_product = summarize(rows, mapping, 'product')
    wb = Workbook()
    overview_sheet(wb, rows, by_ql, period, exported)
    grouped_sheet(wb, 'Theo mặt hàng QL', by_ql, period, by_ql=True)
    grouped_sheet(wb, 'Theo sản phẩm TikTok', by_product, period)
    detail_sheet(wb, rows, mapping, period)
    apply_reference_styles(wb)
    wb.active = 0
    wb.save(out)
    return by_ql


def recalc_with_excel(path):
    """Tính công thức tồn kho trong Excel để xem được trên điện thoại."""
    command = ('$excel = New-Object -ComObject Excel.Application; $excel.DisplayAlerts = $false; '
               'try { $book = $excel.Workbooks.Open($env:V2_QL_PATH); $excel.CalculateFull(); '
               '$book.Save(); $book.Close($false) } '
               'finally { $excel.Quit(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($excel) }')
    result = subprocess.run(['powershell', '-NoProfile', '-NonInteractive', '-Command', command],
                            env=dict(os.environ, V2_QL_PATH=str(path)), capture_output=True)
    return result.returncode == 0


def pick_file():
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    return filedialog.askopenfilename(title='Chọn file Tất cả đơn hàng từ TikTok',
                                      filetypes=[('Excel', '*.xlsx')])


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    if not licensing.ensure_licensed('Báo cáo số lượng TikTok – V2'):
        return 1
    args = [x for x in sys.argv[1:] if x != '--report-only']
    report_only = '--report-only' in sys.argv[1:]
    src = Path(args[0]) if args else Path(pick_file())
    if not src.is_file():
        print('Chưa chọn file đơn hàng hợp lệ.')
        return 1
    if ton_kho.ql_path(ROOT) is None and not report_only:
        print('Không tìm thấy QL- kênh Online.xlsx trong thư mục V2.')
        return 1
    rows = read_orders(src)
    exported = exported_at(src)
    dates = [r['created'] for r in rows]
    REPORT_DIR.mkdir(exist_ok=True)
    out = REPORT_DIR / (f'Báo cáo số lượng TikTok {min(dates):%d.%m.%Y} - '
                        f'{max(dates):%d.%m.%Y} (xuất {exported:%d.%m.%Y %Hh%M}) - theo mẫu.xlsx')
    create_report(rows, src, exported, out)
    print('Đã tạo báo cáo:', out)
    if not report_only:
        for message in ton_kho.cap_nhat(rows, exported, ROOT, recalc_with_excel):
            print(message)
    return 0


if __name__ == '__main__':
    try:
        raise SystemExit(main())
    except (ValueError, PermissionError) as exc:
        print('Lỗi:', exc)
        raise SystemExit(1)
