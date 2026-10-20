"""Định dạng V2 theo mẫu Báo cáo TikTok của bản gốc (Arial, bảng Excel có viền)."""

import math
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.hyperlink import Hyperlink

import ton_kho

NAVY, BLUE, LIGHT, BAND = '1F3864', '1F4E78', 'DDEBF7', 'F5F8FC'
TOTAL, GREEN, TILE, RED, GRAY, LINE = 'D9E1F2', 'E2EFDA', 'F2F6FC', 'C00000', '595959', 'BFBFBF'
NUM = '#,##0;-#,##0;"-"'
PCT = '0.0%;-0.0%;"-"'
THIN = Side(style='thin', color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_BOX = Border(left=THIN, right=THIN, top=Side(style='medium', color=NAVY),
                   bottom=Side(style='medium', color=NAVY))
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
RIGHT = Alignment(horizontal='right', vertical='center')


def _copy_format(target, source):
    """Chỉ chép kiểu trình bày; giữ nguyên giá trị và định dạng số lượng."""
    for attribute in ('font', 'fill', 'border', 'alignment', 'protection'):
        setattr(target, attribute, copy(getattr(source, attribute)))


def apply_reference_styles(wb):
    """Dùng các kiểu ô lấy trực tiếp từ file mẫu, đã lưu riêng không kèm dữ liệu."""
    path = Path(__file__).with_name('Mau-dinh-dang.xlsx')
    if not path.exists():
        raise ValueError('Thiếu Mau-dinh-dang.xlsx. Chép đầy đủ thư mục V2 để dùng format mẫu.')
    reference = load_workbook(path)
    try:
        source = reference['Tong-quan']
        overview = wb.worksheets[0]
        for target, origin in [('B2', 'B2'), ('B3', 'B3'), ('B15', 'B15'), ('B16', 'B16'),
                               ('B23', 'B23'), ('B28', 'B29')]:
            _copy_format(overview[target], source[origin])
        for top_row in (5, 8, 11):
            for col in range(2, 6):
                highlighted = top_row == 5 and col == 2
                _copy_format(overview.cell(top_row, col), source['B5' if highlighted else 'C5'])
                value_style = ('B6' if highlighted else 'D6' if top_row == 8 else
                               'D12' if top_row == 11 and col == 3 else 'C6')
                _copy_format(overview.cell(top_row + 1, col), source[value_style])
        for col in range(2, 7):
            _copy_format(overview.cell(17, col), source['B17'])
            for row in range(18, 21):
                origin = ('F' if col == 6 else 'B' if col == 2 else 'E' if col == 5 else 'C') + ('19' if row % 2 else '18')
                _copy_format(overview.cell(row, col), source[origin])
            origin = 'F21' if col == 6 else 'B21' if col == 2 else 'E21' if col == 5 else 'C21'
            _copy_format(overview.cell(21, col), source[origin])
        for row in (24, 25, 26):
            _copy_format(overview.cell(row, 2), source['B24'])
            _copy_format(overview.cell(row, 3), source['C24'])
        for row in (2, 5, 6, 7, 8, 9, 10, 11, 12, 13, 17, 18, 19, 20, 21):
            overview.row_dimensions[row].height = source.row_dimensions[row].height

        source = reference['Theo-mat-hang']
        for ws in wb.worksheets[1:3]:
            by_ql = ws.max_column == 12
            name_columns = {2} if by_ql else {2, 3}
            shipped_col = 4 if by_ql else 5
            returned_cols = {7, 8} if by_ql else {8, 9}
            _copy_format(ws['A1'], source['A1'])
            _copy_format(ws['A2'], source['A2'])
            for cell in ws[4]:
                if not isinstance(cell, MergedCell):
                    origin = 'D4' if cell.column >= shipped_col else 'A4'
                    _copy_format(cell, source[origin])
            for cell in ws[5]:
                if not isinstance(cell, MergedCell):
                    _copy_format(cell, source['D5'])
            last = range_boundaries(ws.auto_filter.ref)[3]
            for row in range(6, last + 1):
                unknown = any(str(ws.cell(row, col).value or '').startswith('CHƯA GHÉP:')
                              for col in name_columns)
                band = '7' if (row - 6) % 2 else '6'
                for col in range(1, ws.max_column + 1):
                    origin = ('A' if col == 1 else 'B' if col in name_columns else
                              'K' if col == shipped_col else 'J' if col in returned_cols else 'C') + band
                    _copy_format(ws.cell(row, col), source[origin])
                    if unknown:
                        ws.cell(row, col).fill = fill('FFF2CC')
            for col in range(1, ws.max_column + 1):
                origin = ('A' if col == 1 else 'B' if col in name_columns else
                          'J' if col in returned_cols else 'C') + '21'
                _copy_format(ws.cell(last + 1, col), source[origin])
            for row in (1, 4, 5):
                ws.row_dimensions[row].height = source.row_dimensions[row].height

        source = reference['Du-lieu']
        detail = wb.worksheets[3]
        for cell in detail[1]:
            _copy_format(cell, source['A1'])
        body_font = copy(source['A2'].font)
        for row in detail.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
        detail.row_dimensions[1].height = source.row_dimensions[1].height
    finally:
        reference.close()


def font(size=10, bold=False, color='262626', italic=False, underline=None):
    return Font(name='Arial', size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(color):
    return PatternFill('solid', fgColor=color)


def printable(value):
    return int(value) if value == int(value) else value


def row_height(text, width, base=15):
    return base * max(1, math.ceil(len(text) / (width * 0.95))) + 4


def setup_print(ws, title_rows=None):
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    if title_rows:
        ws.print_title_rows = title_rows


def title(ws, text, subtitle):
    ws['A1'] = text
    ws['A1'].font = font(14, True, NAVY)
    ws['A2'] = subtitle
    ws['A2'].font = font(9, color=GRAY, italic=True)
    ws.row_dimensions[1].height = 24


def grouped_sheet(wb, sheet_name, groups, period, by_ql=False):
    ws = wb.create_sheet(sheet_name)
    ws.sheet_properties.tabColor = BLUE
    ws.sheet_view.showGridLines = False
    title(ws, 'SỐ LƯỢNG THEO MẶT HÀNG' if by_ql else 'SỐ LƯỢNG THEO SẢN PHẨM TIKTOK',
          f'{period}  •  Sắp xếp theo số chai đã gửi  •  Nhóm SL tính theo SKU bán trong từng đơn')
    columns = [('stt', 'STT', 6), ('name', 'Mặt hàng' if by_ql else 'Sản phẩm TikTok', 44)]
    if not by_ql:
        columns.append(('ql', 'Mặt hàng QL', 28))
    columns += [('orders', 'Số đơn', 10), ('shipped', 'Chai đã gửi', 11),
                ('sold', 'Chai bán', 10), ('gift', 'Chai tặng', 10),
                ('returned', 'Khách trả lại', 11), ('failed', 'Giao thất bại', 11),
                ('size2', 'Đơn SL bán = 2', 12), ('size3', 'Đơn SL bán = 3', 12),
                ('size4', 'Đơn SL bán = 4', 12), ('size5', 'Đơn SL bán = 5', 12)]
    indexes = {key: i for i, (key, _, _) in enumerate(columns, 1)}
    bands = [('SỐ LƯỢNG (CHAI)', 'shipped', 'gift'),
             ('HÀNG HOÀN VỀ (CHAI)', 'returned', 'failed'),
             ('SỐ ĐƠN THEO SỐ LƯỢNG SKU BÁN', 'size2', 'size5')]
    grouped = set()
    for label, start_key, end_key in bands:
        first, last = indexes[start_key], indexes[end_key]
        for col in range(first, last + 1):
            grouped.add(col)
            ws.cell(4, col).fill, ws.cell(4, col).border = fill(BLUE), BOX
        ws.cell(4, first, label).font = font(10, True, 'FFFFFF')
        ws.cell(4, first).alignment = CENTER
        ws.merge_cells(start_row=4, start_column=first, end_row=4, end_column=last)
    for col, (_, label, width) in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
        if col in grouped:
            cell = ws.cell(5, col, label)
            cell.font, cell.fill, cell.alignment, cell.border = font(10, True, NAVY), fill(LIGHT), CENTER, BOX
        else:
            for row in (4, 5):
                ws.cell(row, col).fill, ws.cell(row, col).border = fill(BLUE), BOX
            cell = ws.cell(4, col, label)
            cell.font, cell.alignment = font(10, True, 'FFFFFF'), CENTER
            ws.merge_cells(start_row=4, start_column=col, end_row=5, end_column=col)
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 42
    entries = sorted(groups.values(), key=lambda entry: (-entry['shipped'], entry['name']))
    first, last = 6, 5 + len(entries)
    for row, entry in enumerate(entries, first):
        values = dict(stt=row - first + 1, name=entry['name'], ql=entry['ql'],
                      orders=len(entry['orders']))
        values.update({key: printable(entry[key]) for key in ('shipped', 'sold', 'gift', 'returned', 'failed')})
        values.update({f'size{size}': sum(qty == size for qty in entry['by_order'].values())
                       for size in range(2, 6)})
        unknown = entry['ql'].startswith('CHƯA GHÉP:')
        for col, (key, _, _) in enumerate(columns, 1):
            cell = ws.cell(row, col, values[key])
            cell.border = BOX
            cell.alignment = LEFT if key in ('name', 'ql') else (CENTER if key == 'stt' else RIGHT)
            if isinstance(values[key], (int, float)):
                cell.number_format = NUM
            cell.font = font(10, key == 'shipped', RED if key in ('returned', 'failed') else '262626')
            if unknown:
                cell.fill = fill('FFF2CC')
            elif key == 'shipped':
                cell.fill = fill(GREEN)
            elif (row - first) % 2:
                cell.fill = fill(BAND)
        height = row_height(entry['name'], 44)
        if not by_ql:
            height = max(height, row_height(entry['ql'], 28))
        ws.row_dimensions[row].height = height
    total_row = last + 1
    ws.cell(total_row, indexes['name'], 'CỘNG THEO DÒNG')
    for key, col in indexes.items():
        if key in ('stt', 'name', 'ql'):
            continue
        ws.cell(total_row, col, sum(ws.cell(row, col).value or 0 for row in range(first, last + 1)))
    for col, (key, _, _) in enumerate(columns, 1):
        cell = ws.cell(total_row, col)
        cell.border, cell.fill = TOTAL_BOX, fill(TOTAL)
        cell.font = font(10, True, RED if key in ('returned', 'failed') else '262626')
        cell.alignment = LEFT if key in ('name', 'ql') else RIGHT
        if isinstance(cell.value, (int, float)):
            cell.number_format = NUM
    ws.row_dimensions[total_row].height = 22
    note = total_row + 2
    ws.cell(note, 2, 'Một đơn có thể chứa nhiều mặt hàng: cột Số đơn ở dòng CỘNG THEO DÒNG có thể tính một đơn nhiều lần.').font = font(9, color=GRAY, italic=True)
    ws.cell(note + 1, 2, 'Đơn SL bán = 2/3/4/5: gộp SKU bán của cùng mặt hàng trong cùng đơn; không cộng quà tặng và không nhân hệ số combo.').font = font(9, color=GRAY, italic=True)
    ws.freeze_panes = f'{get_column_letter(indexes["orders"])}6'
    ws.auto_filter.ref = f'A5:{get_column_letter(len(columns))}{last}'
    setup_print(ws, '4:5')
    ws.print_area = f'A1:{get_column_letter(len(columns))}{note + 1}'


def overview_sheet(wb, rows, groups, period, exported):
    ws = wb.active
    ws.title = 'Tổng quan'
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 2
    for col in 'BCDE':
        ws.column_dimensions[col].width = 28
    ws.column_dimensions['F'].width = 66
    ws['B2'] = 'BÁO CÁO SỐ LƯỢNG TIKTOK SHOP'
    ws['B2'].font = font(18, True, NAVY)
    ws.row_dimensions[2].height = 30
    ws['B3'] = f'{period}  •  Dữ liệu xuất lúc {exported:%d/%m/%Y %H:%M}  •  Đơn vị: chai / đơn'
    ws['B3'].font = font(10, color=GRAY, italic=True)
    valid = {row['order'] for row in rows if not row['cancelled']}
    shipped = {row['order'] for row in rows if row['shipped']}
    failed = {row['order'] for row in rows if row['cancelled'] and row['shipped']}
    before = {row['order'] for row in rows if row['cancelled'] and not row['shipped']}
    returned = {row['order'] for row in rows if not row['cancelled'] and row['returned'] > 0}
    total = lambda key: printable(sum(entry[key] for entry in groups.values()))
    tiles = [
        [('TỔNG CHAI ĐÃ GỬI', total('shipped'), GREEN, NAVY),
         ('Chai bán', total('sold'), TILE, NAVY),
         ('Chai tặng', total('gift'), TILE, NAVY),
         ('Số đơn (không tính đơn hủy)', len(valid), TILE, NAVY)],
        [('Chai khách trả lại', total('returned'), TILE, RED),
         ('Chai giao thất bại', total('failed'), TILE, RED),
         ('Đơn khách trả hàng', len(returned), TILE, RED),
         ('Đơn giao thất bại', len(failed), TILE, RED)],
        [('Đơn đã gửi hàng', len(shipped), TILE, NAVY),
         ('Đơn hủy trước khi gửi', len(before), TILE, GRAY),
         ('Số mặt hàng QL có đơn hợp lệ', sum(bool(entry['orders']) for entry in groups.values()), TILE, NAVY),
         ('Số sản phẩm TikTok trong file', len({row['prod'] for row in rows}), TILE, NAVY)],
    ]
    for top_row, line in zip((5, 8, 11), tiles):
        for col, (label, value, background, color) in enumerate(line, 2):
            top = ws.cell(top_row, col, label)
            val = ws.cell(top_row + 1, col, value)
            top.font, top.fill = font(9, True, GRAY), fill(background)
            top.alignment = Alignment(horizontal='left', vertical='bottom', indent=1)
            val.font, val.fill = font(16, True, color), fill(background)
            val.alignment = Alignment(horizontal='left', vertical='center', indent=1)
            val.number_format = NUM
            top.border = Border(left=THIN, right=THIN, top=THIN)
            val.border = Border(left=THIN, right=THIN, bottom=THIN)
        ws.row_dimensions[top_row].height = 20
        ws.row_dimensions[top_row + 1].height = 34
        ws.row_dimensions[top_row + 2].height = 8
    ws['B15'] = 'SỐ LƯỢNG THEO PHÂN LOẠI ĐƠN'
    ws['B15'].font = font(12, True, NAVY)
    ws['B16'] = 'Số chai đã gửi bao gồm cả quà tặng; đơn hủy sau khi gửi được xếp vào giao thất bại'
    ws['B16'].font = font(9, color=GRAY, italic=True)
    headers = ['Phân loại', 'Số đơn', 'Chai đã gửi', 'Tỷ trọng đơn', 'Ý nghĩa']
    for col, label in enumerate(headers, 2):
        cell = ws.cell(17, col, label)
        cell.font, cell.fill, cell.alignment, cell.border = font(10, True, 'FFFFFF'), fill(BLUE), CENTER, BOX
    ws.row_dimensions[17].height = 24
    count = len(valid | failed | before)
    classifications = [
        ('Không bị hủy', len(valid), total('shipped') - total('failed'), 'Đơn đang xử lý, đang giao, đã giao hoặc khách trả hàng'),
        ('Giao thất bại', len(failed), total('failed'), 'Đơn bị hủy sau khi đã gửi hàng; ghi Nhập khi hàng thực tế quay về kho'),
        ('Hủy trước khi gửi', len(before), 0, 'Hàng chưa rời kho, không trừ tồn kho'),
    ]
    for row, (label, qty, bottles, meaning) in enumerate(classifications, 18):
        for col, value in enumerate((label, qty, bottles, qty / count if count else 0, meaning), 2):
            cell = ws.cell(row, col, value)
            cell.font, cell.border = font(10, col == 2), BOX
            cell.alignment = LEFT if col in (2, 6) else RIGHT
            if isinstance(value, (int, float)):
                cell.number_format = PCT if col == 5 else NUM
            if row % 2:
                cell.fill = fill(BAND)
        ws.row_dimensions[row].height = 22
    for col, value in enumerate(('TỔNG', count, total('shipped'), 1 if count else 0, ''), 2):
        cell = ws.cell(21, col, value)
        cell.font, cell.fill, cell.border = font(10, True), fill(TOTAL), TOTAL_BOX
        cell.alignment = LEFT if col in (2, 6) else RIGHT
        if isinstance(value, (int, float)):
            cell.number_format = PCT if col == 5 else NUM
    ws.row_dimensions[21].height = 22
    ws['B23'] = 'XEM CHI TIẾT'
    ws['B23'].font = font(12, True, NAVY)
    for row, (sheet, description) in enumerate([
        ('Theo mặt hàng QL', 'Số chai bán / tặng / hoàn và nhóm đơn SL 2/3/4/5 theo loại hàng trong QL'),
        ('Theo sản phẩm TikTok', 'Xem riêng từng tên sản phẩm TikTok, kể cả combo'),
        ('Dữ liệu số lượng', 'Toàn bộ dòng SKU để lọc, tra cứu và đối chiếu mã đơn'),
    ], 24):
        cell = ws.cell(row, 2, f'→ {sheet}')
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A1")
        cell.font = font(10, True, '0563C1', underline='single')
        ws.cell(row, 3, description).font = font(10, color=GRAY)
    ws['B28'] = 'Lưu ý: nhóm đơn SL 2/3/4/5 tính theo SKU bán, không cộng quà tặng. Số chai thực tế được nhân hệ số combo.'
    ws['B28'].font = font(9, True, RED, italic=True)
    setup_print(ws)
    ws.print_area = 'A1:F28'


def detail_sheet(wb, rows, mapping, period):
    ws = wb.create_sheet('Dữ liệu số lượng')
    ws.sheet_properties.tabColor = '7F7F7F'
    headers = ['Mã đơn', 'Ngày đặt', 'Ngày gửi', 'Trạng thái', 'Sản phẩm TikTok',
               'Mặt hàng QL', 'SL SKU', 'Hệ số chai', 'Số chai', 'Bán / Tặng', 'SL hoàn', 'Giao thất bại']
    widths = [20, 18, 18, 16, 50, 34, 10, 11, 11, 12, 11, 17]
    ws.append(headers)
    for row in rows:
        label, factor = mapping.get(ton_kho.key(row['prod']), ('', 1))
        factor = factor if factor > 0 else 1
        ws.append([row['order'], row['created'], row['shipped'], row['status'], row['prod'],
                   label or 'CHƯA GHÉP', printable(row['qty']), factor,
                   printable(row['qty'] * factor), 'Tặng' if row['gift'] else 'Bán',
                   printable(row['returned'] * factor), 'Có' if row['cancelled'] and row['shipped'] else ''])
    body_font = font()
    for col, width in enumerate(widths, 1):
        cell = ws.cell(1, col)
        cell.font, cell.fill, cell.alignment, cell.border = font(10, True, 'FFFFFF'), fill(BLUE), CENTER, BOX
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            if isinstance(cell.value, (int, float)):
                cell.number_format = NUM
        row[1].number_format = row[2].number_format = 'dd/mm/yyyy hh:mm'
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:L{ws.max_row}'
    setup_print(ws, '1:1')
